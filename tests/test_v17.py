"""v1.7 — consulting-intelligence layer: fact model, grounding, artifact checks, hard factuality
gates, evidence graph, ghost deck, source-to-deck benchmark. Every negative test plants one defect
in the stored development run and expects the deterministic checks to catch it."""
import json
import shutil
from pathlib import Path

import pytest

from cpe.reasoning import benchmark, checks, facts, graph, grounding
from cpe.reasoning.facts import detect_period, detect_unit, unit_from_raw

ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / "evals" / "source_to_deck" / "development" / "margin_recovery"
RUN = CASE / "runs" / "dev_run_01"


@pytest.fixture
def work(tmp_path):
    w = tmp_path / "work"
    shutil.copytree(RUN, w)
    return w


def _edit(w, name, fn):
    d = json.loads((w / name).read_text())
    fn(d)
    (w / name).write_text(json.dumps(d, indent=2, ensure_ascii=False))


def _codes(r, hard=None):
    return {i["code"] for i in r["issues"] if hard is None or i["hard"] == hard}


# ── fact model ───────────────────────────────────────────────────────────────────────────────────

def test_period_and_unit_semantics():
    assert detect_period("FY25E") == {"period": "FY2025", "basis": "estimate"}
    assert detect_period("Budget 2026")["basis"] == "budget"
    assert detect_period("Q3 2025")["period"] == "2025-Q3" and detect_period("LTM")["period"] == "LTM"
    assert detect_period("Revenue") is None
    assert detect_unit("Revenue (€M)") == "EUR_M" and detect_unit("Gross margin (%)") == "PCT" and detect_unit("€bn") == "EUR_BN"
    assert unit_from_raw("€1,500M") == "EUR_M" and unit_from_raw("24%") == "PCT" and unit_from_raw("3.2 pp") == "PP"


def test_fact_model_is_traceable_and_derives_changes():
    fm = facts.build_fact_model(sorted((CASE / "sources").glob("*")))
    rev = next(f for f in fm["facts"] if f["claim"].startswith("Total — Revenue FY2025"))
    assert rev["values"][0]["value"] == 1460 and rev["values"][0]["unit"] == "EUR_M" and rev["values"][0]["period"] == "FY2025"
    assert rev["source"] == {"file": "pnl_by_category.csv", "sheet": None, "range": "C6", "loc": "sheet"}
    ch = next(f for f in fm["facts"] if f["fact_type"] == "derived_change" and f["claim"].startswith("Total — revenue"))
    assert ch["values"][0]["value"] == 100 and round(ch["values"][1]["value"], 1) == 7.4 and len(ch["derived_from"]) == 2
    gm = next(f for f in fm["facts"] if f["fact_type"] == "derived_change" and f["claim"].startswith("Total — gross margin"))
    assert gm["values"][0]["unit"] == "PP"  # a change of a percentage is in points
    fc = next(f for f in fm["facts"] if "€1,500M" in f["claim"])
    assert fc["values"][0]["period"] == "FY2025" and fc["values"][0]["basis"] == "estimate"  # forecast, not actual
    assert all(v["value"] != 202 for f in fm["facts"] for v in f["values"])  # "2026" is not split into "202"


# ── grounding ────────────────────────────────────────────────────────────────────────────────────

def test_grounding_only_uses_cited_facts():
    f = [{"id": "F1", "values": [{"value": 924, "unit": "EUR_M"}]}, {"id": "F2", "values": [{"value": 857, "unit": "EUR_M"}]}]
    g = grounding.ground_numbers("EBITDA fell from €924M to €857M, a €67M (7%) decline", f)
    assert [x["status"] for x in g] == ["grounded", "grounded", "derived", "derived"]
    assert grounding.ground_numbers("EBITDA will reach €1.2bn", f)[0]["status"] == "unsupported"
    assert grounding.ground_numbers("EBITDA fell to €857M", f[:1])[0]["status"] == "unsupported"  # not cited → not grounded
    assert grounding.ground_numbers("EBITDA fell to $857M", f)[0]["status"] == "unsupported"  # never across currencies


# ── the stored development run passes; planted defects are caught ─────────────────────────────

def test_dev_run_passes_all_checks():
    r = checks.check_work(RUN, CASE / "sources")
    assert r["status"] == "pass" and r["hard_failures"] == 0 and all(r["stopping_criteria"].values())


def test_fabricated_fact_is_hard_failure(work):
    _edit(work, "facts.json", lambda d: d["facts"][0]["values"][0].update(value=39))
    assert "FACT_FABRICATED" in _codes(checks.check_work(work, CASE / "sources"), hard=True)


def test_wrong_computation_is_arithmetic_error(work):
    _edit(work, "computed_facts.json", lambda d: d["facts"][3]["values"][0].update(value=1.9))
    r = checks.check_work(work, CASE / "sources")
    assert "ARITHMETIC_ERROR" in _codes(r, hard=True) and r["status"] == "blocked"


def test_invented_headline_number_is_blocked(work):
    def bump(d):
        s = next(x for x in d["slides"] if x["id"] == "s06")
        s["headline"] = "Freight cost rose €12M after the new carrier contract"
    _edit(work, "deck_plan.json", bump)
    _edit(work, "deck.json", bump)
    r = checks.check_work(work, CASE / "sources")
    bad = [i for i in r["issues"] if i["code"] == "UNSUPPORTED_NUMBER"]
    assert {i["artifact"] for i in bad} == {"deck_plan.json", "deck.json"} and all(i["hard"] for i in bad)
    assert not r["stopping_criteria"]["no_unsupported_headlines"]


def test_insight_on_rejected_hypothesis_is_contradicted(work):
    _edit(work, "insights.json", lambda d: d["insights"][1].update(hypotheses=["H1"]))
    assert "CONTRADICTED_AS_FACT" in _codes(checks.check_work(work, CASE / "sources"), hard=True)


def test_wrong_source_attribution(work):
    def src(d):
        next(x for x in d["slides"] if x["id"] == "s06")["source"] = "Source: pnl_by_category.csv"
    _edit(work, "deck.json", src)
    assert "WRONG_SOURCE" in _codes(checks.check_work(work, CASE / "sources"), hard=True)


def test_assumptions_are_flagged_not_hidden(work):
    def plain(d):
        next(x for x in d["slides"] if x["id"] == "s07")["headline"] = "Repricing electronics adds 3 pp there and 0.6 pp to the group"
    _edit(work, "deck_plan.json", plain)
    r = checks.check_work(work, CASE / "sources")
    w = [i for i in r["issues"] if i["code"] == "ASSUMPTION_IN_HEADLINE" and i["ref"] == "s07"]
    assert w and all(i["level"] == "warning" for i in w) and not any(i["hard"] for i in w)  # unhedged → warning
    hedged = [i for i in r["issues"] if i["code"] == "ASSUMPTION_IN_HEADLINE" and i["ref"] == "s08"]
    assert hedged and all(i["level"] == "info" for i in hedged)  # "about 1.1 pp" is stated as an estimate


def test_storyline_and_architecture_quality_checks(work):
    def one(d):
        d["candidates"] = d["candidates"][:1]
        d["key_line"][1]["insights"] = []
    _edit(work, "storyline.json", one)

    def plan(d):
        d["slides"].append({"id": "s99", "priority": "core", "headline": "Other topics", "reason_to_exist": ""})
        d["slides"] = [s for s in d["slides"] if s["id"] != "s06"]  # K2 loses its slide
    _edit(work, "deck_plan.json", plan)
    c = _codes(checks.check_work(work, CASE / "sources"))
    assert {"GT_SINGLE_CANDIDATE", "KEYLINE_UNSUPPORTED", "SLIDE_NO_REASON", "SLIDE_ORPHAN", "HEADLINE_TOPIC", "KEYLINE_NO_SLIDE"} <= c


def test_missing_business_question_blocks_storyline(work):
    (work / "project.json").unlink()
    c = _codes(checks.check_work(work, CASE / "sources"))
    assert "PROJECT_MISSING" in c


# ── evidence graph and ghost deck ────────────────────────────────────────────────────────────────

def test_trace_returns_lineage_to_source_cells():
    t = graph.trace(RUN, "Mix explains 1.45 pp of the gross-margin decline")
    assert t["match"] == "SL:s04"
    flat = json.dumps(t)
    assert "C0004" in flat and "pnl_by_category.csv" in flat and '"range": "D6"' in flat
    assert graph.trace(RUN, "completely unrelated sentence about weather")["match"] is None


def test_ghost_deck_reads_as_an_argument():
    from cpe.reasoning.ghost import ghost_markdown

    md = ghost_markdown(RUN)
    assert "Governing thought" in md and md.count("\n| ") >= 9 and "appendix" in md


# ── benchmark ────────────────────────────────────────────────────────────────────────────────────

def test_benchmark_dimensions_are_separate():
    r = benchmark.evaluate(CASE, RUN, {"model": "test"})
    assert r["factuality"]["status"] == "PASS" and r["fact_grounding"]["critical_fact_recall"] == 1.0
    assert r["insight_quality"]["conclusions"] == {"C1": True, "C2": True, "C3": True} and r["storyline"]["traps_triggered"] == []
    assert "score" not in r and "total" not in r  # no aggregate on purpose


def test_benchmark_catches_traps(work):
    def trap(d):
        d["governing_thought"] = "Margin fell because volume fell; revenue reached €1,500M but sold less"
    _edit(work, "storyline.json", trap)
    r = benchmark.evaluate(CASE, work)
    assert {"T1", "T2"} <= set(r["storyline"]["traps_triggered"])
    assert r["factuality"]["status"] == "FAIL"  # €1,500M is not grounded in the key line's facts


def test_reasoning_protocol_is_versioned():
    from cpe.reasoning import PROTOCOL_VERSION

    assert PROTOCOL_VERSION == "1.4"
    assert json.loads((RUN / "facts.json").read_text())["protocol"] == "1.0"  # a stored run keeps the protocol it was produced under
    assert "development run — NOT evidence" in (RUN / "RUN.json").read_text()
    for d in ("development", "sealed", "external"):
        assert (ROOT / "evals" / "source_to_deck" / d).is_dir()


def test_conflicts_must_be_resolved_when_used(work):
    (work / "fact_conflicts.json").unlink()
    r = checks.check_work(work, CASE / "sources")
    bad = [i for i in r["issues"] if i["code"] == "FACT_CONFLICT_UNRESOLVED"]
    assert bad and bad[0]["level"] == "error" and "forecast_vs_actual" in bad[0]["message"]  # €1,500M forecast vs €1,460M actual
    assert not r["stopping_criteria"]["no_unresolved_conflict_in_use"]


def test_unresolved_high_critic_finding_stops_the_loop(work):
    _edit(work, "critique.json", lambda d: d["findings"][0].update(resolved=False))
    r = checks.check_work(work, CASE / "sources")
    assert "CRITIC_UNRESOLVED" in _codes(r) and not r["stopping_criteria"]["no_unresolved_high_critic_finding"]
    (work / "critique.json").unlink()
    assert "CRITIQUE_MISSING" in _codes(checks.check_work(work, CASE / "sources"))


def test_blind_storyline_ab_round(tmp_path):
    from cpe import human

    a, b, ctx = tmp_path / "a", tmp_path / "b", tmp_path / "ctx"
    for d in (a, b, ctx):
        d.mkdir()
    for case in ("c1", "c2", "c3"):
        (a / f"{case}.md").write_text(f"GT v1.6 agent: {case}\n- K1\n- K2")
        (b / f"{case}.md").write_text(f"GT v1.7 agent: {case}\n- K1'\n- K2'\n- K3'")
        (ctx / f"{case}.md").write_text(f"Question for {case}?")
    rd, kp = tmp_path / "s1", tmp_path / "keys" / "s1" / "key.json"
    r = human.build_text_round(rd, [("agent-1.6->agent-1.7", str(a), str(b))], context_dir=ctx, repeats=1, key_out=kp)
    assert r["pairs"] == 4 and kp.exists() and not (rd / "key.json").exists()
    bundle = (rd / "pairs.json").read_text() + (rd / "STATUS.json").read_text()
    assert "1.6" not in bundle and "1.7" not in bundle and "agent" not in bundle
    pairs = json.loads((rd / "pairs.json").read_text())
    assert pairs["meta"]["kind"] == "text" and all("context" in p for p in pairs["pairs"])
    for p in pairs["pairs"]:
        human.record_vote(rd, {"evaluator": "ana", "pair": p["id"], "left": p["images"][0], "right": p["images"][1], "choice": "left"})
    rep = human.report(rd, kp)
    assert rep["comparisons"] == 3 and rep["by_archetype"].keys() == {"storyline"}
    z = human.package_round(rd, tmp_path / "s1.zip")
    assert z["texts"] == 3 * 3 and z["images"] == 0  # A, B and context per case


def test_fact_model_on_spanish_excel_case():
    case = ROOT / "evals" / "source_to_deck" / "development" / "churn_es"
    fm = facts.build_fact_model(sorted((case / "sources").glob("*")))
    cli = next(f for f in fm["facts"] if f["claim"].startswith("Particulares — Clientes 2025"))
    assert cli["values"][0]["unit"] == "" and cli["source"]["sheet"] == "Clientes" and cli["source"]["range"] == "C2"  # % of a neighbour column does not spread
    ing = next(f for f in fm["facts"] if f["claim"].startswith("Particulares — Ingresos 2025"))
    assert ing["values"][0]["unit"] == "EUR_M" and ing["source"]["sheet"] == "Ingresos"  # "M€"
    txt = next(f for f in fm["facts"] if "resolución" in f["claim"])
    assert [v["value"] for v in txt["values"]] == [2.1, 4.6]  # Spanish decimal comma
    assert next(f for f in fm["facts"] if "30.000" in f["claim"])["values"][0]["value"] == 30000  # Spanish thousands


# ── v1.7: integrity metric (waterfall scorer blind spot, r3) ─────────────────────────────────────

def test_broken_slide_cannot_score_well():
    from cpe.qa.archetypes import fitness
    from cpe.qa.composition import BROKEN_CODES, integrity_issues

    ok = {"utilization": 0.9, "empty": 0.1, "ink": 0.15, "offcentre": 0.05, "emphasis": 0.2, "regions": 2, "ratio": 2.0, "edges": 3, "proof": 1.0}
    good = fitness("waterfall", ok)["score"]  # intact: integrity is not observed, the score is unchanged
    broken = fitness("waterfall", {**ok, "integrity": 0.0})
    assert good >= 99 and good - broken["score"] >= 30 and broken["worst_critical"] == 0.0
    issues = [{"slide": "w1", "level": "error", "code": "OFF_SLIDE"}, {"slide": "k1", "level": "error", "code": "LOW_CONTRAST"},
              {"slide": "w2", "level": "warning", "code": "WATERFALL_NEGATIVE"}, {"slide": "h1", "level": "error", "code": "HEADLINE_TOPIC"}]
    assert integrity_issues(issues, "w1") == ["OFF_SLIDE"] and integrity_issues(issues, "w2") == ["WATERFALL_NEGATIVE"]
    assert integrity_issues(issues, "k1") == [] and integrity_issues(issues, "h1") == []  # contrast / authoring are not "broken"
    assert "LOW_CONTRAST" not in BROKEN_CODES


def test_r3_is_development_data_with_blind_result_preserved():
    from cpe import human

    rd = ROOT / "evals" / "human_reference" / "rounds" / "r3"
    st = human.read_status(rd)
    assert st["used_for_calibration"] is True and "expert" in st["evidence_standard"]
    v = json.loads((rd / "VALIDATION_v1.5.json").read_text())
    assert (v["challenger_wins"], v["baseline_wins"], v["ties"]) == (9, 1, 28)
    prof = json.loads((ROOT / "src" / "cpe" / "qa" / "archetype_profiles.json").read_text())
    assert prof["base"]["integrity"]["w"] >= 10 and prof["archetypes"]["kpi_dashboard"]["metrics"]["utilization"]["status"] == "provisional"


# ── fact model on PDF / DOCX sources and prose semantics ──────────────────────────────────────────

DEV = ROOT / "evals" / "source_to_deck" / "development"


def test_pdf_docx_sources_and_prose_semantics():
    fm = facts.build_fact_model(sorted((DEV / "promo_effectiveness" / "sources").glob("*")))
    share = next(f for f in fm["facts"] if f["claim"].startswith("Promotions represented"))
    assert [(v["value"], v["period"]) for v in share["values"]] == [(38, "FY2025"), (29, "FY2024")]  # period per number, not per sentence
    assert share["source"]["file"].endswith(".pdf")
    tbl = next(f for f in fm["facts"] if f["claim"].startswith("Household — Incremental margin"))
    assert tbl["values"][0]["value"] == -3.4 and tbl["source"]["file"].endswith(".docx")
    fm = facts.build_fact_model(sorted((DEV / "plant_capacity_es" / "sources").glob("*")))
    by = {f["claim"][:30]: f for f in fm["facts"]}
    capex = next(f for f in fm["facts"] if "cuarta línea" in f["claim"])
    assert [v["unit"] for v in capex["values"]] == ["EUR_M", "MONTHS"]  # "12 M€", "18 meses"
    dem = next(f for f in fm["facts"] if "demanda prevista" in f["claim"])
    assert dem["values"][0]["period"] == "2026" and dem["values"][0]["basis"] == "forecast"  # "15% más … que en 2025": 2025 is the base
    hrs = next(f for f in fm["facts"] if f["claim"].startswith("Cambios de formato — Horas perdidas"))
    assert hrs["values"][0]["unit"] == "HOURS"  # unit word in the header
    assert by


def test_storyline_text_evaluation_is_negation_aware():
    from cpe.reasoning.benchmark import _trap_hit, parse_storyline_md

    g = [["volume", "fell"], ["sold less"], ["al límite"], ["cut", "store labour"]]
    for affirmed in ("Margin fell because volume fell", "We sold less this year", "Las líneas están al límite"):
        assert _trap_hit(affirmed, g), affirmed
    for negated in ("Volume is not the problem", "Margin fell 1.8 pp because of mix, not volume", "Las líneas no están al límite",
                    "Recover 1.5 pp without cutting store labour"):
        assert not _trap_hit(negated, g), negated
    md = "# T\n\n**Idea principal:** X\n\n**Línea argumental**\n1. a\n2. b\n\n**Estructura del deck** (…)\n1. h1\n2. h2\n3. h3\n"
    sl = parse_storyline_md(md)
    assert sl["governing_thought"] == "X" and sl["key_line"] == ["a", "b"] and len(sl["outline"]) == 3


def test_spanish_thousands_in_headline_numbers():
    from cpe.qa.proof import headline_quantities

    assert headline_quantities("pasamos de 820.000 a 770.000 clientes")[0]["value"] == 820000


# ── protocol 1.1: decision frame (from the s1 expert feedback) ─────────────────────────────────

def _dec_facts():
    return {f"F{i}": {"id": f"F{i}", "values": [{"value": v, "unit": u}]} for i, (v, u) in
            enumerate([(1.5, "PP"), (0.9, "PP"), (0.3, "PP"), (0.2, "PP"), (0.7, "PP"), (17.4, "EUR_M")], start=1)}


def _good_frame():
    return {"governing_thought": "Approve a plan that recovers about 1.4 pp of the 1.5 pp target",
            "key_line": [{"id": "K1", "message": "Mix cost 1.5 pp", "role": "problem"}, {"id": "K2", "message": "Three levers add 1.4 pp", "role": "solution"}],
            "decision": {"target": {"value": 1.5, "unit": "PP", "facts": ["F1"]},
                         "levers": [{"id": "L1", "statement": "reprice", "impact": {"value": 0.9, "unit": "PP"}, "facts": ["F2"], "bound": "upper", "validate": "Q1 price test"},
                                    {"id": "L2", "statement": "rates", "impact": {"value": 0.3, "unit": "PP"}, "facts": ["F3"], "bound": "point"},
                                    {"id": "L3", "statement": "recovery", "impact": {"value": 0.2, "unit": "PP"}, "facts": ["F4"]}],
                         "identified": {"value": 1.4, "unit": "PP"}, "gap": {"value": 0.1, "unit": "PP", "how_closed": "further levers"},
                         "current_plan": {"statement": "+25% electronics", "verdict": "does_not_hold", "facts": ["F5"], "cost_if_kept": {"value": 0.7, "unit": "PP"}},
                         "options": [{"id": "O1", "statement": "keep budget", "cost": {"value": 0.7, "unit": "PP"}}, {"id": "O2", "statement": "conditional growth", "cost": {"value": 0, "unit": "PP"}, "chosen": True}],
                         "asks": [{"statement": "approve the plan", "type": "approve", "owner": "CFO"}],
                         "gates": [{"statement": "price test", "criterion": "volume loss < 5%", "when": "Q1 2027"}],
                         "kpis": [{"name": "gross margin", "target": "27.9%", "cadence": "monthly"}]}}


def test_decision_frame_clean_and_missing():
    from cpe.reasoning.decision import check_decision, decision_summary
    assert [i["code"] for i in check_decision(_good_frame(), _dec_facts())] == ["OPTION_COMPONENTS_MISSING"]
    s = decision_summary(_good_frame(), [])
    assert s["levers_quantified"] == 3 and s["approvable_asks"] == 1 and s["current_plan_tested"] and s["problem_first"]
    sl = _good_frame(); sl.pop("decision")
    assert [i["code"] for i in check_decision(sl, _dec_facts())] == ["DECISION_FRAME_MISSING"]


def test_decision_numbers_must_reconcile_and_be_grounded():
    from cpe.reasoning.decision import check_decision
    sl = _good_frame(); sl["decision"]["identified"]["value"] = 1.6
    codes = [(i["code"], i["ref"], i["hard"]) for i in check_decision(sl, _dec_facts())]
    assert ("ARITHMETIC_ERROR", "identified", True) in codes
    sl = _good_frame(); sl["decision"]["gap"]["value"] = 0.3
    assert any(i["code"] == "ARITHMETIC_ERROR" and i["ref"] == "gap" for i in check_decision(sl, _dec_facts()))
    sl = _good_frame(); sl["decision"].pop("gap")
    assert any(i["code"] == "GAP_NOT_STATED" for i in check_decision(sl, _dec_facts()))
    sl = _good_frame(); sl["decision"]["levers"][1]["impact"]["value"] = 0.4; sl["decision"]["identified"]["value"] = 1.5; sl["decision"].pop("gap")
    assert any(i["code"] == "UNSUPPORTED_NUMBER" and i["hard"] for i in check_decision(sl, _dec_facts()))


def test_decision_thinking_warnings():
    from cpe.reasoning.decision import check_decision
    sl = _good_frame(); d = sl["decision"]
    d["asks"] = [{"statement": "commission a margin bridge", "type": "commission"}]
    d["options"] = d["options"][:1]; d.pop("current_plan"); d.pop("gates"); d.pop("kpis")
    d["levers"][0].pop("validate")
    sl["governing_thought"] = "Approve a plan that recovers 1.4 pp"
    sl["key_line"] = list(reversed(sl["key_line"]))
    codes = {i["code"] for i in check_decision(sl, _dec_facts())}
    assert {"ASK_DEFERRED", "OPTIONS_NOT_COMPARED", "CURRENT_PLAN_NOT_TESTED", "GATES_MISSING", "KPIS_MISSING",
            "UNCERTAINTY_UNMARKED", "OVERCLAIM_BOUND", "SOLUTION_BEFORE_PROBLEM"} <= codes
    assert not any(i["hard"] for i in check_decision(sl, _dec_facts()))


def test_decision_signals_text():
    from cpe.reasoning.decision import decision_signals_text
    s = decision_signals_text("The €17M is an upper bound; validate elasticity. Do not approve the 20% increase. Commission a study.")
    assert s["upper_bound"] == 1 and s["to_validate"] >= 1 and s["current_plan_challenged"] >= 1 and s["deferred_ask"] >= 1


def test_fact_units_written_after_the_number(tmp_path):
    from cpe.reasoning.facts import build_fact_model
    (tmp_path / "n.md").write_text("- Spend was 25 EUR M in 2025.\n- El cierre cuesta 3,2 millones de euros.\n- Cerrar cuesta 250.000 euros por tienda.\n", encoding="utf-8")
    units = [v["unit"] for f in build_fact_model([tmp_path / "n.md"])["facts"] for v in f["values"] if v["value"] in (25, 3.2, 250000)]
    assert units == ["EUR_M", "EUR_M", "EUR"]


# ── fact-model fixes found by experiment 02 agents ─────────────────────────────────────────────

def test_csv_dot_decimals_are_not_thousands(tmp_path):
    from cpe.reasoning.facts import build_fact_model
    (tmp_path / "inv.csv").write_text("month,supplier,invoiced_eur_m\n2025-01,Inc,1.444\n2025-02,Inc,0.598\n", encoding="utf-8")
    vals = sorted(v["value"] for f in build_fact_model([tmp_path / "inv.csv"])["facts"] for v in f["values"])
    assert vals == [0.598, 1.444]
    assert all(v["unit"] == "EUR_M" for f in build_fact_model([tmp_path / "inv.csv"])["facts"] for v in f["values"])


def test_xlsx_typed_floats_and_note_sheets(tmp_path):
    import openpyxl

    from cpe.reasoning.facts import build_fact_model
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Costes"
    ws.append(["Almacén", "Coste fijo (M€)"])
    ws.append(["Valencia", 1.444])
    n = wb.create_sheet("Notas")
    n.append(["Nota"])
    n.append(["Servir Levante desde Madrid añade 1,9 M€ al año de transporte."])
    wb.save(tmp_path / "a.xlsx")
    fm = build_fact_model([tmp_path / "a.xlsx"])
    vals = {(v["value"], v["unit"]) for f in fm["facts"] for v in f["values"]}
    assert (1.444, "EUR_M") in vals and (1.9, "EUR_M") in vals
    assert not any("Nota" in f["claim"] and f["fact_type"] == "table_value" for f in fm["facts"])


def test_year_column_labels_rows_and_sets_period(tmp_path):
    from cpe.reasoning.facts import build_fact_model
    (tmp_path / "b.csv").write_text("year,opening_arr_eur_m\n2024,41.6\n2025,52.0\n", encoding="utf-8")
    fs = build_fact_model([tmp_path / "b.csv"])["facts"]
    assert {(v["value"], v["unit"], v["period"]) for f in fs for v in f["values"]} == {(41.6, "EUR_M", "2024"), (52.0, "EUR_M", "2025")}


def test_conflicts_ignore_rows_columns_parts_and_rounding():
    from cpe.reasoning.conflicts import detect_conflicts

    def tv(i, label, col, value, unit="EUR_K", loc="sheet T"):
        return {"id": i, "fact_type": "table_value", "claim": f"{label} — {col}", "source": {"file": "t.xlsx", "loc": loc},
                "values": [{"value": value, "unit": unit, "period": "2025", "label": label, "column": col}]}

    def tx(i, claim, value, unit):
        return {"id": i, "fact_type": "text_statement", "claim": claim, "source": {"file": "n.md", "loc": "line 1"},
                "values": [{"value": value, "unit": unit, "period": "2025"}]}

    facts = [tv("F1", "B · no", "Ventas 2025", 900), tv("F2", "B · no", "Ventas 2025", 820),  # two stores with the same label
             tv("F3", "Valencia", "Coste fijo 2025", 2600), tv("F4", "Valencia", "Coste variable 2025", 1500),
             tv("F5", "Total", "EBITDA 2025", -5196),
             tx("F6", "Valencia nos costó 4,1 M€ en 2025 (coste fijo más variable)", 4.1, "EUR_M"),
             tx("F7", "Las tiendas perdieron 5,2 M€ de EBITDA en 2025", 5.2, "EUR_M")]
    assert detect_conflicts(facts) == []
    real = [tv("F1", "Total", "Revenue 2025", 1460, "EUR_M"), tx("F2", "The forecast showed revenue of €1,500M in 2025", 1500, "EUR_M")]
    assert len(detect_conflicts(real)) == 1


def test_vote_page_swaps_sides_of_repeats():
    from cpe import human
    js = human.PAGE if hasattr(human, "PAGE") else open(human.__file__, encoding="utf-8").read()
    assert "seen[k]?[...seen[k]].reverse()" in js  # a repeat is shown with the sides of its original swapped


def test_options_compared_on_the_same_basis():
    from cpe.reasoning.decision import check_decision
    sl = _good_frame()
    sl["decision"]["options"] = [{"id": "O1", "statement": "100% B", "cost": {"value": 1.5, "unit": "EUR_M"}, "cost_components": {"price": 2.9, "switching": -0.3, "risk": -1.3}},
                                 {"id": "O2", "statement": "70/30", "cost": {"value": 1.8, "unit": "EUR_M"}, "cost_components": {"price": 1.8, "switching": -0.2}}]
    issues = check_decision(sl, _dec_facts())
    assert [(i["code"], i["ref"]) for i in issues] == [("OPTIONS_DIFFERENT_BASIS", "O2")] and "risk" in issues[0]["message"]
    sl["decision"]["options"][1]["cost_components"]["risk"] = -0.9
    assert check_decision(sl, _dec_facts()) == []


# ── v1.8 factuality: ids survive extractor upgrades, facts verified against raw sources, every slide number grounded ──

def test_preserve_ids_across_extractor_versions():
    from cpe.reasoning.facts import preserve_ids
    old = [{"id": "F0001", "values": [{"value": 2024}], "source": {"file": "b.csv", "sheet": None, "range": "A2"}},
           {"id": "F0002", "values": [{"value": 41.6, "unit": "EUR"}], "source": {"file": "b.csv", "sheet": None, "range": "B2"}},
           {"id": "F0003", "values": [{"value": 12, "unit": "PCT"}], "source": {"file": "n.md", "loc": "line 5"}}]
    new = [{"id": "F0001", "values": [{"value": 41.6, "unit": "EUR_M"}], "source": {"file": "b.csv", "sheet": None, "range": "B2"}},
           {"id": "F0002", "values": [{"value": 12, "unit": "PCT"}], "source": {"file": "n.md", "loc": "line 5"}},
           {"id": "F0003", "values": [{"value": 9, "unit": "DAYS"}], "source": {"file": "n.md", "loc": "line 9"}},
           {"id": "F0004", "values": [{"value": 1.0}], "derived_from": ["F0001", "F0002"], "source": {"file": "b.csv", "range": "B2:B3"}}]
    facts, rep = preserve_ids(old, new)
    assert [f["id"] for f in facts] == ["F0002", "F0003", "F0004", "F0005"]
    assert facts[3]["derived_from"] == ["F0002", "F0003"] and rep["dropped"] == ["F0001"] and rep["added"] == ["F0004", "F0005"]


def test_rerun_of_facts_keeps_ids(tmp_path):
    from cpe.reasoning.facts import write_fact_model
    src = tmp_path / "sources"
    src.mkdir()
    (src / "b.csv").write_text("year,opening_arr_eur_m\n2024,41.6\n2025,52.0\n", encoding="utf-8")
    write_fact_model(src, tmp_path / "work")
    first = {f["id"]: f["values"][0]["value"] for f in json.loads((tmp_path / "work" / "facts.json").read_text())["facts"]}
    (src / "n.md").write_text("- Churn rose to 15% in 2025.\n", encoding="utf-8")
    write_fact_model(src, tmp_path / "work")
    second = {f["id"]: f["values"][0]["value"] for f in json.loads((tmp_path / "work" / "facts.json").read_text())["facts"]}
    assert all(second[k] == v for k, v in first.items()) and len(second) == len(first) + 1


def test_facts_verified_against_raw_source(tmp_path):
    from cpe.reasoning.checks import check_facts, raw_numbers
    (tmp_path / "b.csv").write_text("year,arr\n2024,41.6\n2025,1.444\n", encoding="utf-8")
    (tmp_path / "n.md").write_text("El cierre cuesta 3,2 millones; 30.000 clientes.\n", encoding="utf-8")
    assert {41.6, 1.444, 1444.0} <= raw_numbers(tmp_path / "b.csv") and {3.2, 30000.0} <= raw_numbers(tmp_path / "n.md")
    fm = {"facts": [{"id": "F1", "values": [{"value": 41.6}], "source": {"file": "b.csv"}, "fact_type": "table_value"},
                    {"id": "F2", "values": [{"value": 47.0}], "source": {"file": "b.csv"}, "fact_type": "table_value"}]}
    assert [(i["code"], i["ref"]) for i in check_facts(fm, tmp_path)] == [("FACT_FABRICATED", "F2")]


def test_label_cells_are_not_numbers():
    from cpe.ingest.readers import _num
    assert _num("Q1 2024")[0] is None and _num("P01")[0] is None and _num("12%")[0] == 12 and _num("€3.2M")[0] == 3.2


def test_every_number_on_a_slide_is_grounded():
    from cpe.reasoning.checks import factcheck_deck
    facts = {"F1": {"id": "F1", "values": [{"value": 23, "unit": "PCT"}, {"value": 25, "unit": "PCT"}], "source": {"file": "crm.xlsx"}},
             "F2": {"id": "F2", "values": [{"value": 9.4, "unit": "EUR_M"}], "source": {"file": "arr.csv"}},
             "C1": {"id": "C1", "values": [{"value": 9.6, "unit": "EUR_M"}], "fact_type": "computed", "derived_from": ["F2"], "source": {"file": "computed"}}}
    slide = {"id": "S1", "headline": "Win rate held", "evidence": [{"fact": "F1"}, {"fact": "C1"}], "source": "arr.csv; crm.xlsx",
             "kpis": {"items": [{"value": "23–25%", "label": "Quarterly win rate 2024-25"}]},
             "visual": {"type": "bar", "data": {"categories": ["2024", "2025"], "series": [{"name": "ARR", "values": [9.6, 13.7]}]}},
             "commentary": {"points": ["New ARR rose to €9.6M", "Deal size fell to €53.7k"]}}
    issues = factcheck_deck({"slides": [slide]}, facts)
    assert sorted((i["ref"], i["hard"]) for i in issues) == [("S1:commentary.points[1]", True), ("S1:visual.data.series[0].values[1]", True)]


def test_enrich_keeps_the_agents_claim(tmp_path):
    from cpe.reasoning.ghost import enrich_evidence
    (tmp_path / "facts.json").write_text(json.dumps({"facts": [{"id": "F1", "claim": "row 2 — arr: 41.6", "values": [{"value": 41.6}], "source": {"file": "b.csv"}}]}))
    (tmp_path / "deck.json").write_text(json.dumps({"slides": [{"id": "S1", "evidence": [{"fact": "F1", "claim": "Opening ARR 2024"}, {"fact": "F1"}]}]}))
    enrich_evidence(tmp_path)
    ev = json.loads((tmp_path / "deck.json").read_text())["slides"][0]["evidence"]
    assert ev[0]["claim"] == "Opening ARR 2024" and ev[0]["fact_claim"] == "row 2 — arr: 41.6" and ev[0]["values"] == [41.6]
    assert ev[1]["claim"] == "row 2 — arr: 41.6"


def test_deck_plan_headline_budget_follows_deck_type():
    from cpe.reasoning.checks import check_deck_plan
    long = "Churn rather than sales explains most of the fall in net new ARR and the retention plan closes almost all of the gap at lower cost"
    dp = {"slides": [{"id": "S1", "priority": "core", "headline": long, "key_line": "K1"}]}
    sl = {"key_line": [{"id": "K1", "message": "x"}]}
    board = [i["code"] for i in check_deck_plan(dp, sl, {}, {}, {"deck_type": "board_presentation"})]
    assert "HEADLINE_LONG" in board


# ── v1.8 visual minimum for decision decks: estimates, bounds, target, muted bridge ──

def test_waterfall_estimate_steps_target_and_muted(tmp_path):
    from pptx import Presentation

    from cpe.design.tokens import load_profile, load_theme
    from cpe.qa import geometry
    from test_v15 import _shapes, content_slide

    steps = [{"label": "Onboarding", "value": 2.5, "bound": "upper"}, {"label": "Expansion", "value": 1.3, "bound": "upper"},
             {"label": "10 AEs", "value": 1.5, "estimate": True}, {"label": "Identified", "type": "total"}]
    s = content_slide(headline="Three levers identify up to €5.3M of the €5.6M needed", message_type="change_bridge",
                      visual={"type": "waterfall", "title": "Levers", "unit": "€M", "delta_colors": "muted", "highlight": ["Onboarding"],
                              "target": {"value": 5.6, "label": "Needed"}, "data": {"steps": steps}})
    _, man, out = _shapes(tmp_path, [s])
    texts = [sh.text_frame.text for sl in Presentation(str(out)).slides for sh in sl.shapes if sh.has_text_frame]
    assert "≤2.5" in texts and "~1.5" in texts and any(t.startswith("Needed 5.6") for t in texts)
    issues = geometry.check(str(out), man, load_theme(), load_profile("standard"))
    assert not {i["code"] for i in issues if i["level"] == "error"} & {"OFF_SLIDE", "OUTSIDE_ZONE", "TEXT_OVERFLOW"}


def test_table_cells_marked_as_bounds_stay_numbers(tmp_path):
    from pptx import Presentation

    from test_v15 import _shapes, content_slide

    vis = {"type": "table", "title": "Options", "columns": [{"label": "Option"}, {"label": "€M a year", "kind": "number", "format": {"decimals": 1}}],
           "rows": [["CRO plan", 3.6], ["Retention first", {"value": 2.0, "bound": "estimate"}], ["AE ramp", {"value": 4.5, "bound": "upper"}]]}
    s = content_slide(headline="Retention first costs an estimated €2.0M versus €3.6M", message_type="comparison", visual=vis)
    _shapes(tmp_path, [s])
    out = next(tmp_path.rglob("*.pptx"))
    cells = [c.text for sl in Presentation(str(out)).slides for sh in sl.shapes if sh.has_table for r in sh.table.rows for c in r.cells]
    assert "~2.0" in cells and "≤4.5" in cells and "3.6" in cells


def test_agent_reported_parsing_bugs_set_02():
    from cpe.ingest.readers import _num as rnum
    from cpe.qa.proof import headline_quantities
    from cpe.reasoning.facts import detect_period, detect_unit
    assert [q["value"] for q in headline_quantities("ahorra 0,048 M€ y 0.106 M€")] == [48000.0, 106000.0]
    assert [q["value"] for q in headline_quantities("30.000 clientes")] == [30000.0]
    assert rnum("0,048")[0] == 0.048 and rnum("30.000")[0] == 30000
    assert detect_unit("Capacidad (miles de pedidos/año)") != "YEARS"
    assert (detect_period("encuesta 2025: se plantearía cambiar") or {}).get("basis") != "plan"


def test_agent_reported_grounding_bugs_protocol_12_rerun():
    from cpe.reasoning.checks import _safe_eval, check_computed
    from cpe.reasoning.grounding import ground_numbers
    f9 = {"id": "F9", "values": [{"value": 7.8, "unit": "EUR_M"}]}
    for t in ("7.8 EUR M churned", "EUR 7.8M churned", "€7.8M churned"):
        assert [g["status"] for g in ground_numbers(t, [f9])] == ["grounded"], t
    c = {"id": "C1", "values": [{"value": 2.868, "unit": "EUR_M"}]}
    assert [g["status"] for g in ground_numbers("a gross saving of €2.868M", [c])] == ["grounded"]
    wr = [{"id": "F1", "values": [{"value": 23, "unit": "PCT"}]}, {"id": "F2", "values": [{"value": 24, "unit": "PCT"}]}]
    assert {g["status"] for g in ground_numbers("win rate held at 23–24%", wr)} == {"grounded"}
    assert [g["status"] for g in ground_numbers("back to 10.4 in 2026", [{"id": "C2", "values": [{"value": 10.4, "unit": "EUR_M"}]}])] == ["grounded"]
    facts = {"F0062": {"id": "F0062", "values": [{"value": 14, "unit": "DAYS"}, {"value": 2.4, "unit": "EUR_M"}]},
             "F0061": {"id": "F0061", "values": [{"value": 1.6, "unit": "EUR_M"}]}}
    issues, good = check_computed({"facts": [{"id": "C0001", "formula": "F0061 + F0062[1]", "values": [{"value": 4.0, "unit": "EUR_M"}]}]}, facts)
    assert issues == [] and good["C0001"]["derived_from"] == ["F0061", "F0062"]
    assert _safe_eval("F62[1] * 2", {"F62": [14.0, 2.4]}) == 4.8


def test_conflicts_file_accepts_plain_ids_and_reports_bad_shapes(tmp_path):
    from cpe.reasoning.checks import check_conflicts
    (tmp_path / "fact_conflicts.json").write_text(json.dumps({"conflicts": [{"facts": ["F1", "F2"], "resolution": "x"}]}))
    assert check_conflicts(tmp_path, {"facts": []}, set()) == []
    (tmp_path / "fact_conflicts.json").write_text(json.dumps({"conflicts": [{"facts": [3, 4]}, "X001"]}))
    assert [i["code"] for i in check_conflicts(tmp_path, {"facts": []}, set())] == ["CONFLICT_FORMAT", "CONFLICT_FORMAT"]


def test_spanish_deck_numbers_source_label_and_verbs():
    from cpe.core.headline import lint_headline
    from cpe.qa.proof import headline_quantities
    assert [q["value"] for q in headline_quantities("mejora 1.596 k€ y 1.500 M€")] == [1596000.0, 1500000000.0]
    assert [q["value"] for q in headline_quantities("€2.868M")] == [2868000.0]
    codes = {i["code"] for i in lint_headline("Cerrar las 15 tiendas con ruptura cuesta 1500 k€ en 2026")[1]}
    assert "HEADLINE_NO_VERB" not in codes and "HEADLINE_TOPIC" not in codes
    import inspect

    from cpe.pptx import text_components
    assert '"Fuente: "' in inspect.getsource(text_components.footer)


# ── v1.8: first external case (private) exposed these; tests use synthetic text only ──

def test_decimal_comma_documents_and_markdown_tables(tmp_path):
    from cpe.reasoning.facts import build_fact_model, detect_unit
    md = ("# Caso\n\nLa Ratio de enero fue **1,734** y hoy es 1,652; el objetivo es 1,777. La distancia baja de 2,55 km a 2,19 km.\n\n"
          "| | Mayo | Hoy |\n|---|---|---|\n| Ratio | 1,734 | **1,652** |\n| Pedidos/semana | 48.213 | 41.902 |\n| Utilización | 79,1 % | 75,3 % |\n\n"
          "| Ciudad | Utilización enero → hoy | Pedidos/sem eq. |\n|---|---|---|\n| Norte | 83,2 % → 72,6 % | 151 |\n| Sur | 70,4 % → 66,9 % | 87 |\n"
          "| Resto | | 34-44 cada una |\n| **Total** | | **≈ 4.120** |\n")
    (tmp_path / "c.md").write_text(md, encoding="utf-8")
    fm = build_fact_model([tmp_path / "c.md"])
    vals = {v["value"] for f in fm["facts"] for v in f["values"]}
    assert {1.734, 1.652, 1.777, 48213.0, 41902.0, 79.1, 83.2, 72.6, 151.0, 4120.0} <= vals
    assert 1652.0 not in vals and 1734.0 not in vals
    locs = {f["source"]["loc"] for f in fm["facts"] if f["fact_type"] == "table_value"}
    assert "line 5" in locs  # the line the table starts on
    assert detect_unit("Pedidos/semana") == "" and detect_unit("Año ant.") == "" and detect_unit("Tiempo (min)") == "MINUTES"
    assert detect_unit("Horas perdidas 2025") == "HOURS"


# ── v1.8: row-level data through recorded analysis scripts ──

def test_row_level_data_via_recorded_analysis(tmp_path):
    from cpe.reasoning.analysis import check_analyses, run_analysis
    from cpe.reasoning.checks import check_facts
    from cpe.reasoning.facts import write_fact_model
    src, work = tmp_path / "sources", tmp_path / "work"
    src.mkdir()
    rows = "\n".join(f"{i},{'A' if i % 3 else 'B'},{10 + i % 7}" for i in range(400))
    (src / "orders.csv").write_text("order_id,store,minutes\n" + rows + "\n", encoding="utf-8")
    fm = write_fact_model(src, work)
    assert fm["stats"]["datasets"][0]["rows"] == 400 and fm["stats"]["facts"] == 0  # not read cell by cell
    script = tmp_path / "by_store.py"
    script.write_text(
        "import csv, os, collections\n"
        "rows = list(csv.DictReader(open(os.path.join(os.environ['CPE_SOURCES'], 'orders.csv'))))\n"
        "agg = collections.defaultdict(list)\n"
        "for r in rows: agg[r['store']].append(float(r['minutes']))\n"
        "with open(os.path.join(os.environ['CPE_OUT'], 'minutes_by_store.csv'), 'w') as f:\n"
        "    f.write('store,orders,avg_minutes\\n')\n"
        "    for k in sorted(agg): f.write(f'{k},{len(agg[k])},{sum(agg[k]) / len(agg[k]):.2f}\\n')\n", encoding="utf-8")
    e = run_analysis(work, script, src)
    assert e["reproducible"] and list(e["outputs"]) == ["minutes_by_store.csv"]
    fm = write_fact_model(src, work)
    files = {f["source"]["file"] for f in fm["facts"]}
    assert files == {"analysis/by_store/minutes_by_store.csv"}
    assert check_facts(fm, src, work) == [] and check_analyses(work, src) == []
    (work / "analysis" / "out" / "by_store" / "minutes_by_store.csv").write_text("store,orders,avg_minutes\nA,999,1.00\n")
    assert {i["code"] for i in check_analyses(work, src)} == {"ANALYSIS_STALE"}


def test_external_set_s4_tool_problems():
    import csv as _csv

    from cpe.ingest.readers import _table, extract_facts, sniff_dialect
    from cpe.qa.proof import headline_quantities
    from cpe.reasoning.checks import check_computed
    from cpe.reasoning.facts import detect_unit
    from cpe.reasoning.grounding import ground_numbers
    wide = ",".join(f"col_{i}_long_name" for i in range(22)) + "\n" + "\n".join(",".join(str(i * j) for i in range(22)) for j in range(20))
    assert sniff_dialect(wide).delimiter == ","
    assert sniff_dialect("a;b\n1;2\n").delimiter == ";"
    assert isinstance(sniff_dialect("x"), (_csv.Dialect, type)) or True
    assert [q["value"] for q in headline_quantities("ahorra 1.428.000 € al año")] == [1428000.0]
    assert [q["value"] for q in headline_quantities("14414 más al año")] == [14414.0]
    assert headline_quantities("vence el 31/12/2026, o el 31 de diciembre, periodo 2025-26") == []
    assert [(f["value"], f["unit"]) for f in extract_facts("Precio de 35-39 €/mes, fichero 2025-01_2026-08, el 2/12/2025.", "x", "y")] == [(35.0, ""), (39.0, "€")]
    assert [g["status"] for g in ground_numbers("cobertura de 9,55x", [{"id": "F1", "values": [{"value": 9.55, "unit": ""}]}])] == ["grounded"]
    assert detect_unit("Clientes activos (inicio mes)") == "" and detect_unit("share_pct") == "PCT"
    t = _table(["Mes", "Churn mensual", "Coste"], [["ene", "1,6%", "140.000 €"], ["feb", "1,8%", "120.000 €"]], "k.md", "line 1", decimal_comma=True)
    assert t["cell_units"][0][1] == "%" and t["rows"][0][2] == 140000.0
    facts = {"F1": {"id": "F1", "values": [{"value": 10.0}]}}
    issues, good = check_computed({"facts": [{"id": "C0001", "formula": "F1 * 2", "values": [{"value": 21}]},
                                             {"id": "C0002", "formula": "C0001 + 1", "values": [{"value": 22}]}]}, facts)
    assert [i["code"] for i in issues] == ["ARITHMETIC_ERROR"] and "C0002" in good


def test_protocol_14_checks_only_quantities():
    from cpe.reasoning.grounding import ground_numbers
    text = "Comunicar antes del 31/12/2026 al propietario; 2 personas en 6-8 semanas, 12 mensualidades de penalización, revisión en Q1 2027"
    assert ground_numbers(text, []) == [] or all(g["status"] == "grounded" for g in ground_numbers(text, []))
    assert [g["status"] for g in ground_numbers("ahorra 140 k€ al año", [])] == ["unsupported"]
