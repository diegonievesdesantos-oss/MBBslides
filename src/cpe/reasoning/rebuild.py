"""Rebuild the slides whose message no longer holds, in the deck's own style (v2.0, item 5).

Patching numbers keeps a slide's argument; when the argument itself changed, the slide has to be
rebuilt. `rewrite_spec` writes a deck.json with only those slides, taken from the old deck's native
rebuild (`old_deck_spec.json`) with every APPROVED value already in place and the proposed headline.
The reviewer (or the agent) edits that spec: the headline, the exhibit, the commentary. Then
`cpe update --rebuild` builds it with the brand learned from the old deck itself (`brand ingest`,
v1.9 D2) and proposes a `replace_slide` edit per slide: on `--apply`, the new slide takes the old
one's place in the ORIGINAL file, on the old slide's own layout.

The transplant (`transplant_slide`) copies the built slide's shapes into a new slide of the original:
- placeholders become plain shapes with their geometry and text style written out (they would
  otherwise inherit from a layout the original does not have);
- shapes the built slide got from its layout but the original layout lacks (learned repeated shapes)
  are copied too;
- charts (with their embedded workbook) and pictures are copied as new parts.
"""
from __future__ import annotations

import copy
import io
import json
import re
from pathlib import Path

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
SHAPES = {f"{P}sp", f"{P}pic", f"{P}graphicFrame", f"{P}grpSp", f"{P}cxnSp"}


# ── the spec ───────────────────────────────────────────────────────────────────

def rewrite_spec(work: str | Path, slides: list[int] | None = None) -> dict:
    """deck.json for the slides to rebuild (default: every slide whose message no longer holds)."""
    work = Path(work)
    old = json.loads((work / "old_deck_spec.json").read_text(encoding="utf-8"))
    rev = json.loads((work / "messages.json").read_text(encoding="utf-8"))
    edits = json.loads((work / "edits.json").read_text(encoding="utf-8"))
    want = set(slides or [r["slide"] for r in rev if r["verdict"] == "no longer holds"])
    approved = [e for e in edits.get("edits") or [] if e.get("approved") and e.get("op", "number") == "number"]
    heads = {e["slide"]: e["text"] for e in edits.get("edits") or [] if e.get("op") == "set_headline"}
    props = {r["slide"]: r["proposal"] for r in rev if r.get("proposal")}
    out = []
    for s in old["slides"]:
        n = s.get("_old_slide") or int(s["id"][1:])
        if n not in want:
            continue
        s = copy.deepcopy(s)
        for e in approved:
            if e["slide"] == n:
                _apply_to_spec(s, e)
        if s.get("headline") is not None:
            s["headline"] = heads.get(n) or props.get(n) or s["headline"]
        s["purpose"] = f"(rebuilt: slide {n} of the old deck, its message no longer holds; edit headline and exhibit, then `cpe update --rebuild`)"
        out.append(s)
    meta = {**old.get("meta", {}), "brand": str((work / "brand").resolve())}
    words = " ".join(str(s.get("headline") or "") for s in old["slides"]).lower().split()
    if not meta.get("language") and sum(w in ("de", "la", "el", "los", "las", "y", "en", "con", "por", "que") for w in words) > 0.08 * max(1, len(words)):
        meta["language"] = "es"  # the old deck's language: "Fuente:", decimal comma
    story = dict(old.get("storyline") or {})  # a partial deck: its storyline is the rebuilt slides' own messages
    heads = [s["headline"] for s in out if s.get("headline")]
    story["governing_thought"] = story.get("governing_thought") or (heads[0] if heads else meta.get("title", ""))
    story["key_line"] = story.get("key_line") or [{"id": f"K{i + 1}", "message": h} for i, h in enumerate(heads[:5] or [story["governing_thought"]])]
    return {"meta": meta, "storyline": story, "slides": out}


def _apply_to_spec(s: dict, e: dict) -> None:
    where, find, rep = e["where"], str(e.get("find") or ""), str(e["replace"])
    pat = re.compile(rf"(?:(?<![\w.,])[-−])?(?<![\d.,]){re.escape(find.lstrip('-−'))}(?![\d])")
    m = re.match(r"exhibit\[\d+\]\.rows\[(\d+)\]\[(\d+)\]", where)
    v = s.get("visual") or {}
    if m and v.get("type") == "table":
        r, c = int(m.group(1)), int(m.group(2))
        if r < len(v.get("rows") or []) and c < len(v["rows"][r]):
            v["rows"][r][c] = pat.sub(rep, str(v["rows"][r][c]), count=1)
        return
    m = re.match(r"exhibit\[\d+\]\.(.+)\[(.*)\]$", where)
    if m and (v.get("data") or {}).get("series"):
        from .deck_update import _parse_core

        cats = [str(c) for c in v["data"].get("categories") or []]
        for ser in v["data"]["series"]:
            if ser.get("name") == m.group(1) and m.group(2) in cats:
                x = _parse_core(rep.lstrip("-−"))
                if x is not None:
                    ser["values"][cats.index(m.group(2))] = -x if rep.startswith(("-", "−")) else x
        return
    if where == "title" and s.get("headline"):
        s["headline"] = pat.sub(rep, s["headline"], count=1)
    pts = (s.get("commentary") or {}).get("points") or []
    m = re.match(r"body\[(\d+)\]", where)
    if m and int(m.group(1)) < len(pts):
        pts[int(m.group(1))] = pat.sub(rep, pts[int(m.group(1))], count=1)


# ── the transplant ─────────────────────────────────────────────────────────────

def _inherited(ph_el, src_layout) -> tuple:
    """(xfrm, lstStyle) a placeholder takes from its layout (or the layout's master)."""
    from pptx.oxml.ns import qn

    ph = ph_el.find(f".//{P}nvPr/{P}ph")
    typ, idx = (ph.get("type") or "body"), ph.get("idx")
    match = None
    for lp in src_layout.placeholders:
        f = lp.placeholder_format
        lt = (lp._element.find(f".//{P}nvPr/{P}ph").get("type") or "body")
        if (idx is not None and str(f.idx) == idx) or (lt == typ and typ in ("title", "ctrTitle", "subTitle")):
            match = lp._element
            break
    xfrm = match.find(f"{P}spPr/{A}xfrm") if match is not None else None
    lst = match.find(f"{P}txBody/{A}lstStyle") if match is not None else None
    if (lst is None or not len(lst)) and typ in ("title", "ctrTitle"):
        ms = src_layout.slide_master._element.find(f"{P}txStyles/{P}titleStyle")
        if ms is not None:
            lst = copy.deepcopy(ms)
            lst.tag = qn("a:lstStyle")
    return xfrm, lst


def _deplaceholder(el, src_layout) -> None:
    ph = el.find(f".//{P}nvPr/{P}ph")
    if ph is None:
        return
    xfrm, lst = _inherited(el, src_layout)
    sppr = el.find(f"{P}spPr")
    if sppr is not None and sppr.find(f"{A}xfrm") is None and xfrm is not None:
        sppr.insert(0, copy.deepcopy(xfrm))
    body = el.find(f"{P}txBody")
    if body is not None and lst is not None:
        old = body.find(f"{A}lstStyle")
        if old is None or not len(old):
            if old is not None:
                body.remove(old)
            body.insert(1, copy.deepcopy(lst))
    ph.getparent().remove(ph)


def _copy_rels(el, src_part, dst_part) -> None:
    """Pictures and charts the shape refers to, copied into the destination package."""
    from pptx.opc.package import Part, XmlPart
    from pptx.opc.packuri import PackURI

    for e in el.iter():
        for at in (f"{R}embed", f"{R}id", f"{R}link"):
            rid = e.get(at)
            if not rid or rid not in src_part.rels:
                continue
            rel = src_part.rels[rid]
            if rel.is_external:
                e.set(at, dst_part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref))
                continue
            tgt = rel.target_part
            if rel.reltype.endswith("/image"):
                _, new_rid = dst_part.get_or_add_image_part(io.BytesIO(tgt.blob))
            elif rel.reltype.endswith("/chart"):
                pkg = dst_part.package
                name = pkg.next_partname("/ppt/charts/chart%d.xml")
                new = XmlPart(name, tgt.content_type, pkg, copy.deepcopy(tgt._element))
                for crid, crel in tgt.rels.items():  # its embedded workbook and style parts
                    if crel.is_external:
                        continue
                    sub = crel.target_part
                    ext = str(sub.partname).rsplit(".", 1)[-1]
                    base = str(sub.partname).rsplit("/", 1)[0]
                    sub_new = Part(PackURI(_free(pkg, f"{base}/copied%d.{ext}")), sub.content_type, pkg, sub.blob)
                    new_id = new.relate_to(sub_new, crel.reltype)
                    for x in new._element.iter():
                        for a2 in (f"{R}id", f"{R}embed"):
                            if x.get(a2) == crid:
                                x.set(a2, new_id)
                new_rid = dst_part.relate_to(new, rel.reltype)
            else:
                continue
            e.set(at, new_rid)


def _free(pkg, tmpl: str) -> str:
    names = {str(p.partname) for p in pkg.iter_parts()}
    i = 1
    while tmpl % i in names:
        i += 1
    return tmpl % i


CHROME_PH = ("sldNum", "ftr", "dt")


def transplant_slide(dst_prs, src_slide, index: int, layout, old_slide=None) -> object:
    """A copy of `src_slide` (another presentation) inserted at 0-based `index` of `dst_prs`, on `layout`.
    The old slide's own page number / footer / date placeholders are kept; the built deck's page
    number (it counts the partial deck) is dropped."""
    new = dst_prs.slides.add_slide(layout)
    tree = new.shapes._spTree
    for sh in list(new.placeholders):
        sh._element.getparent().remove(sh._element)
    if old_slide is not None:
        for sh in old_slide.placeholders:
            ph = sh._element.find(f".//{P}nvPr/{P}ph")
            if ph is not None and ph.get("type") in CHROME_PH:
                tree.append(copy.deepcopy(sh._element))
    src_layout = src_slide.slide_layout
    dst_keys = {(sh.name, sh.left, sh.top) for sh in layout.shapes}
    for sh in src_layout.shapes:  # what the built slide got from its (learned) layout but this layout lacks
        if not sh.is_placeholder and (sh.name, sh.left, sh.top) not in dst_keys:
            el = copy.deepcopy(sh._element)
            _copy_rels(el, src_layout.part, new.part)
            tree.append(el)
    for el in src_slide.shapes._spTree:
        if el.tag not in SHAPES:
            continue
        name = (el.find(f".//{P}cNvPr").get("name") if el.find(f".//{P}cNvPr") is not None else "") or ""
        text = "".join(t.text or "" for t in el.iter(f"{A}t")).strip()
        if name.startswith("cpe|chrome|") and text.isdigit():
            continue  # the partial deck's page number
        el = copy.deepcopy(el)
        _deplaceholder(el, src_layout)
        _copy_rels(el, src_slide.part, new.part)
        tree.append(el)
    if src_slide.has_notes_slide and src_slide.notes_slide.notes_text_frame.text.strip():
        new.notes_slide.notes_text_frame.text = src_slide.notes_slide.notes_text_frame.text
    lst = dst_prs.slides._sldIdLst
    el = list(lst)[-1]
    lst.remove(el)
    lst.insert(index, el)
    return new
