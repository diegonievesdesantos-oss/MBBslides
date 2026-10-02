"""v2.0: derived figures (totals, ratios, repeated figures) and the single `cpe update` command."""
import json


def _bc_deck(path):
    """A business case: a KPI on the summary repeating the table's total, a table with a total column,
    a net row and a payback row, and a sentence stating a total with its parts."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    t.text_frame.text = "Invertir 5,6 M€ en dos fases se recupera pronto"
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
    k = s.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(3), Inches(1))
    k.text_frame.text = "5,6 M€"
    lab = s.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(3), Inches(1))
    lab.text_frame.text = "Inversión total del proyecto"
    b = s.shapes.add_textbox(Inches(0.5), Inches(4), Inches(9), Inches(1))
    b.text_frame.text = "El proyecto reduce la plantilla en 39 FTE (21 en la fase 1 y 18 en la fase 2)."
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    t.text_frame.text = "Las dos fases se pagan en menos de 5 años"
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
    rows = [("Concepto", "Fase 1", "Fase 2", "Total"), ("Inversión (capex, k€)", "3.200", "2.400", "5.600"),
            ("Reducción de plantilla (FTE)", "21", "18", "39"), ("Ahorro de personal (k€)", "809", "693", "1.502"),
            ("Mantenimiento (k€)", "-150", "-110", "-260"), ("Ahorro neto anual (k€)", "659", "583", "1.242"),
            ("Payback simple (años)", "4,9", "4,1", "4,5")]
    tbl = s.shapes.add_table(len(rows), 4, Inches(0.5), Inches(1.5), Inches(9), Inches(4)).table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            tbl.cell(i, j).text = v
    prs.save(str(path))


def test_relations_found_in_the_old_deck_and_recomputed(tmp_path):
    from cpe.reasoning.deck_update import ingest_deck, update_plan
    from cpe.reasoning.derive import derive, relations

    _bc_deck(tmp_path / "old.pptx")
    plan = update_plan(ingest_deck(tmp_path / "old.pptx"), [])
    rels = relations(plan)
    how = {(r["target"].split("|")[1], r["target"].split("|")[3]): r for r in rels}
    assert how[("exhibit[0].rows[0][3]", "5.600")]["how"] == "row total"
    assert how[("exhibit[0].rows[4][1]", "659")]["how"] == "column total"  # net = personal + maintenance
    assert how[("exhibit[0].rows[5][1]", "4,9")]["op"] == "ratio"  # payback = capex / net savings
    assert how[("body[0]", "€5,6 M")]["op"] == "same"
    assert any(r["how"] == "sum stated in the text" and r["target"].split("|")[3] == "39" for r in rels)
    ids = {q["raw"] + "@" + q["where"]: q["id"] for s in plan["slides"] for q in s["numbers"]}
    known = {ids["3.200@exhibit[0].rows[0][1]"]: 4030.0, ids["2.400@exhibit[0].rows[0][2]"]: 2616.0,
             ids["809@exhibit[0].rows[2][1]"]: 590.0, ids["-150@exhibit[0].rows[3][1]"]: -210.0}
    got = derive(plan, rels, known)
    assert abs(got[ids["5.600@exhibit[0].rows[0][3]"]][0] - 6646) < 1e-6
    assert abs(got[ids["€5,6 M@body[0]"]][0] - 6.646e6) < 1  # the KPI follows the table, in its own scale
    assert abs(got[ids["659@exhibit[0].rows[4][1]"]][0] - 380) < 1e-6
    assert abs(got[ids["4,9@exhibit[0].rows[5][1]"]][0] - 4030 / 380) < 1e-6
    assert ids["39@body[2]"] not in got  # its parts have no new value: nothing is guessed


def test_cpe_update_end_to_end(tmp_path):
    from pptx import Presentation

    from cpe.cli import main

    _bc_deck(tmp_path / "old.pptx")
    (tmp_path / "src").mkdir()
    (tmp_path / "src" / "nota.md").write_text("# Nota\n\nLa plantilla de picking se ha reducido en 12 FTE desde junio.\n", encoding="utf-8")
    work, out = tmp_path / "work", tmp_path / "out" / "new.pptx"
    assert main(["update", str(tmp_path / "old.pptx"), str(tmp_path / "src"), "-o", str(work)]) == 0
    e = json.loads((work / "edits.json").read_text(encoding="utf-8"))
    plan = json.loads((work / "update_plan.json").read_text(encoding="utf-8"))
    ids = {(s["slide"], q["where"], q["raw"]): q for s in plan["slides"] for q in s["numbers"]}

    def edit(slide, where, raw, new):
        q = ids[(slide, where, raw)]
        e["edits"].append({"op": "number", "id": q["id"], "slide": slide, "where": where, "find": q["core"], "occ": q["occ"],
                           "replace": new, "old": raw, "approved": True})

    edit(2, "exhibit[0].rows[0][1]", "3.200", "4.030")  # the reviewer's corrections
    edit(2, "exhibit[0].rows[0][2]", "2.400", "2.616")
    (work / "edits.json").write_text(json.dumps(e, ensure_ascii=False), encoding="utf-8")
    assert main(["update", "--apply", str(work), "-o", str(out), "--accept-derived"]) == 0
    rep = json.loads((out.parent / "update_report.json").read_text(encoding="utf-8"))
    assert rep["rewrite_headlines"] and rep["rewrite_headlines"][0]["slide"] == 1  # "5,6 M€" in the headline changed
    new = Presentation(str(out))
    t = new.slides[1].shapes[1].table
    assert t.cell(1, 1).text == "4.030" and t.cell(1, 3).text == "6.646"  # the total follows the approved parts
    texts = [sh.text_frame.text for sh in new.slides[0].shapes if sh.has_text_frame]
    assert "6,65 M€" in texts and any(x.startswith("Invertir 6,65 M€") for x in texts)
    assert "39 FTE (21" in texts[3]  # nothing approved about head count: left as it was, listed for review
    assert any(x["number"] == "39" for x in rep["left_unchanged"])
    # re-running the preparation keeps the review
    assert main(["update", str(tmp_path / "old.pptx"), str(tmp_path / "src"), "-o", str(work)]) == 0
    e2 = json.loads((work / "edits.json").read_text(encoding="utf-8"))
    assert sum(1 for x in e2["edits"] if x.get("approved") and not x.get("derived")) == 2
