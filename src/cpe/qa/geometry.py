"""Geometric QA on the generated .pptx (independent of any renderer).

Reads every shape back with python-pptx, re-measures all text with real font
metrics and checks it against the build manifest (zones) and the design system.

Codes (severity):
  OFF_SLIDE (error)            shape extends beyond the slide
  OUTSIDE_SAFE_AREA (error)    shape outside the margins (except full-bleed)
  OUTSIDE_ZONE (error)         shape escapes the layout zone it belongs to
  TEXT_OVERFLOW (error)        text needs more height than its box
  TEXT_COLLISION (error)       the ink of two text boxes overlaps
  CONNECTOR_THROUGH_TEXT (warning) a connector crosses text
  FONT_TOO_SMALL (error/warn)  below 8 pt / below 9 pt outside the footer
  FONT_FAMILY (warning)        font other than the theme font
  COLOR_OFF_PALETTE (warning)  text colour not in the theme palette
  LOW_CONTRAST (error/warn)    contrast < 3.0 / < 4.5 against the fill behind
  PLACEHOLDER_TEXT (error)     TODO / lorem / xx / [..] left in the deck
  HEADLINE_LINES (error)       headline wraps to more than 2 lines
  HEADLINE_WIDOW (warning)     last headline line is a single short word
  MISALIGNED (warning)         left edges that almost (but not exactly) align
  SHAPE_COUNT (warning)        more shapes than the density profile allows
  TINY_ELEMENT (warning)       non-line shape smaller than 0.03 in
  EMPTY_ZONE (warning)         a layout zone received no content
  FIT_SHRUNK (info)            text was shrunk towards the floor to fit
"""
from __future__ import annotations

import re
from dataclasses import dataclass

from pptx import Presentation
from pptx.util import Emu

from ..design import text_metrics as tm
from ..design.tokens import GRID, SLIDE_H, SLIDE_W, Theme, contrast_ratio
from ..layout.engine import Box
from ..spec import issue

EMU = 914400
PLACEHOLDER_RE = re.compile(r"\bTODO\b|\bTBD\b|lorem ipsum|\bxx+\b|\[[^\]]{1,40}\]|\?\?\?", re.IGNORECASE)
TOL = 0.03


@dataclass
class ShapeInfo:
    name: str
    zone: str
    kind: str
    box: Box
    shape: object
    order: int
    text: str = ""
    ink: Box | None = None
    need_h: float = 0.0
    lines: int = 0
    fill: str | None = None
    vertical: bool = False


def _parse_name(name: str) -> tuple[str, str]:
    parts = (name or "").split("|")
    if len(parts) >= 4 and parts[0] == "cpe":
        return parts[1], parts[2]
    return "?", "?"


def _box(sh) -> Box:
    return Box(Emu(sh.left).inches, Emu(sh.top).inches, Emu(sh.width).inches, Emu(sh.height).inches)


def _fill_hex(sh) -> str | None:
    try:
        if sh.fill.type == 1:  # solid
            return str(sh.fill.fore_color.rgb)
    except Exception:
        return None
    return None


def measure_text_frame(sh) -> tuple[float, int, float, list[str]]:
    """(needed height, line count, widest line width, lines) of a shape's text frame."""
    tf = sh.text_frame
    vertical = tf._txBody.bodyPr.get("vert") in ("vert270", "vert", "eaVert")
    ml = Emu(tf.margin_left or 0).inches
    mr = Emu(tf.margin_right or 0).inches
    box_w = Emu(sh.width).inches if not vertical else Emu(sh.height).inches
    width = max(0.05, box_w - ml - mr)
    total_h = 0.0
    nlines = 0
    widest = 0.0
    all_lines: list[str] = []
    paras = tf.paragraphs
    for i, p in enumerate(paras):
        runs = p.runs
        text = "".join(r.text for r in runs)
        sizes = [r.font.size.pt for r in runs if r.font.size is not None] or [18.0]
        size = max(sizes)
        bold_chars = sum(len(r.text) for r in runs if r.font.bold)
        bold = bold_chars > len(text) / 2 if text else False
        pPr = p._p.pPr
        indent = 0.0
        if pPr is not None and pPr.get("marL"):
            indent = int(pPr.get("marL")) / EMU
        spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        lines = tm.wrap_lines(text, max(0.05, width - indent), size, bold) if text else [""]
        all_lines += lines
        nlines += len(lines)
        total_h += len(lines) * tm.line_height_in(size, spacing)
        widest = max([widest] + [tm.text_width_in(l, size, bold) + indent for l in lines])
        if i < len(paras) - 1 and p.space_after is not None:
            total_h += p.space_after.pt / 72
    return total_h, nlines, widest, all_lines


def ink_box(sh, box: Box, need_h: float, widest: float) -> Box:
    tf = sh.text_frame
    anchor = tf._txBody.bodyPr.get("anchor", "t")
    ml = Emu(tf.margin_left or 0).inches
    mt = Emu(tf.margin_top or 0).inches
    inner_h = box.h - 2 * mt
    h = min(need_h, max(inner_h, need_h))
    if anchor == "ctr":
        y = box.y + mt + (inner_h - h) / 2
    elif anchor == "b":
        y = box.b - mt - h
    else:
        y = box.y + mt
    align = None
    if tf.paragraphs:
        a = tf.paragraphs[0].alignment
        align = str(a) if a is not None else None
    w = min(widest, box.w - 2 * ml) if widest else 0
    if align and "CENTER" in align:
        x = box.x + (box.w - w) / 2
    elif align and "RIGHT" in align:
        x = box.r - ml - w
    else:
        x = box.x + ml
    return Box(x, y, w, h)


def collect(slide) -> list[ShapeInfo]:
    out = []
    for i, sh in enumerate(slide.shapes):
        zone, kind = _parse_name(sh.name)
        b = _box(sh)
        info = ShapeInfo(sh.name, zone, kind, b, sh, i, fill=_fill_hex(sh))
        if sh.has_text_frame and sh.text_frame.text.strip():
            info.text = sh.text_frame.text
            need, lines, widest, _ = measure_text_frame(sh)
            info.vertical = sh.text_frame._txBody.bodyPr.get("vert") in ("vert270", "vert", "eaVert")
            info.need_h, info.lines = need, lines
            info.ink = b if info.vertical else ink_box(sh, b, need, widest)
        out.append(info)
    return out


def _all_run_props(sh):
    if not sh.has_text_frame:
        return
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            yield r


def check_slide(slide, idx: int, manifest: dict | None, theme: Theme, profile: dict) -> list[dict]:
    sid = (manifest or {}).get("slide_id") or f"#{idx}"
    out: list[dict] = []
    shapes = collect(slide)
    zones = {n: Box(z["x"], z["y"], z["w"], z["h"]) for n, z in ((manifest or {}).get("zones") or {}).items()}
    safe_l, safe_r = GRID.margin_l - TOL, SLIDE_W - GRID.margin_r + TOL
    safe_t, safe_b = GRID.tracker_y - TOL, GRID.footer_y + GRID.footer_h + TOL
    palette = theme.palette() | {"FFFFFF", "000000"}
    used_zones = set()
    texts = [s for s in shapes if s.ink is not None]
    for s in shapes:
        b = s.box
        used_zones.add(s.zone)
        if b.x < -TOL or b.y < -TOL or b.r > SLIDE_W + TOL or b.b > SLIDE_H + TOL:
            out.append(issue("error", "OFF_SLIDE", f"{s.name} extends beyond the slide ({b.x:.2f},{b.y:.2f},{b.r:.2f},{b.b:.2f})", sid, shape=s.name))
            continue
        if s.kind != "bleed" and (b.x < safe_l or b.r > safe_r or b.y < safe_t or b.b > safe_b):
            out.append(issue("error", "OUTSIDE_SAFE_AREA", f"{s.name} crosses the margins", sid, shape=s.name))
        if s.zone in zones and s.kind not in ("bleed",):
            zb = zones[s.zone]
            # labels may sit a little outside plot areas but never outside the zone
            test = s.ink if (s.ink is not None and s.kind in ("text", "label")) else b
            if s.kind == "connector" or s.kind == "line":
                test = b
            if not zb.contains(test, 0.06):
                out.append(issue("error", "OUTSIDE_ZONE", f"{s.name} escapes zone '{s.zone}'", sid, shape=s.name))
        if s.kind not in ("line", "connector") and (b.w < 0.03 and b.h < 0.03):
            out.append(issue("warning", "TINY_ELEMENT", f"{s.name} is {b.w:.3f}×{b.h:.3f} in", sid, shape=s.name))
        # text checks
        if s.ink is not None:
            inner_h = (b.w if s.vertical else b.h) - 2 * Emu(s.shape.text_frame.margin_top or 0).inches
            if s.need_h > inner_h + max(0.03, 0.03 * inner_h):
                out.append(issue("error", "TEXT_OVERFLOW", f"Text needs {s.need_h:.2f} in but box is {inner_h:.2f} in: '{s.text[:50]}'", sid, shape=s.name, zone=s.zone))
            if PLACEHOLDER_RE.search(s.text):
                out.append(issue("error", "PLACEHOLDER_TEXT", f"Placeholder text left: '{s.text[:60]}'", sid, shape=s.name))
            in_footer = b.y >= GRID.footer_y - 0.05
            for r in _all_run_props(s.shape):
                if r.font.size is not None:
                    pt = r.font.size.pt
                    if pt < 8 - 1e-6:
                        out.append(issue("error", "FONT_TOO_SMALL", f"{pt:g} pt text: '{r.text[:30]}'", sid, shape=s.name))
                        break
                    if pt < 9 - 1e-6 and not in_footer:
                        out.append(issue("warning", "FONT_TOO_SMALL", f"{pt:g} pt text outside the footer: '{r.text[:30]}'", sid, shape=s.name))
                        break
                if r.font.name and r.font.name != theme.font_latin:
                    out.append(issue("warning", "FONT_FAMILY", f"Font '{r.font.name}' is not the theme font", sid, shape=s.name))
                    break
                try:
                    col = str(r.font.color.rgb) if r.font.color and r.font.color.type is not None else None
                except AttributeError:
                    col = None
                if col and col.upper() not in palette:
                    out.append(issue("warning", "COLOR_OFF_PALETTE", f"Text colour #{col} not in theme palette", sid, shape=s.name))
                if col:
                    bg = _background_at(shapes, s, theme)
                    cr = contrast_ratio(col, bg)
                    if cr < 3.0:
                        out.append(issue("error", "LOW_CONTRAST", f"Contrast {cr:.1f}:1 for '{r.text[:25]}' on #{bg}", sid, shape=s.name))
                        break
                    if cr < 4.5 and (r.font.size is None or r.font.size.pt < 14) and not r.font.bold:
                        out.append(issue("warning", "LOW_CONTRAST", f"Contrast {cr:.1f}:1 for small text '{r.text[:25]}'", sid, shape=s.name))
                        break
    # text collisions (ink vs ink)
    for i in range(len(texts)):
        for j in range(i + 1, len(texts)):
            a, c = texts[i], texts[j]
            ia, ic = a.ink, c.ink
            inter = ia.intersection(ic)
            if inter <= 0:
                continue
            small = min(ia.w * ia.h, ic.w * ic.h) or 1e-6
            if inter / small > 0.08 and inter > 0.004:
                out.append(issue("error", "TEXT_COLLISION", f"'{a.text[:30]}' overlaps '{c.text[:30]}'", sid, shape=a.name, other=c.name))
    # connectors crossing text
    for s in shapes:
        if s.kind != "connector":
            continue
        lb = s.box
        seg = Box(lb.x - 0.01, lb.y - 0.01, lb.w + 0.02, lb.h + 0.02)
        for t in texts:
            if t.kind == "label" and seg.intersection(t.ink) > 0 and t.zone == s.zone:
                continue  # annotation labels sit on their own leader lines by design
            cx, cy = t.ink.x + t.ink.w / 2, t.ink.y + t.ink.h / 2
            if any(o.fill and s.order < o.order < t.order and o.box.x <= cx <= o.box.r and o.box.y <= cy <= o.box.b for o in shapes):
                continue  # the connector runs behind a filled shape that carries the text
            if seg.intersection(t.ink.inset(0.02, 0.02, 0.02, 0.02)) > 0.0005:
                out.append(issue("warning", "CONNECTOR_THROUGH_TEXT", f"{s.name} crosses '{t.text[:30]}'", sid, shape=s.name))
                break
    # headline
    head = next((s for s in texts if s.zone == "chrome" and abs(s.box.y - GRID.headline_y) < 0.01), None)
    if head is not None:
        _, n, _, lines = measure_text_frame(head.shape)
        if n > profile.get("headline_max_lines", 2):
            out.append(issue("error", "HEADLINE_LINES", f"Headline wraps to {n} lines", sid))
        if n >= 2 and len(lines[-1].split()) == 1 and len(lines[-1]) < 12:
            out.append(issue("warning", "HEADLINE_WIDOW", f"Headline's last line is the single word '{lines[-1]}'", sid))
    # alignment near-misses among body text left edges
    roles = {n: z.get("role") for n, z in ((manifest or {}).get("zones") or {}).items()}
    lefts = sorted({round(s.box.x, 3) for s in texts if roles.get(s.zone) in ("commentary", "column", "statements", "kpis", "takeaway") and s.kind == "text"})
    for a, c in zip(lefts, lefts[1:]):
        if 0.012 < c - a < 0.06:
            out.append(issue("warning", "MISALIGNED", f"Left edges at {a:.3f} and {c:.3f} in almost align", sid))
            break
    # density
    maxs = profile.get("max_shapes_per_slide", 70)
    if len(shapes) > maxs:
        out.append(issue("warning", "SHAPE_COUNT", f"{len(shapes)} shapes (profile max {maxs}): the slide is busy", sid))
    for z in zones:
        if z not in used_zones:
            out.append(issue("warning", "EMPTY_ZONE", f"Zone '{z}' received no content", sid))
    for f in (manifest or {}).get("fits", []):
        if f.get("overflow"):
            continue  # reported as TEXT_OVERFLOW by measurement
        if f.get("chosen_pt") and f.get("base_pt") and f["chosen_pt"] < f["base_pt"] - 0.01:
            out.append(issue("info", "FIT_SHRUNK", f"'{f.get('what', '')[:40]}' shrunk {f['base_pt']}→{f['chosen_pt']} pt to fit", sid, zone=f.get("zone")))
    for w in (manifest or {}).get("warnings", []):
        out.append(issue("warning", w["code"], w["message"], sid, zone=w.get("zone")))
    return out


def _background_at(shapes: list[ShapeInfo], s: ShapeInfo, theme: Theme) -> str:
    """Topmost filled shape drawn before `s` that contains its ink centre."""
    if s.fill:  # the text box itself is filled
        return s.fill
    cx = s.ink.x + s.ink.w / 2
    cy = s.ink.y + s.ink.h / 2
    bg = theme.c("background")
    for o in shapes:
        if o.order >= s.order:
            break
        if o.fill and o.box.x <= cx <= o.box.r and o.box.y <= cy <= o.box.b and o.kind in ("fill", "bleed", "marker"):
            bg = o.fill
    return bg


def check(pptx_path: str, manifests: list[dict], theme: Theme, profile: dict) -> list[dict]:
    prs = Presentation(pptx_path)
    out: list[dict] = []
    for i, slide in enumerate(prs.slides, start=1):
        man = manifests[i - 1] if i - 1 < len(manifests) else None
        out += check_slide(slide, i, man, theme, profile)
    return out
