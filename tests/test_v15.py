"""v1.5: r2 governance, profile statuses, KPI dashboards, derived proof, render edges, QA semantics,
human-key privacy, multi-rater statistics, robustness signals, reporting and provenance."""
import hashlib
import json
from pathlib import Path

import pytest
from pptx import Presentation

from conftest import content_slide, mini_spec
from cpe import human
from cpe.core.planner import plan
from cpe.design.tokens import contrast_ratio, load_profile, load_theme
from cpe.pptx.builder import build
from cpe.qa import geometry

ROOT = Path(__file__).resolve().parents[1]
R2 = ROOT / "evals" / "human_reference" / "rounds" / "r2"


# ── r2 governance ────────────────────────────────────────────────────────────────────────────────

def test_r2_historical_result_is_preserved_and_recomputable():
    val = json.loads((R2 / "VALIDATION_v1.4.json").read_text())
    r = human.report(R2)
    res = val["result"]
    assert (r["challenger_wins"], r["baseline_wins"], r["ties"]) == (res["challenger_wins"], res["baseline_wins"], res["ties"]) == (35, 1, 2)
    assert r["evaluators"] == 1 == res["evaluators"]
    # the historical scorer agreement is the v1.4 scorer's (key.json), not a re-scored one
    pa = r["pair_agreement"]
    assert (pa["agree"], pa["decisive_pairs"]) == (val["scorer_agreement"]["agree"], val["scorer_agreement"]["decisive_pairs"]) == (33, 34)
    h = val["hashes"]  # the votes, key and pairs it was computed from are unchanged
    for name, digest in [*((f"votes/{k}", v) for k, v in h["votes"].items()), ("key.json", h["key.json"]), ("pairs.json", h["pairs.json"])]:
        assert hashlib.sha256((R2 / name).read_bytes()).hexdigest() == digest, name


def test_r2_is_development_data_after_mark_used():
    st = human.read_status(R2)
    assert st["used_for_calibration"] is True and st["blind"] is False
    assert "v1.5" in json.dumps(st["calibration_changes"])
    assert "VALIDATION_v1.4.json" in st["historical_validation"]


def test_profile_statuses_v15():
    prof = json.loads((ROOT / "src" / "cpe" / "qa" / "archetype_profiles.json").read_text())
    for arch in ("process", "comparison"):
        for m in ("utilization", "empty"):
            assert prof["archetypes"][arch]["metrics"][m]["status"] == "human_supported_single_rater", (arch, m)
    kpi = prof["archetypes"]["kpi_dashboard"]["metrics"]
    assert all(kpi[m].get("status") == "provisional" for m in ("utilization", "empty"))
    gates = json.loads((ROOT / "evals" / "archetype_gates.json").read_text())
    assert "provisional" in json.dumps(gates.get("mean_floor", gates)).lower() or "provisional" in json.dumps(gates).lower()


# ── KPI dashboards ───────────────────────────────────────────────────────────────────────────────

def _kpi_slide(n, long_labels=False, long_values=False):
    items = [{"value": ("€1,234.5M" if long_values else f"{10 + i}%"), "label": ("Share of customers retained after the first twelve months" if long_labels else f"Metric {i}"),
              "delta": "0 pts" if i == 0 else "+2 pts"} for i in range(n)]
    return content_slide(id=f"k{n}", headline="Retention improved on every measure we track this year", message_type="kpi_dashboard",
                         visual={"type": "kpi", "data": {"items": items}}, evidence=[{"claim": "kpis"}])


def _shapes(tmp_path, slides):
    res, _ = plan(mini_spec(slides))
    out = tmp_path / "d.pptx"
    man = build(res, out)
    return Presentation(str(out)), man, out


@pytest.mark.parametrize("n,rows", [(3, 1), (4, 1), (6, 2)])
def test_kpi_cards_adaptive_rows(tmp_path, n, rows):
    prs, man, out = _shapes(tmp_path, [_kpi_slide(n)])
    slide = prs.slides[-1]
    cards = [sh for sh in slide.shapes if "|fill|" in sh.name and sh.width < prs.slide_width * 0.5]
    tops = sorted({round(c.top / 914400, 1) for c in cards})
    assert len(cards) >= n and len(tops) == rows
    codes = {i["code"] for i in geometry.check(str(out), man, load_theme(), load_profile("standard")) if i["level"] == "error"}
    assert not codes & {"OFF_SLIDE", "OUTSIDE_ZONE", "TEXT_OVERFLOW", "LOW_CONTRAST"}


def test_kpi_long_labels_and_values_stay_inside(tmp_path):
    prs, man, out = _shapes(tmp_path, [_kpi_slide(5, long_labels=True, long_values=True)])
    codes = {i["code"] for i in geometry.check(str(out), man, load_theme(), load_profile("standard")) if i["level"] == "error"}
    assert not codes & {"OFF_SLIDE", "OUTSIDE_ZONE", "TEXT_OVERFLOW"}


def test_kpi_text_is_legible_on_the_card_fill():
    from cpe.design.tokens import activate

    theme = load_theme()
    activate(theme)

    class P:  # the painter's colour logic only
        from cpe.pptx.painter import Painter as _P
        legible = _P.legible

        def color(self, t):
            return theme.c(t)

    card = theme.c("surface")
    for token in ("neutral", "text_muted", "positive", "negative"):
        c = P().legible(token, 10, False, bg=card)
        assert contrast_ratio(c, card) >= 4.5, token


# ── derived proof ────────────────────────────────────────────────────────────────────────────────

def _proof(headline, visual):
    from cpe.qa.proof import derive, headline_quantities, slide_series

    qs = headline_quantities(headline)
    ser = slide_series({"visual": visual})
    return [derive(q, ser) for q in qs]


def test_derived_sum_percent_change_pp_share():
    bar = {"type": "bar", "unit": "€M", "data": {"categories": ["a", "b", "c"], "series": [{"name": "s", "values": [24, 19, 18]}]}}
    p = _proof("A €61M opportunity across three levers", bar)[0]
    assert p["status"] == "derived_proof" and p["operation"] == "sum"
    assert set(p) >= {"headline_value", "status", "operation", "series", "operands", "normalized_result", "tolerance"}
    line = {"type": "line", "unit": "€M", "data": {"categories": ["2023", "2025"], "series": [{"name": "cost", "values": [100, 82]}]}}
    assert _proof("Costs fell 18% in two years", line)[0]["operation"].startswith("percent change")
    pct = {"type": "line", "unit": "%", "data": {"categories": ["2023", "2025"], "series": [{"name": "m", "values": [12.1, 14.5]}]}}
    assert _proof("Margin up 2.4 pp", pct)[0]["operation"] == "percentage-point change"
    share = {"type": "bar", "unit": "€M", "data": {"categories": list("abcd"), "series": [{"name": "s", "values": [30, 25, 25, 20]}]}}
    assert _proof("Two segments make 55% of revenue", share)[0]["status"] == "derived_proof"


def test_derived_unit_normalisation_and_incompatible_units():
    bn = {"type": "bar", "unit": "€M", "data": {"categories": list("ab"), "series": [{"name": "s", "values": [700, 500]}]}}
    assert _proof("A €1.2bn programme", bn)[0]["status"] == "derived_proof"
    usd = {"type": "bar", "unit": "$M", "data": {"categories": list("abc"), "series": [{"name": "s", "values": [24, 19, 18]}]}}
    assert _proof("A €61M opportunity", usd)[0]["status"] == "unknown"  # never across currencies
    pct = {"type": "bar", "unit": "%", "data": {"categories": list("abc"), "series": [{"name": "s", "values": [24, 19, 18]}]}}
    assert _proof("A €61M opportunity", pct)[0]["status"] == "unknown"  # never money from percentages


def test_derived_proof_is_conservative():
    bar = {"type": "bar", "data": {"categories": list("abc"), "series": [{"name": "s", "values": [3, 2, 2]}]}}
    assert _proof("We see 7 levers", bar)[0]["status"] == "unknown"  # small plain integers: coincidence
    wf = {"type": "waterfall", "unit": "€bn", "data": {"steps": [{"label": "2024", "value": 1.0, "type": "total"}, {"label": "a", "value": 0.3},
                                                                {"label": "b", "value": 0.3}, {"label": "c", "value": -0.3}, {"label": "2025", "type": "total"}]}}
    p = _proof("A €0.3bn bridge", wf)[0]
    assert p["status"] == "unknown" and "ambiguous" in p["reason"]
    bar2 = {"type": "bar", "unit": "€M", "data": {"categories": list("abcde"), "series": [{"name": "s", "values": [11, 13, 17, 19, 23]}]}}
    assert _proof("A €30M gap", bar2)[0]["status"] == "unknown"  # 13+17 would need a subset search: not done


# ── render edges ─────────────────────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("steps", [
    [{"label": "2024", "value": -3, "type": "total"}, {"label": "Volume", "value": -6}, {"label": "Price", "value": 2}, {"label": "2025", "type": "total"}],
    [{"label": "2024", "value": -25, "type": "total"}, {"label": "Operating improvement", "value": 30}, {"label": "Capex", "value": -2}, {"label": "2025", "type": "total"}],
    [{"label": "Inicio", "value": 10, "type": "total"}, {"label": "Traspaso al canal online de las tiendas", "value": -14}, {"label": "Fin", "type": "total"}],
])
def test_waterfall_negative_and_zero_crossing_stays_in_zone(tmp_path, steps):
    s = content_slide(headline="The result changes sign over the period as costs move", message_type="change_bridge",
                      visual={"type": "waterfall", "title": "Bridge", "unit": "€M", "data": {"steps": steps}})
    _, man, out = _shapes(tmp_path, [s])
    issues = geometry.check(str(out), man, load_theme(), load_profile("standard"))
    assert not {i["code"] for i in issues if i["level"] == "error"} & {"OFF_SLIDE", "OUTSIDE_ZONE"}
    assert "WATERFALL_NEGATIVE" not in {i["code"] for i in issues}


def test_long_statement_fallback_ladder(tmp_path):
    from cpe.pptx import text_components as tc

    assert tc.STATEMENT_SIZES[0] == 30 and tc.STATEMENT_SIZES[-1] >= 22  # readability floor
    long = " ".join(["word"] * 70)
    st = {"id": "st1", "section": "K1", "kind": "statement", "headline": "Decision requested", "text": long}
    _, man, out = _shapes(tmp_path, [st])
    assert any(w["code"] == "STATEMENT_TOO_LONG" for m in man for w in m.get("warnings", []))
    issues = geometry.check(str(out), man, load_theme(), load_profile("standard"))
    assert not {i["code"] for i in issues if i["level"] == "error"} & {"TEXT_OVERFLOW", "OFF_SLIDE", "OUTSIDE_ZONE"}


def test_render_spans_ignore_trailing_space_advance():
    import inspect

    from cpe.qa import render_checks

    assert "rawdict" in inspect.getsource(render_checks._spans)


# ── QA semantics ─────────────────────────────────────────────────────────────────────────────────

def test_visual_and_authoring_qa_are_separate():
    from cpe.qa.report import is_authoring, qa_semantics

    assert is_authoring("HEADLINE_TOPIC") and is_authoring("STORY_NO_ASK") and is_authoring("INTENT_EVIDENCE")
    assert not is_authoring("HEADLINE_LINES") and not is_authoring("RENDER_TEXT_SPILL")
    s = qa_semantics([{"level": "error", "code": "HEADLINE_TOPIC"}, {"level": "warning", "code": "OUTSIDE_ZONE"}])
    assert s["visual_qa_passed"] and not s["authoring_qa_passed"]
    s = qa_semantics([{"level": "error", "code": "OUTSIDE_ZONE"}])
    assert not s["visual_qa_passed"] and s["authoring_qa_passed"]


# ── human-key privacy and multi-rater statistics ────────────────────────────────────────────────

def _fake_runs(tmp_path):
    from PIL import Image

    roots = []
    for name, color, score in (("base", "red", 60.0), ("chal", "blue", 70.0)):
        d = tmp_path / name / "deck1"
        (d / "renders").mkdir(parents=True)
        slides = ["s1", "s2", "s3", "s4"]
        (d / "resolved.json").write_text(json.dumps({"slides": [{"id": s, "kind": "content", "_plan": {"layout": {"id": "exhibit_full"}}} for s in slides]}))
        (d / "qa_report.json").write_text(json.dumps({"composition": {"slides": [{"slide_id": s, "score": score, "archetype": "chart"} for s in slides]}}))
        for i, _ in enumerate(slides, 1):
            Image.new("RGB", (40, 22), color).save(d / "renders" / f"slide-{i:02d}.png")
        roots.append(tmp_path / name)
    return roots


def test_private_key_round_bundle_reveals_nothing(tmp_path):
    base, chal = _fake_runs(tmp_path)
    rd, kp = tmp_path / "r9", tmp_path / "private" / "r9" / "key.json"
    r = human.build_round(rd, [("v1.4.0->v1.5.0", str(base), str(chal))], n=4, repeats=1, purpose="v1.4.0 vs v1.5.0", key_out=kp)
    assert r["pairs"] == 4 and kp.exists() and not (rd / "key.json").exists()
    bundle = "".join((rd / f).read_text() for f in ("STATUS.json", "pairs.json", "index.html"))
    for leak in ("v1.4", "v1.5", "baseline", "challenger", "exhibit_full", "score", "chart", "deck1", "s1"):
        assert leak not in bundle, leak
    assert {p.name for p in rd.iterdir()} <= set(human.BUNDLE)
    assert (rd / "key.sha256").read_text().split()[0] == hashlib.sha256(kp.read_bytes()).hexdigest()
    with pytest.raises(SystemExit):
        human.report(rd)  # no key in the round and none at the default location → explicit --key
    pairs = json.loads((rd / "pairs.json").read_text())["pairs"]
    for e in ("ana", "bo"):
        for p in pairs:
            human.record_vote(rd, {"evaluator": e, "pair": p["id"], "left": p["images"][0], "right": p["images"][1], "choice": "tie" if e == "bo" else "left"})
    rep = human.report(rd, kp)
    assert rep["evaluators"] == 2 and set(rep["by_rater"]) == {"ana", "bo"}
    assert rep["by_rater"]["bo"]["ties"] == 3 and "fleiss_kappa" in rep["inter_rater"]
    tampered = tmp_path / "other.json"
    tampered.write_text(kp.read_text().replace("chart", "table"))
    with pytest.raises(SystemExit):
        human.report(rd, tampered)  # the commitment protects against a swapped key
    st = human.close_round(rd, kp)
    assert (rd / "key.json").exists() and st["blind"] is False


def test_one_rater_voting_twice_is_one_rater(tmp_path):
    base, chal = _fake_runs(tmp_path)
    rd = tmp_path / "r8"
    human.build_round(rd, [("a->b", str(base), str(chal))], n=4, repeats=0, key_out=tmp_path / "k" / "key.json")
    p = json.loads((rd / "pairs.json").read_text())["pairs"][0]
    for choice in ("left", "right"):
        human.record_vote(rd, {"evaluator": "ana", "pair": p["id"], "left": p["images"][0], "right": p["images"][1], "choice": choice})
    rep = human.report(rd, tmp_path / "k" / "key.json")
    assert rep["comparisons"] == 1 and rep["evaluators"] == 1 and "note" in rep["inter_rater"]


def test_votes_folder_is_created_and_utf8(tmp_path):
    base, chal = _fake_runs(tmp_path)
    rd = tmp_path / "r7"
    human.build_round(rd, [("a->b", str(base), str(chal))], n=2, repeats=0, key_out=tmp_path / "k7" / "key.json")
    (rd / "votes").rmdir()  # a fresh git clone has no empty folders
    p = json.loads((rd / "pairs.json").read_text(encoding="utf-8"))["pairs"][0]
    human.record_vote(rd, {"evaluator": "ana", "pair": p["id"], "left": p["images"][0], "right": p["images"][1], "choice": "left"})
    assert human.done_pairs(rd, "ana") == [p["id"]]
    assert "encoding=\"utf-8\"" in Path(human.__file__).read_text()


def test_r1_r2_history_not_rewritten():
    for r in ("r1", "r2"):
        assert (ROOT / "evals" / "human_reference" / "rounds" / r / "key.json").exists()


# ── robustness ───────────────────────────────────────────────────────────────────────────────────

def _row(delta, layout=False, font=0, qa=(), arch="process", pert="more_steps"):
    return {"seed": "s", "perturbation": pert, "archetype": arch, "seed_score": 90, "score": 90 + delta, "delta": delta, "layout_change": layout,
            "font_drop": font, "new_visual_qa": list(qa), "new_flags": [], "catastrophic": delta <= -25 or bool(qa)}


def test_robustness_signals_and_gates():
    from cpe import robustness as rb

    rows = [_row(0), _row(-6, layout=True), _row(-12), _row(-1, layout=True), _row(0, font=2, arch="table", pert="more_rows")]
    s = rb.summarize(rows)
    assert s["meaningful_drop_rate"] == 0.4 and s["large_drop_rate"] == 0.2 and s["max_drop"] == 12
    assert s["layout_change_rate"] == 0.4 and s["layout_change_with_quality_drop_rate"] == 0.2  # a layout change alone is not bad
    assert s["font_drop_rate"] == 0.2 and set(s["by_archetype"]) == {"process", "table"} and "p95_drop" in s
    base = {"variants": 5, "catastrophic": 0, "new_visual_errors": 0, "p90_drop": 0.0, "large_drop_rate": 0.0, "font_drop_rate": 0.0,
            "layout_change_with_quality_drop_rate": 0.0}
    assert rb.compare(s, base) == []
    assert rb.provisional_breaches(s, base)  # reported, not enforced
    s2 = rb.summarize(rows + [_row(-2, qa=["RENDER_TEXT_SPILL"])])
    assert any("new visual QA" in x for x in rb.compare(s2, base))
    assert any("coverage" in x for x in rb.compare(rb.summarize(rows[:3]), base))


# ── reporting and provenance ─────────────────────────────────────────────────────────────────────

def test_slide_counts_and_provenance_terms(tmp_path):
    from cpe import environment
    from cpe.evals import _aggregate
    from cpe.results_report import render_block

    agg = _aggregate([{"slides_authored": 3, "slides_resolved": 4, "slides": {"a": {"flags": [], "score": 90}}}])
    assert (agg["slides_authored"], agg["slides_resolved"], agg["slides_measured"]) == (3, 4, 1)
    pv = environment.provenance()
    assert {"evaluated_source_commit", "evaluated_source_dirty", "working_tree_commit", "working_tree_dirty"} <= set(pv)
    block = render_block({"regression": {"provenance": {"evaluated_source_commit": "abc", "evaluated_source_dirty": False, "dirty": False}}})
    assert "-dirty" not in block and "DIRTY" not in block and "evaluated source commit" in block
    block = render_block({"regression": {"provenance": {"evaluated_source_commit": "abc", "evaluated_source_dirty": True}}})
    assert "evaluated source DIRTY" in block


def test_result_files_do_not_make_the_source_dirty(monkeypatch):
    from cpe import environment

    monkeypatch.setattr(environment, "_run", lambda cmd: "evals/results/latest.json\nREADME.md" if "diff" in cmd else "0" * 40)
    assert environment.dirty_paths() == []
    assert not environment.git_commit().endswith("-dirty")


def test_intake_status_never_substitutes(tmp_path, monkeypatch):
    from cpe import private_holdout as ph

    assert ph.external_status(tmp_path / "ext")["status"] == "EXTERNAL HOLDOUT: AWAITING INPUT"
    assert ph.corporate_status(tmp_path / "corp")["status"] == "UNSEEN CORPORATE TEMPLATE: AWAITING USER-SUPPLIED TEMPLATE"
    (tmp_path / "ext" / "decks").mkdir(parents=True)
    (tmp_path / "ext" / "decks" / "d.json").write_text("{}")
    assert "NOT INDEPENDENT" in ph.external_status(tmp_path / "ext")["status"]  # no attested outside author
    (tmp_path / "corp").mkdir()
    (tmp_path / "corp" / "template.pptx").write_bytes(b"known")
    monkeypatch.setattr(ph, "known_development_hashes", lambda: {hashlib.sha256(b"known").hexdigest()})
    assert "REFUSED" in ph.corporate_status(tmp_path / "corp")["status"]


def test_ci_never_runs_private_or_human_validation():
    ci = (ROOT / ".github" / "workflows" / "ci.yml").read_text()
    for forbidden in ("holdout v15", "v15-external", "v15-corporate", "holdout private", "holdout external", "human report", "holdout_v2"):
        assert forbidden not in ci, forbidden


def test_benchmark_specs_are_editorially_valid():
    from cpe.core.headline import lint_headline
    from cpe.core.storyline import lint_storyline

    for h in ("Customers rate the new app 4.7 out of 5", "La rotación de plantilla bajó al 9,8%", "La plataforma de datos tiene cuatro capas"):
        assert "HEADLINE_TOPIC" not in {i["code"] for i in lint_headline(h)[1]}, h
    for h in ("Lessons from the pilot", "Revenue evolution"):
        assert "HEADLINE_TOPIC" in {i["code"] for i in lint_headline(h)[1]}, h  # lint coverage kept
    slides = [{"id": f"s{i}", "kind": "content", "headline": "Revenue grew 12%"} for i in range(6)]
    spec = {"storyline": {"framework": "SCR", "governing_thought": "x", "key_line": []}, "slides": slides}
    assert "STORY_NO_EXEC_SUMMARY" in {i["code"] for i in lint_storyline(spec)}
    spec["storyline"]["collection"] = True
    assert "STORY_NO_EXEC_SUMMARY" not in {i["code"] for i in lint_storyline(spec)}
    for f in (ROOT / "evals" / "regression" / "cases").glob("2*_battery_*.json"):
        assert json.loads(f.read_text())["storyline"]["collection"] is True
