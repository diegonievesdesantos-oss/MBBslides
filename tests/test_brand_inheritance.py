"""v1.3: the brand is read through PowerPoint's inheritance chain, not from theme slot names."""
import pytest
from pptx import Presentation

from conftest import content_slide, mini_spec
from cpe.brand.fixtures import make_inherited_styles
from cpe.brand.ingest import ingest
from cpe.brand.model import analyse, family_name
from cpe.brand.rules import deck_advice
from cpe.design.tokens import activate, contrast_ratio, load_theme


@pytest.fixture(scope="module")
def tpl(tmp_path_factory):
    return make_inherited_styles(tmp_path_factory.mktemp("inh") / "inh.pptx")


@pytest.fixture(autouse=True)
def _restore():
    yield
    activate(load_theme("meridian"))


def test_weight_names_belong_to_one_family():
    assert family_name("Inter Black") == family_name("Inter-SemiBold") == family_name("Inter ExtraBold") == "Inter"
    assert family_name("Inter Tight") == "Inter Tight"  # a different family, not a weight
    assert family_name("Arial") == "Arial"


def test_inherited_font_wins_over_theme_and_conflict_is_reported(tpl):
    ty = analyse(Presentation(tpl), set())["typography"]
    assert ty["primary"] == "Inter"
    assert ty["roles"]["heading"]["font"] == "Inter" and ty["roles"]["body"]["font"] == "Inter"
    assert ty["conflict"]["declared_theme_font"] == "Arial" and ty["conflict"]["observed_primary_font"] == "Inter"
    assert ty["inherited_from"].get("master placeholder", 0) > 0  # no direct formatting anywhere


def test_colour_roles_come_from_what_is_drawn_not_slot_names(tpl, tmp_path):
    m = analyse(Presentation(tpl), set())
    assert m["palette"]["text"] == "1F2933" and m["palette"]["page"] == "FAFAF7"
    rep = ingest(tpl, tmp_path / "b")
    c = rep["color_mapping"]
    assert c["background"] == "FAFAF7" and c["text"] == "1F2933"
    for role in ("text", "text_muted"):
        assert contrast_ratio(c[role], c["background"]) >= 4.5
    assert contrast_ratio("FFFFFF", c["secondary"]) >= 4.5 and contrast_ratio("FFFFFF", c["primary"]) >= 4.5


def test_closing_learned_from_the_last_slide_and_bookend_advice(tpl, tmp_path):
    m = analyse(Presentation(tpl), set())
    sec = next(c for c in m["layouts"] if c["layout"] == "Section Header")
    assert any(x["type"] == "closing" for x in sec["classification"])
    assert m["rules"]["bookend"]["first_and_last_in_brand_colour"] is True
    ingest(tpl, tmp_path / "b2")
    theme = load_theme(str(tmp_path / "b2"))
    spec = mini_spec([content_slide(id=f"s{i}") for i in range(5)])
    adv = deck_advice({"slides": [{"id": "c", "kind": "cover"}] + spec["slides"]}, [], theme)
    assert {a["code"] for a in adv} >= {"BRAND_BOOKEND"}
    ok = deck_advice({"slides": [{"id": "c", "kind": "cover"}, *spec["slides"], {"id": "e", "kind": "closing"}]},
                     [{"slide_id": "e", "corporate": {"mode": "native"}}, {"slide_id": "c", "corporate": {"mode": "native"}}], theme)
    assert "BRAND_BOOKEND" not in {a["code"] for a in ok}


def test_installed_fonts_are_measured_with_the_real_font():
    from cpe.design import text_metrics as tm

    assert tm.family_for("Arial") == "LiberationSans"  # metric twins keep priority
    fam = tm.family_for("DejaVu Sans")  # installed everywhere the engine runs
    assert tm._font_path(fam, False) and tm._font_path(fam, False).endswith(".ttf")
