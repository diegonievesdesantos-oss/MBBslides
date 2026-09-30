"""Corporate template intelligence on SYNTHETIC templates (no corporate material in the repository)."""
import json

import pytest
from pptx import Presentation

from conftest import content_slide, mini_spec
from cpe.brand.fixtures import make_multimaster
from cpe.brand.ingest import ingest
from cpe.brand.model import analyse
from cpe.core.planner import plan
from cpe.design.tokens import activate, load_theme
from cpe.pptx.builder import build


@pytest.fixture(scope="module")
def mm(tmp_path_factory):
    d = tmp_path_factory.mktemp("mm")
    return make_multimaster(d / "mm.pptx"), d


@pytest.fixture(scope="module")
def model(mm):
    return analyse(Presentation(mm[0]), set())


@pytest.fixture(autouse=True)
def _restore_grid():
    yield
    activate(load_theme("meridian"))


def _cls(model, name):
    return next(c for c in model["layouts"] if c["layout"] == name)["classification"]


def test_every_master_and_layout_is_detected_no_index_zero_assumption(model):
    assert [m["name"] for m in model["masters"]] == ["Executive", "Analytical", "Narrative"]
    assert len(model["layouts"]) == 11
    assert {c["master_id"] for c in model["layouts"]} == {"m1", "m2", "m3"}
    themes = {m["theme"]["fonts"]["majorFont"] for m in model["masters"]}
    assert themes == {"Georgia", "Arial", "Calibri"}  # one theme per master, all read


def test_layout_classification_by_geometry_not_names(model):
    assert _cls(model, "Cover")[0]["type"] == "cover"
    assert _cls(model, "Layout 2")[0]["type"] == "section"  # poor name, section by geometry + background + usage
    assert _cls(model, "CUSTOM_4_1_2")[0]["type"] == "table"  # poor name, table placeholder
    assert _cls(model, "Matrix")[0]["type"] == "matrix"
    assert _cls(model, "Process")[0]["type"] == "process"
    assert _cls(model, "Title only")[0]["type"] == "content"
    for c in model["layouts"]:
        assert all(0 < x["confidence"] <= 1 for x in c["classification"])


def test_layout_meaning_is_learned_from_example_slides(model):
    c = next(c for c in model["layouts"] if c["layout"] == "CUSTOM_3_1_1")
    assert c["usage"] == {"image_split": 3}
    assert c["classification"][0]["type"] == "image_split"  # geometry alone says two columns; usage says picture + text
    assert any(x["type"] == "two_column" for x in c["classification"])  # ambiguity is kept, with a lower confidence


def test_declared_theme_font_vs_observed_usage_conflict(model):
    ty = model["typography"]
    assert ty["primary"] == "Inter"
    assert ty["conflict"]["declared_theme_font"] in ("Arial", "Georgia", "Calibri")
    assert ty["conflict"]["observed_primary_font"] == "Inter"
    assert ty["conflict"]["style_guide_mentions"].get("Inter")
    assert ty["roles"]["body"]["conflict"] is True


def test_missing_corporate_font_is_reported_never_silent(mm):
    rep = ingest(mm[0], mm[1] / "brand", name="Synthetic")
    f = rep["fonts"]["body"]
    assert f["font"] == "Inter" and not f["installed"] and f["measurement"] == "approximate"
    assert "FONT WARNING" in f["warning"] and f["render_fallback"]
    md = (mm[1] / "brand" / "compatibility.md").read_text()
    assert "Conflict detected" in md and "FONT WARNING" in md and "Masters: **3**" in md


def test_assets_logos_and_rules(model):
    assert model["assets"]["logo_positions"] == {"top-right": 1, "bottom-right": 1}
    assert model["rules"]["headline_case"]["dominant"] == "sentence"
    assert model["rules"]["shapes"]["chevrons"] >= 3
    assert model["rules"]["bookend"]["first_and_last_share_a_colour"] is True


def test_rescaled_16_9_template_uses_its_masters(tmp_path):
    p = make_multimaster(tmp_path / "mm10.pptx", width_in=10.0)
    rep = ingest(p, tmp_path / "b10", name="S10")
    assert rep["masters_used"] and rep["slide_size"]["rescaled"]["factor"] == pytest.approx(1.3333, abs=1e-3)
    prs = Presentation(str(tmp_path / "b10" / "template.pptx"))
    assert round(prs.slide_width / 914400, 2) == 13.33 and len(prs.slides) == 0 and len(prs.slide_masters) == 3


def test_layout_matching_picks_layouts_from_different_masters(mm):
    brand = mm[1] / "brand_m"
    ingest(mm[0], brand, name="Synthetic")
    spec = mini_spec([content_slide(id="s1"), {"id": "d1", "kind": "divider", "title": "Where the market is going"}])
    spec["slides"].insert(0, {"id": "c1", "kind": "cover", "title": "Brand test"})
    spec["slides"].append({"id": "e1", "kind": "closing", "title": "Thank you"})
    spec["meta"]["brand"] = str(brand)
    res, _ = plan(spec)
    man = build(res, mm[1] / "deck.pptx")
    modes = {m["slide_id"]: m["corporate"] for m in man}
    assert modes["c1"]["mode"] == modes["d1"]["mode"] == modes["e1"]["mode"] == "native"
    assert modes["c1"]["layout"] == "Cover" and modes["e1"]["layout"] == "End"
    assert modes["s1"]["mode"] == "adaptive"
    # layouts whose artwork crosses the engine footer are rejected with a reason, not used silently
    assert any("footer band" in r for r in modes["s1"].get("rejected", []))
    prs = Presentation(str(mm[1] / "deck.pptx"))
    masters = {s.slide_layout.slide_master.name for s in prs.slides}
    assert len(masters) >= 2
    cover = prs.slides[0]
    assert cover.shapes.title.text == "Brand test"  # native placeholder filled, template typography inherited
    assert not any(ph.has_text_frame and not ph.text_frame.text.strip() for ph in cover.placeholders)


def test_single_master_template_still_works(tmp_path):
    import subprocess
    import sys

    from conftest import ROOT

    out = tmp_path / "tpl.pptx"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "make_sample_template.py"), str(out)], check=True, capture_output=True)
    rep = ingest(out, tmp_path / "k")
    assert len(rep["masters"]) == 1 and rep["masters_used"]
    assert json.loads((tmp_path / "k" / "layout_catalog.json").read_text())
