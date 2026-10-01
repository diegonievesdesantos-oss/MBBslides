"""Synthetic corporate templates for tests and docs (no real corporate material).

    make_multimaster(path, width_in=13.333)

builds a template with THREE slide masters, each with its own theme (fonts + colours),
background and artwork, deliberately poor layout names on some layouts, and example slides
that show how some layouts are really used:

    Master A — executive   Cover · "Layout 2" (a section divider) · End (closing)
    Master B — analytical  Chart · Chart + commentary · "CUSTOM_4_1_2" (a table) · Matrix
    Master C — narrative   Title and text · "CUSTOM_3_1_1" (two content areas; used with a picture
                           + text on the examples) · Process · Title only

Typography is contradictory on purpose: every theme declares a font, but the example slides set
most text in "Inter" by direct formatting and a style-guide slide names it — the case where
`theme XML ≠ actual brand`.

python-pptx cannot add masters, so the parts are created at the OPC level: master, layouts and
theme parts related to each other and registered in presentation.xml with unique ids.
"""
from __future__ import annotations

import copy
import io

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.parts.slide import SlideLayoutPart, SlideMasterPart
from pptx.util import Emu, Inches, Pt

NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PH_TYPE = {"title": 'type="title"', "ctrTitle": 'type="ctrTitle"', "subTitle": 'type="subTitle" idx="1"', "body": 'type="body" idx="{i}"',
           "obj": 'idx="{i}"', "pic": 'type="pic" idx="{i}"', "chart": 'type="chart" idx="{i}"', "tbl": 'type="tbl" idx="{i}"',
           "dt": 'type="dt" sz="half" idx="10"', "ftr": 'type="ftr" sz="quarter" idx="11"', "sldNum": 'type="sldNum" sz="quarter" idx="12"'}


def _emu(v: float, k: float) -> int:
    return int(round(v * k * 914400))


def _ph(sid: int, kind: str, x, y, w, h, k: float, i: int = 1) -> str:
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{kind} {sid}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph {PH_TYPE[kind].format(i=i)}/></p:nvPr></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{_emu(x, k)}" y="{_emu(y, k)}"/><a:ext cx="{_emu(w, k)}" cy="{_emu(h, k)}"/></a:xfrm></p:spPr>'
            f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>')


def _rect(sid: int, name: str, x, y, w, h, color: str, k: float, prst: str = "rect") -> str:
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{_emu(x, k)}" y="{_emu(y, k)}"/><a:ext cx="{_emu(w, k)}" cy="{_emu(h, k)}"/></a:xfrm>'
            f'<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')


def _tree(items: str) -> str:
    return ('<p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>' + items + '</p:spTree>')


def _bg(color: str | None) -> str:
    return f'<p:bg><p:bgPr><a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>' if color else ""


MASTERS = [
    {"name": "Executive", "major": "Georgia", "minor": "Arial", "bg": "1F2A44", "text": "FFFFFF", "accent": "E8A33D",
     "art": [], "title_sz": 4000,
     "layouts": [
         {"name": "Cover", "type": "title", "ph": [("ctrTitle", 0.8, 2.6, 11.7, 1.5), ("subTitle", 0.8, 4.2, 11.7, 0.8)]},
         {"name": "Layout 2", "type": "cust", "ph": [("title", 0.8, 3.0, 11.7, 1.2)], "art": [("Rule", 0.8, 2.7, 1.2, 0.08, "E8A33D")]},
         {"name": "End", "type": "cust", "ph": [("title", 0.8, 3.0, 11.7, 1.2)]},
     ]},
    {"name": "Analytical", "major": "Arial", "minor": "Arial", "bg": "FFFFFF", "text": "1B1B1B", "accent": "0B6E4F",
     "art": [("Logo", 12.3, 0.2, 0.7, 0.4, "0B6E4F"), ("Footer bar", 0.0, 7.25, 13.333, 0.25, "0B6E4F")], "title_sz": 2400,
     "layouts": [
         {"name": "Chart", "type": "cust", "ph": [("title", 0.6, 0.4, 11.4, 0.9), ("chart", 0.6, 1.6, 12.1, 5.3), ("sldNum", 12.0, 6.9, 0.8, 0.3)]},
         {"name": "Chart + commentary", "type": "cust", "ph": [("title", 0.6, 0.4, 11.4, 0.9), ("chart", 0.6, 1.6, 7.8, 5.3), ("body", 8.7, 1.6, 4.0, 5.3)]},
         {"name": "CUSTOM_4_1_2", "type": "cust", "ph": [("title", 0.6, 0.4, 11.4, 0.9), ("tbl", 0.6, 1.6, 12.1, 5.3)]},
         {"name": "Matrix", "type": "cust", "ph": [("title", 0.6, 0.4, 11.4, 0.9), ("obj", 0.6, 1.6, 5.9, 2.5), ("obj", 6.8, 1.6, 5.9, 2.5),
                                                   ("obj", 0.6, 4.3, 5.9, 2.5), ("obj", 6.8, 4.3, 5.9, 2.5)]},
     ]},
    {"name": "Narrative", "major": "Calibri", "minor": "Calibri", "bg": "F5F3EF", "text": "222222", "accent": "C0392B",
     "art": [("Logo", 12.3, 6.9, 0.7, 0.4, "C0392B")], "title_sz": 2800,
     "layouts": [
         {"name": "Title and text", "type": "obj", "ph": [("title", 0.6, 0.4, 12.1, 0.9), ("body", 0.6, 1.6, 12.1, 5.1)]},
         {"name": "CUSTOM_3_1_1", "type": "cust", "ph": [("title", 0.6, 0.4, 12.1, 0.9), ("obj", 0.6, 1.6, 5.9, 5.1), ("obj", 6.8, 1.6, 5.9, 5.1)]},
         {"name": "Process", "type": "cust", "ph": [("title", 0.6, 0.4, 12.1, 0.9)] + [("body", 0.6 + i * 2.46, 2.2, 2.26, 3.2) for i in range(5)]},
         {"name": "Title only", "type": "titleOnly", "ph": [("title", 0.6, 0.4, 12.1, 0.9)]},
     ]},
]


def _theme_xml(base_blob: bytes, m: dict, name: str) -> bytes:
    root = etree.fromstring(base_blob)
    root.set("name", name)
    cs = root.find(f".//{A}clrScheme")
    cs.set("name", name)
    for slot, val in (("dk1", m["text"] if m["bg"] != "1F2A44" else "1B1B1B"), ("lt1", "FFFFFF"), ("accent1", m["accent"])):
        el = cs.find(f"{A}{slot}")
        for ch in list(el):
            el.remove(ch)
        etree.SubElement(el, f"{A}srgbClr", val=val)
    fs = root.find(f".//{A}fontScheme")
    fs.find(f"{A}majorFont/{A}latin").set("typeface", m["major"])
    fs.find(f"{A}minorFont/{A}latin").set("typeface", m["minor"])
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


BRAND_FONT = ["Inter"]


def make_multimaster(path, width_in: float = 13.333, with_examples: bool = True, brand_font: str = "Inter") -> str:
    """brand_font: the typeface set by direct formatting on the examples (pass a non-existent name
    to test the missing-font path independently of the fonts installed)."""
    BRAND_FONT[0] = brand_font
    k = width_in / 13.333
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(_emu(13.333, k)), Emu(_emu(7.5, k))
    pkg = prs.part.package
    base_master_part = prs.slide_masters[0].part
    base_theme = base_master_part.part_related_by(RT.THEME)
    theme_blob = base_theme.blob
    # detach the default master from the presentation (unreachable parts are not saved)
    lst = prs.part._element.find(f"{P}sldMasterIdLst")
    for el in list(lst):
        prs.part.drop_rel(el.get(f"{R}id"))
        lst.remove(el)
    next_id = 2147483648
    for mi, m in enumerate(MASTERS, 1):
        art = "".join(_rect(100 + j, a[0], *a[1:5], a[5], k) for j, a in enumerate(m["art"]))
        phs = _ph(2, "title", 0.6, 0.4, 12.1, 0.9, k) + _ph(3, "body", 0.6, 1.6, 12.1, 5.1, k)
        master_xml = (f'<p:sldMaster {NS}><p:cSld name="{m["name"]}">{_bg(m["bg"])}{_tree(phs + art)}</p:cSld>'
                      '<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" '
                      'accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/><p:sldLayoutIdLst/>'
                      f'<p:txStyles><p:titleStyle><a:lvl1pPr><a:defRPr sz="{m["title_sz"]}" b="1"><a:solidFill><a:srgbClr val="{m["text"]}"/></a:solidFill>'
                      '<a:latin typeface="+mj-lt"/></a:defRPr></a:lvl1pPr></p:titleStyle>'
                      f'<p:bodyStyle><a:lvl1pPr><a:defRPr sz="1600"><a:solidFill><a:srgbClr val="{m["text"]}"/></a:solidFill><a:latin typeface="+mn-lt"/></a:defRPr></a:lvl1pPr></p:bodyStyle>'
                      '<p:otherStyle/></p:txStyles></p:sldMaster>')
        mpart = SlideMasterPart(PackURI(f"/ppt/slideMasters/slideMaster{20 + mi}.xml"), CT.PML_SLIDE_MASTER, pkg, parse_xml(master_xml.encode()))
        tpart = Part(PackURI(f"/ppt/theme/theme{20 + mi}.xml"), CT.OFC_THEME, pkg, _theme_xml(theme_blob, m, f"{m['name']} theme"))
        mpart.relate_to(tpart, RT.THEME)
        lid_lst = mpart._element.find(f"{P}sldLayoutIdLst")
        for li, lay in enumerate(m["layouts"], 1):
            items = ""
            n = 2
            counters = {}
            for ph in lay["ph"]:
                kind = ph[0]
                counters[kind] = counters.get(kind, 0) + 1
                items += _ph(n, kind, *ph[1:5], k, i=n)
                n += 1
            for a in lay.get("art", []):
                items += _rect(n, a[0], *a[1:5], a[5], k)
                n += 1
            lay_xml = (f'<p:sldLayout {NS} type="{lay["type"]}" preserve="1"><p:cSld name="{lay["name"]}">{_tree(items)}</p:cSld>'
                       '<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>')
            lpart = SlideLayoutPart(PackURI(f"/ppt/slideLayouts/slideLayout{100 + mi * 20 + li}.xml"), CT.PML_SLIDE_LAYOUT, pkg, parse_xml(lay_xml.encode()))
            lpart.relate_to(mpart, RT.SLIDE_MASTER)
            rid = mpart.relate_to(lpart, RT.SLIDE_LAYOUT)
            next_id += 1
            etree.SubElement(lid_lst, f"{P}sldLayoutId", {"id": str(next_id), f"{R}id": rid})
        rid = prs.part.relate_to(mpart, RT.SLIDE_MASTER)
        next_id += 1
        el = etree.SubElement(lst, f"{P}sldMasterId", {"id": str(next_id), f"{R}id": rid})
        lst.insert(mi - 1, el)
    if with_examples:
        _examples(prs, k)
    prs.save(str(path))
    return str(path)


def _layout(prs, master_name: str, layout_name: str):
    for m in prs.slide_masters:
        if m.name == master_name:
            for lay in m.slide_layouts:
                if lay.name == layout_name:
                    return lay
    raise KeyError((master_name, layout_name))


def _text(slide, x, y, w, h, text, k, size=14, font=None, bold=False):
    font = font or BRAND_FONT[0]
    tb = slide.shapes.add_textbox(Inches(x * k), Inches(y * k), Inches(w * k), Inches(h * k))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    for p in tb.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size * k)
            r.font.name = font
            r.font.bold = bold
    return tb


def _set_title(slide, text, font=None):
    font = font or BRAND_FONT[0]
    t = slide.shapes.title
    t.text_frame.text = text
    for r in t.text_frame.paragraphs[0].runs:
        r.font.name = font


def _png(color=(200, 80, 60)) -> io.BytesIO:
    from PIL import Image

    b = io.BytesIO()
    Image.new("RGB", (64, 48), color).save(b, "PNG")
    b.seek(0)
    return b


def _examples(prs, k: float) -> None:
    s = prs.slides.add_slide(_layout(prs, "Executive", "Cover"))
    _set_title(s, "Annual strategy review")
    s.placeholders[1].text_frame.text = "Board meeting, March"
    s = prs.slides.add_slide(_layout(prs, "Executive", "Layout 2"))
    _set_title(s, "Where the market is going")
    s = prs.slides.add_slide(_layout(prs, "Analytical", "Chart"))
    _set_title(s, "Revenue grew in every region last year")
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    cd = CategoryChartData()
    cd.categories = ["North", "South", "East"]
    cd.add_series("Revenue", (10, 12, 9))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6 * k), Inches(1.6 * k), Inches(12 * k), Inches(5 * k), cd)
    s = prs.slides.add_slide(_layout(prs, "Analytical", "CUSTOM_4_1_2"))
    _set_title(s, "Three plants explain most of the cost gap")
    tbl = s.shapes.add_table(4, 3, Inches(0.6 * k), Inches(1.6 * k), Inches(12 * k), Inches(3 * k)).table
    for r in range(4):
        for c in range(3):
            tbl.cell(r, c).text = f"r{r}c{c}"
    for i in range(3):  # the poorly named two-area layout is used for picture + text
        s = prs.slides.add_slide(_layout(prs, "Narrative", "CUSTOM_3_1_1"))
        _set_title(s, f"Customers value speed over price in segment {i + 1}")
        s.shapes.add_picture(_png(), Inches(0.6 * k), Inches(1.6 * k), Inches(5.9 * k), Inches(5.0 * k))
        _text(s, 6.8, 1.6, 5.9, 4.0, "Interviews with forty customers show that delivery time decides the purchase.", k)
    s = prs.slides.add_slide(_layout(prs, "Narrative", "Process"))
    _set_title(s, "The programme runs in four steps")
    for i in range(4):
        sh = s.shapes.add_shape(52 if i else 51, Inches((0.6 + i * 3.0) * k), Inches(2.0 * k), Inches(2.9 * k), Inches(0.6 * k))  # homePlate, chevron
        sh.text_frame.text = f"Step {i + 1}"
    s = prs.slides.add_slide(_layout(prs, "Narrative", "Title only"))
    _set_title(s, "Our typography")
    f = BRAND_FONT[0]
    _text(s, 0.6, 1.6, 12, 1.0, f"Typography: {f} is our primary typeface for titles and text.", k, size=20)
    _text(s, 0.6, 3.0, 12, 2.0, f"Use {f} Bold for headlines and {f} Regular for body copy. Never use more than two weights on a slide.", k)
    s = prs.slides.add_slide(_layout(prs, "Narrative", "Title and text"))
    _set_title(s, "What we recommend")
    body = next(ph for ph in s.placeholders if ph.placeholder_format.idx != 0)
    body.text_frame.text = "Invest in delivery speed before cutting prices."
    for r in body.text_frame.paragraphs[0].runs:
        r.font.name = BRAND_FONT[0]
    s = prs.slides.add_slide(_layout(prs, "Executive", "End"))
    _set_title(s, "Thank you")


def clone(path_in, path_out) -> None:
    """Byte copy helper for tests (keeps fixtures immutable)."""
    import shutil

    shutil.copy(path_in, path_out)


__all__ = ["make_multimaster", "MASTERS", "clone", "copy"]


def make_inherited_styles(path, n_content: int = 6) -> str:
    """Synthetic template where the brand lives in the INHERITANCE chain, not in the theme:

    * theme fonts say Arial; the master's title / body placeholders are styled "Inter Black" /
      "Inter" and the example text carries no direct formatting (it inherits them);
    * colour slots are used unconventionally: dk1 is a vivid accent, lt1 the dark text colour,
      the master background is dk2 (light) — so slot names cannot be trusted;
    * the deck opens on an accent-coloured cover and closes on an accent-coloured divider layout.
    """
    from pptx.dml.color import RGBColor  # noqa: F401

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(_emu(13.333, 1)), Emu(_emu(7.5, 1))
    master = prs.slide_masters[0]
    theme_part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    cs = root.find(f".//{A}clrScheme")
    for slot, val in (("dk1", "E8590C"), ("lt1", "1F2933"), ("dk2", "FAFAF7"), ("lt2", "F1EFEA"), ("accent1", "EDE9E3"), ("accent2", "F4C542"), ("accent3", "B8D8E0")):
        el = cs.find(f"{A}{slot}")
        for ch in list(el):
            el.remove(ch)
        etree.SubElement(el, f"{A}srgbClr", val=val)
    fs = root.find(f".//{A}fontScheme")
    fs.find(f"{A}majorFont/{A}latin").set("typeface", "Arial")
    fs.find(f"{A}minorFont/{A}latin").set("typeface", "Arial")
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    mel = master._element
    cSld = mel.find(f"{P}cSld")
    bg = etree.fromstring(f'<p:bg {NS}><p:bgPr><a:solidFill><a:schemeClr val="dk2"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>')
    cSld.insert(0, bg)
    for sh in master.shapes:
        if not sh.is_placeholder:
            continue
        t = sh.placeholder_format.type
        face = "Inter Black" if t == 1 else "Inter"
        txb = sh._element.find(f"{P}txBody")
        lst = txb.find(f"{A}lstStyle")
        for ch in list(lst):
            lst.remove(ch)
        lv = etree.SubElement(lst, f"{A}lvl1pPr")
        d = etree.SubElement(lv, f"{A}defRPr", sz="2000" if t == 1 else "1400")
        sf = etree.SubElement(d, f"{A}solidFill")
        etree.SubElement(sf, f"{A}schemeClr", val="lt1")
        etree.SubElement(d, f"{A}latin", typeface=face)
    lay = {l.name: l for l in master.slide_layouts}
    accent = "E8590C"

    def coloured(slide):
        el = slide._element.find(f"{P}cSld")
        el.insert(0, etree.fromstring(f'<p:bg {NS}><p:bgPr><a:solidFill><a:srgbClr val="{accent}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'))

    s = prs.slides.add_slide(lay["Title Slide"])
    coloured(s)
    s.shapes.title.text_frame.text = "Annual review"
    for i in range(n_content):
        s = prs.slides.add_slide(lay["Title and Content"])
        s.shapes.title.text_frame.text = f"Revenue grew in region {i + 1} as volume rose"
        body = next(ph for ph in s.placeholders if ph.placeholder_format.idx != 0)
        body.text_frame.text = "Volumes rose in every channel and prices held, so revenue grew faster than the market."
    s = prs.slides.add_slide(lay["Section Header"])
    coloured(s)
    s.shapes.title.text_frame.text = "Thank you"
    prs.save(str(path))
    return str(path)
