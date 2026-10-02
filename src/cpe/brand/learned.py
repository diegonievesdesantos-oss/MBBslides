"""Layouts learned from slide geometry (v1.9, DEBT_V18 U7).

Many real decks are built with free text boxes: their slide layouts carry no title or body
placeholder, so the layout classifier finds nothing to match and every slide falls back to the
engine's own layout (only colours and fonts come through). The slides themselves still show the
house layout: where the action title sits and how big it is, where the subtitle of a cover goes,
which shapes repeat on every page.

`learn_layouts` reads that geometry and writes a copy of the deck with one real layout per slide
family (cover, statement, content), each a clone of the layout those slides use plus:
  - a title placeholder at the median title box, with its size, weight, colour and font;
  - a subtitle placeholder on covers (the second text box);
  - a body placeholder over the area the slides fill below the title;
  - the shapes that repeat on most slides of the family (a bar, a logo, a constant label).
The example slides are re-linked to their learned layout, so the usage evidence counts too.
Everything downstream (classification, native / adaptive matching, fidelity) is unchanged.
"""
from __future__ import annotations

import copy
import re
import statistics
from pathlib import Path

from lxml import etree

EMU = 914400
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
DATE_RE = re.compile(r"\b\d{1,2}[/.-]\d{1,2}[/.-](?:19|20)?\d{2}\b|\b(?:19|20)\d{2}\b")
R_ATTRS = (f"{{{NS['r']}}}embed", f"{{{NS['r']}}}id", f"{{{NS['r']}}}link")


def _has_title_ph(shapes) -> bool:
    for sh in shapes:
        if getattr(sh, "is_placeholder", False):
            t = str(sh.placeholder_format.type)
            if "TITLE" in t and "SUB" not in t:
                return True
    return False


def needs_learning(prs) -> bool:
    """True when the example slides set their titles in free text boxes: no slide with text in a
    title placeholder, and none of the layouts they use has one."""
    slides = list(prs.slides)
    if len(slides) < 2:
        return False
    used = {s.slide_layout.part.partname: s.slide_layout for s in slides}
    if any(_has_title_ph(lay.placeholders) for lay in used.values()):
        return False
    return not any(_has_title_ph(s.placeholders) and any(p.has_text_frame and p.text_frame.text.strip() for p in s.placeholders) for s in slides)


def _texts(slide) -> list[dict]:
    out = []
    for sh in slide.shapes:
        if getattr(sh, "is_placeholder", False) or not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        runs = [r for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()]
        sizes = [r.font.size.pt for r in runs if r.font.size]
        r0 = runs[0] if runs else None
        color = None
        try:
            if r0 is not None and r0.font.color and r0.font.color.type is not None and r0.font.color.rgb is not None:
                color = str(r0.font.color.rgb)
        except (AttributeError, TypeError):
            color = None
        out.append({"shape": sh, "text": " ".join(sh.text_frame.text.split()), "size": max(sizes, default=0.0),
                    "bold": bool(r0 is not None and r0.font.bold), "color": color, "font": r0.font.name if r0 is not None else None,
                    "x": sh.left or 0, "y": sh.top or 0, "w": sh.width or 0, "h": sh.height or 0})
    return out


def _title_of(texts: list[dict]) -> dict | None:
    """The action title: the top-most sentence in a large size, else the largest text (as deck ingest)."""
    sentences = [t for t in texts if len(t["text"].split()) >= 3 and t["size"] >= 16]
    if sentences:
        return min(sentences, key=lambda t: t["y"])
    return max(texts, key=lambda t: (t["size"], -t["y"])) if texts else None


def _median_box(items: list[dict]) -> dict:
    return {k: int(statistics.median(i[k] for i in items)) for k in ("x", "y", "w", "h")}


def _style(items: list[dict]) -> dict:
    sizes = [i["size"] for i in items if i["size"]]
    colors = [i["color"] for i in items if i["color"]]
    fonts = [i["font"] for i in items if i["font"]]
    return {"size": statistics.median(sizes) if sizes else None, "bold": sum(i["bold"] for i in items) * 2 > len(items),
            "color": max(set(colors), key=colors.count) if colors else None, "font": max(set(fonts), key=fonts.count) if fonts else None}


def _ph_xml(sid: int, name: str, ph: str, box: dict, style: dict | None, idx: int | None = None, prompt: str = "", anchor: str = "t") -> etree._Element:
    st = style or {}
    rpr = []
    if st.get("size"):
        rpr.append(f'sz="{int(round(st["size"] * 100))}"')
    if st.get("bold"):
        rpr.append('b="1"')
    fill = f'<a:solidFill><a:srgbClr val="{st["color"]}"/></a:solidFill>' if st.get("color") else ""
    latin = f'<a:latin typeface="{st["font"]}"/>' if st.get("font") else ""
    lst = f'<a:lstStyle><a:lvl1pPr><a:defRPr {" ".join(rpr)}>{fill}{latin}</a:defRPr></a:lvl1pPr></a:lstStyle>' if (rpr or fill or latin) else "<a:lstStyle/>"
    idx_attr = f' idx="{idx}"' if idx is not None else ""
    type_attr = f' type="{ph}"' if ph != "body" else ""
    xml = (f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}"><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
           f'<p:nvPr><p:ph{type_attr}{idx_attr}/></p:nvPr></p:nvSpPr>'
           f'<p:spPr><a:xfrm><a:off x="{box["x"]}" y="{box["y"]}"/><a:ext cx="{box["w"]}" cy="{box["h"]}"/></a:xfrm></p:spPr>'
           f'<p:txBody><a:bodyPr anchor="{anchor}"/>{lst}<a:p><a:r><a:rPr lang="en-US"/><a:t>{prompt}</a:t></a:r></a:p></p:txBody></p:sp>')
    return etree.fromstring(xml)


def _repeated(group: list, title_shapes: set) -> list:
    """Shapes (not text with changing content, not data) at the same place on ≥ 60% of the family's slides."""
    if len(group) < 3:
        return []
    seen: dict = {}
    for slide in group:
        keys = set()
        for sh in slide.shapes:
            if getattr(sh, "is_placeholder", False) or id(sh) in title_shapes or getattr(sh, "has_chart", False) or getattr(sh, "has_table", False):
                continue
            txt = " ".join(sh.text_frame.text.split()) if sh.has_text_frame else ""
            key = (round((sh.left or 0) / EMU, 1), round((sh.top or 0) / EMU, 1), round((sh.width or 0) / EMU, 1), round((sh.height or 0) / EMU, 1), txt, sh.shape_type)
            if key not in keys:
                keys.add(key)
                seen.setdefault(key, (0, slide, sh))
                seen[key] = (seen[key][0] + 1, seen[key][1], seen[key][2])
    return [(sl, sh) for n, sl, sh in seen.values() if n >= 0.6 * len(group)]


def _next_partname(package, tmpl: str):
    from pptx.opc.packuri import PackURI

    names = {str(p.partname) for p in package.iter_parts()}
    i = 1
    while tmpl % i in names:
        i += 1
    return PackURI(tmpl % i)


def _clone_layout(prs, src_layout, name: str, ooxml_type: str):
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.parts.slide import SlideLayoutPart

    src = src_layout.part
    element = copy.deepcopy(src._element)
    part = SlideLayoutPart(_next_partname(src.package, "/ppt/slideLayouts/slideLayout%d.xml"), src.content_type, src.package, element)
    remap = {}
    for rid, rel in src.rels.items():
        if rel.is_external:
            remap[rid] = part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            remap[rid] = part.relate_to(rel.target_part, rel.reltype)
    for el in element.iter():
        for at in R_ATTRS:
            if el.get(at) in remap:
                el.set(at, remap[el.get(at)])
    cSld = element.find(f"{{{NS['p']}}}cSld")
    cSld.set("name", name)
    element.set("type", ooxml_type)
    element.set("preserve", "1")
    master = src_layout.slide_master
    rid = master.part.relate_to(part, RT.SLIDE_LAYOUT)
    lst = master._element.get_or_add_sldLayoutIdLst()
    used = [int(x.get("id")) for x in prs.part._element.iter() if x.tag in (f"{{{NS['p']}}}sldMasterId",)]
    used += [int(x.get("id")) for m in prs.slide_masters for x in m._element.iter(f"{{{NS['p']}}}sldLayoutId")]
    new = etree.SubElement(lst, f"{{{NS['p']}}}sldLayoutId")
    new.set("id", str(max(used + [2147483648]) + 1))
    new.set(f"{{{NS['r']}}}id", rid)
    return part


def _copy_shape(src_slide, sh, part, tree) -> None:
    """A slide shape into a layout's shape tree, its picture / link relationships carried over."""
    el = copy.deepcopy(sh._element)
    for e in el.iter():
        for at in R_ATTRS:
            rid = e.get(at)
            if rid and rid in src_slide.part.rels:
                rel = src_slide.part.rels[rid]
                e.set(at, part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref) if rel.is_external else part.relate_to(rel.target_part, rel.reltype))
    tree.append(el)


def learn_layouts(src: str | Path, dst: str | Path) -> dict:
    """Write `dst`: the deck at `src` with learned layouts added and its slides re-linked to them.
    Returns what was learned (for the compatibility report)."""
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    prs = Presentation(str(src))
    slides = list(prs.slides)
    families: dict = {}
    cover_layout = slides[0].slide_layout.part.partname if slides else None  # a later slide on the cover's layout, title low: a statement
    for i, s in enumerate(slides):
        lay = s.slide_layout
        t = _title_of(_texts(s))
        top = t is not None and t["y"] < prs.slide_height * 0.12  # an action title at the very top: a content page
        if i == 0:
            fam = "cover"
        elif lay.part.partname == cover_layout and not top:
            fam = "statement"
        else:
            fam = f"content|{lay.part.partname}"
        families.setdefault(fam, []).append(s)
    learned = []
    content_n = sum(1 for f in families if f.startswith("content|"))
    for fam, group in families.items():
        info = [(s, _texts(s)) for s in group]
        titles = [(s, _title_of(t), t) for s, t in info]
        titles = [(s, t, ts) for s, t, ts in titles if t]
        if not titles:
            continue
        lay = group[0].slide_layout
        kind = fam.split("|")[0]
        label = {"cover": "Learned · Cover", "statement": "Learned · Statement", "content": "Learned · Title and content"}[kind]
        if kind == "content" and content_n > 1:
            label += f" ({lay.name})"
        ooxml = {"cover": "title", "statement": "titleOnly", "content": "obj"}[kind]
        part = _clone_layout(prs, lay, label, ooxml)
        tree = part._element.find(f"{{{NS['p']}}}cSld/{{{NS['p']}}}spTree")
        ids = [int(e.get("id")) for e in tree.iter(f"{{{NS['p']}}}cNvPr") if (e.get("id") or "").isdigit()]
        sid = max(ids + [1]) + 1
        tbox, tstyle = _median_box([t for _, t, _ in titles]), _style([t for _, t, _ in titles])
        if kind == "cover":  # room for a two-line title: the box grows upwards, the text stays on the learned baseline
            top = max(int(prs.slide_height * 0.08), tbox["y"] - tbox["h"])
            tbox = {**tbox, "y": top, "h": tbox["h"] + (tbox["y"] - top)}
        tree.append(_ph_xml(sid, "Title 1", "ctrTitle" if kind == "cover" else "title", tbox, tstyle, prompt="Title", anchor="b" if kind == "cover" else "t"))
        entry = {"layout": label, "from_layout": lay.name, "slides": [slides.index(s) + 1 for s in group], "kind": kind,
                 "title": {k: round(v / EMU, 2) for k, v in tbox.items()} | {"size": tstyle["size"], "bold": tstyle["bold"]}}
        if kind == "cover":
            subs = [min(below_s, key=lambda x: x["y"]) for _, t, ts in titles if (below_s := [x for x in ts if x is not t and x["y"] >= t["y"]])]
            if subs:
                sbox = _median_box(subs)
                tree.append(_ph_xml(sid + 1, "Subtitle 2", "subTitle", sbox, _style(subs), idx=1, prompt="Subtitle"))
                entry["subtitle"] = {k: round(v / EMU, 2) for k, v in sbox.items()}
        elif kind == "content":
            body_items = [x for s, t, ts in titles for x in ts if x is not t and x["y"] > t["y"] + t["h"] * 0.5 and x["size"] > 10]
            others = [{"x": sh.left or 0, "y": sh.top or 0, "w": sh.width or 0, "h": sh.height or 0} for s, t, _ in titles for sh in s.shapes
                      if not getattr(sh, "is_placeholder", False) and (getattr(sh, "has_chart", False) or getattr(sh, "has_table", False))]
            area = body_items + others
            if area:
                x0, y0 = min(a["x"] for a in area), max(tbox["y"] + tbox["h"], min(a["y"] for a in area))
                x1, y1 = max(a["x"] + a["w"] for a in area), max(a["y"] + a["h"] for a in area)
                prs_h = prs.slide_height
                y1 = min(y1, int(prs_h * 0.88))
                if y1 - y0 > prs_h * 0.25:
                    bbox = {"x": x0, "y": y0, "w": x1 - x0, "h": y1 - y0}
                    tree.append(_ph_xml(sid + 1, "Content Placeholder 2", "body", bbox, None, idx=1, prompt="Text"))
                    entry["body"] = {k: round(v / EMU, 2) for k, v in bbox.items()}
        title_ids = {id(t["shape"]) for _, t, _ in titles}
        rep = _repeated(group, title_ids)
        for s, sh in rep:
            _copy_shape(s, sh, part, tree)
        entry["repeated_shapes"] = len(rep)
        for s in group:  # the example slides now use the learned layout (usage evidence)
            for rel in s.part.rels.values():
                if rel.reltype == RT.SLIDE_LAYOUT:
                    rel._target = part
        learned.append(entry)
    prs.save(str(dst))
    dated = sorted({" ".join(sh.text_frame.text.split()) for m in prs.slide_masters for holder in [m, *m.slide_layouts] for sh in holder.shapes
                    if not getattr(sh, "is_placeholder", False) and sh.has_text_frame and DATE_RE.search(sh.text_frame.text)})
    return {"learned_layouts": learned, "dated_layout_text": dated}
