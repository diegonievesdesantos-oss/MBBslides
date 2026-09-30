import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

from conftest import ROOT, content_slide, mini_spec
from cpe.brand.ingest import ingest
from cpe.core.planner import plan
from cpe.design.tokens import GRID, load_theme
from cpe.pptx.builder import build


def _template(tmp_path) -> Path:
    out = tmp_path / "tpl.pptx"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "make_sample_template.py"), str(out)], check=True, capture_output=True)
    return out


def test_ingest_reports_fonts_colours_layouts_and_reserved_areas(tmp_path):
    rep = ingest(_template(tmp_path), tmp_path / "brand", name="Kestrel")
    assert rep["slide_size"]["supported"] and rep["masters_used"]
    assert rep["fonts"]["heading"]["font"] == "Georgia" and rep["fonts"]["body"]["font"] == "Calibri"
    assert rep["fonts"]["body"]["measurement"] == "exact"  # Calibri → Carlito (metric-compatible)
    assert rep["fonts"]["heading"]["measurement"] == "approximate" or rep["fonts"]["heading"]["installed"]
    assert rep["color_mapping"]["primary"] == "1B3A2F" and rep["color_mapping"]["highlight"] == "C8A24A"
    assert rep["base_layout"] == "Blank"
    assert any(a["name"] == "Kestrel logo" for a in rep["reserved_areas"])
    assert any("gradient" in u for u in rep["unsupported"])
    assert rep["grid_overrides"]["margin_l"] == 0.7
    assert (tmp_path / "brand" / "compatibility.md").exists()


def test_deck_is_built_on_the_template_masters(tmp_path):
    ingest(_template(tmp_path), tmp_path / "brand", name="Kestrel")
    spec = mini_spec([content_slide()])
    spec["meta"]["brand"] = str(tmp_path / "brand")
    res, _ = plan(spec)
    out = tmp_path / "d.pptx"
    man = build(res, out)
    prs = Presentation(str(out))
    assert len(prs.slides) == 1  # the template's sample slide was dropped
    s = prs.slides[0]
    # a content slide is carried by the template's own content layout (adaptive corporate composition)
    assert man[0]["corporate"]["mode"] == "adaptive" and s.slide_layout.name == man[0]["corporate"]["layout"] == "Title Only"
    assert not list(s.placeholders)  # no empty "click to add" boxes
    fonts = {r.font.name for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
    assert {"Georgia", "Calibri"} <= fonts
    assert abs(GRID.margin_l - 0.7) < 1e-6
    theme = load_theme(str(tmp_path / "brand"))
    assert json.loads((tmp_path / "brand" / "theme.json").read_text())["extras"]["template"] == "template.pptx"
    assert theme.font_for("headline") == "Georgia" and theme.font_for("body") == "Calibri"
    # restore the default grid for the other tests
    load_theme("meridian")
    from cpe.design.tokens import activate

    activate(load_theme("meridian"))
