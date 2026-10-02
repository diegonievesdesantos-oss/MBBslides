"""v3.0 validation kit: an owner's blind key for a real update case, sealed before the tool runs.

    python scripts/validation_kit.py template OLD.pptx -o clave.xlsx      # one row per number of the old deck
    python scripts/validation_kit.py seal clave.xlsx [...] -o SEALED      # sha256 of each filled key, not opened
    python scripts/validation_kit.py score clave.xlsx WORK [-v]           # the tool's plan against the key

The template lists every number `cpe update` reads in the old deck (slide, number as written, context,
where). The owner fills three columns without seeing the tool's output:
    estado      desactualizada | vigente | sin fuente | ignorar
    valor nuevo the new value in the deck's own units and scale (only when desactualizada)
    de dónde    optional: the file that gives it, or "calculado"
The key is sealed (sha256) before the tool runs on the case, and the tool runs once.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

COLS = ["id", "diapositiva", "cifra", "contexto", "dónde", "estado", "valor nuevo", "de dónde", "nota"]
STATES = ("desactualizada", "vigente", "sin fuente", "ignorar")
LABEL = {"desactualizada": "outdated", "vigente": "valid", "sin fuente": "unsourced", "ignorar": "ignore",
         "outdated": "outdated", "valid": "valid", "unsourced": "unsourced", "ignore": "ignore"}


def _where(w: str) -> str:
    if w == "title":
        return "título"
    if ".rows[" in w:
        return "tabla"
    if w.startswith("exhibit["):
        return "gráfico"
    return "texto"


def template(pptx: str, out: str) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    from cpe.reasoning.deck_update import ingest_deck

    inv = ingest_deck(pptx)
    wb = Workbook()
    ws = wb.active
    ws.title = "Clave"
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True)
    n = 0
    for s in inv["slides"]:
        for q in s["numbers"]:
            n += 1
            ws.append([f"{s['n']}|{q['where']}|{q.get('occ', 0)}|{q['raw']}", s["n"], q["raw"], (q.get("context") or "")[:90], _where(q["where"]),
                       "", "", "", ""])
    fill = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=9):
        for c in row:
            c.fill = fill
    dv = DataValidation(type="list", formula1='"' + ",".join(STATES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F2:F{n + 1}")
    for col, w in zip("ABCDEFGHI", (8, 11, 14, 60, 10, 16, 14, 24, 30)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["A"].hidden = True
    ws.freeze_panes = "C2"
    help_ = wb.create_sheet("Instrucciones")
    for line in (
        "Una fila por cifra del deck anterior. Rellena solo las columnas amarillas, sin mirar lo que proponga la herramienta.",
        "estado: desactualizada (las fuentes nuevas implican otro valor) · vigente (una fuente confirma el mismo valor) ·",
        "        sin fuente (nada en las fuentes nuevas dice cómo está hoy) · ignorar (no es una cifra de negocio).",
        "valor nuevo: solo si está desactualizada, en las mismas unidades y escala del deck ('3,2 M€' → 3,52; celda '3.200' en k€ → 3.520).",
        "Un total, ratio o payback recalculado con datos nuevos es 'desactualizada' con su valor recalculado.",
        "Una oferta o presupuesto antiguo que repite la cifra no la confirma: 'sin fuente', salvo que otra fuente la confirme.",
        "de dónde (opcional): el archivo que da el valor, o 'calculado'. nota: cualquier duda.",
    ):
        help_.append([line])
    help_.column_dimensions["A"].width = 130
    wb.save(out)
    return n


def seal(keys: list[str], out: str) -> None:
    lines = [f"{hashlib.sha256(Path(k).read_bytes()).hexdigest()}  {Path(k).name}" for k in keys]
    Path(out).write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))


def _num(s) -> float | None:
    if s is None or s == "":
        return None
    if isinstance(s, (int, float)):
        return float(s)
    t = str(s).strip().replace("−", "-").replace(" ", "")
    t = re.sub(r"[^\d,.\-]", "", t)
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".") if re.search(r",\d{1,2}$", t) or t.count(",") == 1 and not re.search(r",\d{3}$", t) else t.replace(",", "")
    elif t.count(".") == 1 and re.search(r"\.\d{3}$", t) and not t.startswith("0."):
        t = t.replace(".", "")  # 3.200 → 3200
    try:
        return float(t)
    except ValueError:
        return None


def read_key(path: str) -> list[dict]:
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)["Clave"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "").strip() for h in rows[0]]
    out = []
    for r in rows[1:]:
        d = dict(zip(head, r))
        lab = LABEL.get(str(d.get("estado") or "").strip().lower())
        if not d.get("id") or not lab:
            continue
        out.append({"id": d["id"], "label": lab, "new": _num(d.get("valor nuevo")), "note": d.get("nota") or ""})
    return out


def score(key_path: str, work: str, verbose: bool = False) -> dict:
    plan = json.loads((Path(work) / "update_plan.json").read_text(encoding="utf-8"))
    mark = plan.get("decimal_mark") or ","
    by_id = {q["id"]: q for s in plan["slides"] for q in s["numbers"] if q.get("id")}
    c = {k: 0 for k in ("keyed", "matched", "out_gold", "out_found", "out_flag", "out_flag_ok", "cur_flag", "cur_ok", "val_n", "val_ok")}
    bad = []
    for k in read_key(key_path):
        c["keyed"] += 1
        q = by_id.get(k["id"])
        if q is None:
            bad.append(("NOT IN PLAN", k["id"]))
            continue
        c["matched"] += 1
        lab, st = k["label"], q["status"]
        c["out_gold"] += lab == "outdated"
        c["out_found"] += lab == "outdated" and st == "outdated"
        if st == "outdated":
            c["out_flag"] += 1
            c["out_flag_ok"] += lab == "outdated"
            if lab == "outdated" and k["new"] is not None:
                c["val_n"] += 1
                chart = q["where"].startswith("exhibit[") and ".rows[" not in q["where"]
                shown = str(q.get("new_display") or "")
                if chart or mark == ".":  # chart values and point-decimal decks: "1.302" is one point three
                    try:
                        v = float(shown.replace(",", "."))
                    except ValueError:
                        v = _num(shown)
                else:
                    v = _num(shown)
                from cpe.reasoning.derive import _step

                tol = _step(q) * abs(k["new"]) / max(abs(q["value"]), 1e-12) + 1e-9 if q["value"] else 0.5
                ok = v is not None and abs(v - k["new"]) <= max(tol, 0.005 * abs(k["new"]))
                c["val_ok"] += ok
                if not ok:
                    bad.append(("VALUE", k["id"], q.get("new_display"), k["new"]))
        if st == "current":
            c["cur_flag"] += 1
            c["cur_ok"] += lab == "valid"
        if (lab == "outdated") != (st == "outdated") or (st == "current" and lab != "valid"):
            bad.append((lab, st, k["id"], k["note"]))
    print(f"keyed {c['keyed']} matched {c['matched']} | outdated found {c['out_found']}/{c['out_gold']} | 'outdated' right {c['out_flag_ok']}/{c['out_flag']}"
          f" | 'current' right {c['cur_ok']}/{c['cur_flag']} | value right {c['val_ok']}/{c['val_n']}")
    if verbose:
        for b in bad:
            print("  ", b)
    return c


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    t = sub.add_parser("template")
    t.add_argument("pptx")
    t.add_argument("-o", "--out", required=True)
    s = sub.add_parser("seal")
    s.add_argument("keys", nargs="+")
    s.add_argument("-o", "--out", required=True)
    r = sub.add_parser("score")
    r.add_argument("key")
    r.add_argument("work")
    r.add_argument("-v", action="store_true")
    a = ap.parse_args()
    if a.cmd == "template":
        print(f"{template(a.pptx, a.out)} numbers → {a.out}")
    elif a.cmd == "seal":
        seal(a.keys, a.out)
    else:
        score(a.key, a.work, a.v)


if __name__ == "__main__":
    main()
