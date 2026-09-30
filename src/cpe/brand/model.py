"""Corporate template intelligence: understand a PowerPoint's visual SYSTEM, not just its theme.

    analyse(prs) → brand model

Everything iterates over ALL slide masters (never `slide_masters[0]`):

    Presentation ─ Master 1 ─ theme, placeholders, artwork, layouts ─ placeholders, artwork
                 ─ Master 2 ─ …
                 ─ example slides (the strongest signal of how the system is really used)

The model combines evidence from several sources with explicit confidence and reports
contradictions instead of silently picking one:

  typography   theme (declared) + master text styles + layouts + actual runs on the example
               slides (direct formatting, weighted by characters) + explicit style-guide slides
  palette      theme colour scheme + area-weighted fills and character-weighted text colours
               actually used on layouts and slides
  grid         left/right edges of placeholders and content shapes → margins and the column
               system that best explains them
  assets       pictures and vector artwork on masters / layouts / slides: logos, icons,
               reserved artwork
  layouts      geometric features per layout → semantic classes with confidences, refined by
               how example slides actually use each layout
  brand rules  sentence case, bookend openings/closings, share of brand-colour slides,
               headline position, logo position, rounded shapes, chevrons, alignment

No rule here is specific to any company: every conclusion is derived from the file.
"""
from __future__ import annotations

import collections
import colorsys
import hashlib
import re
from dataclasses import dataclass, field

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
SCHEME_SLOTS = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]
SCHEME_ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
CLASSES = ["cover", "section", "content", "statement", "one_column", "two_column", "three_column", "image_split", "chart", "chart_commentary",
           "table", "comparison", "matrix", "process", "roadmap", "timeline", "conclusion", "closing", "special", "unknown"]


def _in(v) -> float:
    return round(Emu(v or 0).inches, 3)


# ── theme ───────────────────────────────────────────────────────────────────────

def theme_of_master(master) -> dict:
    part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    colors = {}
    cs = root.find(f".//{A}clrScheme")
    for slot in SCHEME_SLOTS:
        el = cs.find(f"{A}{slot}") if cs is not None else None
        if el is None:
            continue
        c = el.find(f"{A}srgbClr")
        if c is not None:
            colors[slot] = c.get("val").upper()
        else:
            s = el.find(f"{A}sysClr")
            if s is not None:
                colors[slot] = (s.get("lastClr") or ("000000" if s.get("val") == "windowText" else "FFFFFF")).upper()
    fonts = {}
    fs = root.find(f".//{A}fontScheme")
    if fs is not None:
        for kind in ("majorFont", "minorFont"):
            lat = fs.find(f"{A}{kind}/{A}latin")
            fonts[kind] = lat.get("typeface") if lat is not None else None
    return {"part": str(part.partname), "name": root.get("name"), "scheme_name": cs.get("name") if cs is not None else None, "colors": colors, "fonts": fonts}


def _resolve_font(face: str | None, theme: dict) -> str | None:
    if not face:
        return None
    if face.startswith("+mj"):
        return theme["fonts"].get("majorFont")
    if face.startswith("+mn"):
        return theme["fonts"].get("minorFont")
    return face


def _color_of(el, theme: dict) -> str | None:
    """First solid colour under el (srgbClr or schemeClr resolved through the theme)."""
    if el is None:
        return None
    c = el.find(f".//{A}srgbClr")
    s = el.find(f".//{A}schemeClr")
    if c is not None and (s is None or _depth(c) <= _depth(s)):
        return c.get("val").upper()
    if s is not None:
        return theme["colors"].get(SCHEME_ALIAS.get(s.get("val"), s.get("val")))
    return None


def _depth(el) -> int:
    d = 0
    while el is not None:
        el = el.getparent()
        d += 1
    return -d


# ── colour helpers ─────────────────────────────────────────────────────────────

def rgb(h: str) -> tuple[int, int, int]:
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def lum(h: str) -> float:
    r, g, b = rgb(h)
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255


def sat(h: str) -> float:
    r, g, b = [c / 255 for c in rgb(h)]
    return colorsys.rgb_to_hsv(r, g, b)[1]


def hue(h: str) -> float:
    r, g, b = [c / 255 for c in rgb(h)]
    return colorsys.rgb_to_hsv(r, g, b)[0] * 360


def is_neutral(h: str) -> bool:
    return sat(h) < 0.15


# ── shapes ─────────────────────────────────────────────────────────────────────

def shape_kind(sh) -> str:
    if sh.is_placeholder:
        return "placeholder"
    try:
        st = sh.shape_type
    except NotImplementedError:
        st = None
    if st == MSO_SHAPE_TYPE.PICTURE or sh._element.tag == f"{P}pic":
        return "picture"
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if getattr(sh, "has_chart", False) and sh.has_chart:
        return "chart"
    if getattr(sh, "has_table", False) and sh.has_table:
        return "table"
    if sh._element.tag.endswith("graphicFrame"):
        return "graphic_frame"
    if sh._element.tag.endswith("cxnSp"):
        return "connector"
    if sh.has_text_frame and sh.text_frame.text.strip():
        return "text"
    return "shape"


def ph_type(sh) -> str:
    t = str(sh.placeholder_format.type).split(".")[-1].split(" ")[0]
    return t


def box(sh, sw: float, shh: float) -> dict:
    x, y, w, h = _in(sh.left), _in(sh.top), _in(sh.width), _in(sh.height)
    return {"x": x, "y": y, "w": w, "h": h, "nx": round(x / sw, 4), "ny": round(y / shh, 4), "nw": round(w / sw, 4), "nh": round(h / shh, 4)}


def prst(sh) -> str | None:
    g = sh._element.find(f".//{A}prstGeom")
    return g.get("prst") if g is not None else None


def image_hash(sh) -> str | None:
    try:
        return hashlib.sha1(sh.image.blob).hexdigest()[:12]
    except Exception:
        blip = sh._element.find(f".//{A}blip")
        return blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed") if blip is not None else None


def _fill(sh, theme: dict) -> str | None:
    sp = sh._element.find(f"{P}spPr")
    if sp is None:
        sp = sh._element.find(f".//{P}spPr")
    if sp is None or sp.find(f"{A}noFill") is not None:
        return None
    sf = sp.find(f"{A}solidFill")
    if sf is not None:
        return _color_of(sf, theme)
    return None


def background(el, theme: dict) -> dict | None:
    """Background of a master / layout / slide element: kind and colour when solid."""
    bg = el.find(f"{P}cSld/{P}bg")
    if bg is None:
        return None
    if bg.find(f".//{A}blipFill") is not None:
        return {"kind": "picture"}
    if bg.find(f".//{A}gradFill") is not None:
        return {"kind": "gradient", "color": _color_of(bg.find(f".//{A}gradFill"), theme)}
    ref = bg.find(f"{P}bgRef")
    col = _color_of(bg, theme)
    if col is None and ref is not None:
        col = theme["colors"].get("lt1")
    return {"kind": "solid", "color": col}


def _full_bleed_fill(shapes, sw, shh, theme) -> str | None:
    """A non-placeholder shape covering ≥ 85% of the slide acts as the background."""
    for sh in shapes:
        try:
            if sh.is_placeholder:
                continue
            if _in(sh.width) >= 0.85 * sw and _in(sh.height) >= 0.85 * shh:
                if shape_kind(sh) == "picture":
                    return "picture"
                c = _fill(sh, theme)
                if c:
                    return c
        except Exception:
            continue
    return None


def slide_background(slide, th: dict, sw: float, shh: float) -> str | None:
    """Effective background colour of a slide: slide → layout → master, explicit bg or a full-bleed shape."""
    lay = slide.slide_layout
    for el, shapes in ((slide._element, slide.shapes), (lay._element, lay.shapes), (lay.slide_master._element, lay.slide_master.shapes)):
        bleed = _full_bleed_fill(shapes, sw, shh, th)
        if bleed:
            return bleed
        b = background(el, th)
        if b:
            return b.get("color") if b.get("kind") != "picture" else "picture"
    return th["colors"].get("lt1")


# ── layouts ─────────────────────────────────────────────────────────────────────

@dataclass
class LayoutInfo:
    master_id: str
    layout_id: str
    name: str
    ooxml_type: str | None
    placeholders: list = field(default_factory=list)
    artwork: list = field(default_factory=list)
    background: dict | None = None
    shows_master_artwork: bool = True
    features: dict = field(default_factory=dict)
    geometry_scores: dict = field(default_factory=dict)
    usage: dict = field(default_factory=dict)
    classification: list = field(default_factory=list)


CONTENT_PH = {"BODY", "OBJECT", "PICTURE", "CHART", "TABLE", "MEDIA_CLIP", "ORG_CHART", "SLIDE_IMAGE", "CLIP_ART", "BITMAP", "VERTICAL_BODY", "VERTICAL_OBJECT"}
CHROME_PH = {"DATE", "FOOTER", "SLIDE_NUMBER", "HEADER"}


def _x_clusters(boxes: list[dict], tol: float = 0.04) -> int:
    xs = sorted(b["nx"] for b in boxes)
    n = 0
    last = None
    for x in xs:
        if last is None or x - last > tol:
            n += 1
        last = x
    return n


def layout_features(li: LayoutInfo, sw: float, shh: float, master_bg: dict | None, master_art: list, theme: dict) -> dict:
    phs = li.placeholders
    title = next((p for p in phs if p["type"] in ("TITLE", "CENTER_TITLE")), None)
    sub = next((p for p in phs if p["type"] == "SUBTITLE"), None)
    content = [p for p in phs if p["type"] in CONTENT_PH]
    bodies = [p for p in content if p["type"] in ("BODY", "OBJECT", "VERTICAL_BODY", "VERTICAL_OBJECT")]
    pics = [p for p in content if p["type"] in ("PICTURE", "BITMAP", "CLIP_ART")]
    big = [p for p in content if p["nw"] * p["nh"] >= 0.04]
    small_heads = [p for p in bodies if p["nh"] < 0.1]  # column headers
    big_bodies = [p for p in bodies if p["nh"] >= 0.1]
    art = li.artwork + (master_art if li.shows_master_artwork else [])
    bg = li.background or master_bg or {}
    bg_color = bg.get("color")
    bleed = li.features.get("_bleed")
    if bleed and bleed != "picture":
        bg_color = bleed
    dark = bg_color is not None and lum(bg_color) < 0.45
    coloured = bg_color is not None and not is_neutral(bg_color) and lum(bg_color) >= 0.25
    title_top = title is not None and title["ny"] < 0.22
    title_mid = title is not None and 0.22 <= title["ny"] + title["nh"] / 2 <= 0.78 and title["ny"] >= 0.2
    free_top = (title["ny"] + title["nh"]) if title else 0.12
    covered = sum(p["nw"] * p["nh"] for p in big)
    art_in_body = [a for a in art if a["ny"] > free_top and a["ny"] + a["nh"] < 0.92 and a["nw"] * a["nh"] > 0.01 and a["nw"] < 0.85]
    cols = _x_clusters(big_bodies + pics) if (big_bodies or pics) else 0
    widths = [p["nw"] for p in big_bodies + pics]
    rows_ = len({round(p["ny"], 1) for p in big_bodies + pics})
    symmetric = bool(widths) and max(widths) - min(widths) < 0.05
    f = {
        "master_id": li.master_id, "layout_id": li.layout_id, "layout_name": li.name, "ooxml_type": li.ooxml_type,
        "placeholder_count": len(phs), "placeholder_types": sorted({p["type"] for p in phs}),
        "title_geometry": {k: title[k] for k in ("x", "y", "w", "h", "nx", "ny", "nw", "nh")} if title else None,
        "title_position": ("top" if title_top else "middle" if title_mid else "bottom") if title else None,
        "subtitle": sub is not None,
        "body_geometry": _union(content) if content else None,
        "image_placeholders": len(pics), "chart_placeholders": sum(1 for p in content if p["type"] == "CHART"),
        "table_placeholders": sum(1 for p in content if p["type"] == "TABLE"),
        "table_compatible_regions": sum(1 for p in big_bodies if p["nw"] >= 0.5 and p["nh"] >= 0.35),
        "content_placeholders": len(content), "column_headers": len(small_heads),
        "columns": cols, "rows": rows_, "symmetric": symmetric,
        "background": {"kind": bg.get("kind", "master" if li.background is None else None), "color": bg_color, "dark": dark, "coloured": coloured,
                       "picture": bg.get("kind") == "picture" or bleed == "picture"},
        "content_capacity": round(max(0.0, (0.92 - free_top)) * 0.9 - sum(a["nw"] * a["nh"] for a in art_in_body), 3),
        "content_coverage": round(covered, 3),
        "reserved_artwork": [{"name": a["name"], "kind": a["kind"], **{k: a[k] for k in ("x", "y", "w", "h")}} for a in art if a["nw"] * a["nh"] < 0.5],
        "artwork_in_body": len(art_in_body),
        "footer": any(p["type"] in CHROME_PH for p in phs),
        "logo": any(a.get("logo") for a in art),
        "orientation": "horizontal" if cols >= 2 and rows_ <= 1 else ("vertical" if rows_ >= 2 and cols <= 1 else ("grid" if cols >= 2 and rows_ >= 2 else "single")),
        "name_tokens": sorted(set(re.findall(r"[a-záéíóúñ]+", li.name.lower()))),
    }
    return f


def _union(ps: list[dict]) -> dict:
    x0 = min(p["nx"] for p in ps)
    y0 = min(p["ny"] for p in ps)
    x1 = max(p["nx"] + p["nw"] for p in ps)
    y1 = max(p["ny"] + p["nh"] for p in ps)
    return {"nx": round(x0, 4), "ny": round(y0, 4), "nw": round(x1 - x0, 4), "nh": round(y1 - y0, 4)}


NAME_HINTS = {
    "cover": ("cover", "title slide", "portada", "front", "opening", "titulo", "título"),
    "section": ("section", "divider", "separator", "chapter", "divisor", "sección", "seccion", "header"),
    "closing": ("closing", "end", "thank", "cierre", "final", "back cover", "contact", "gracias"),
    "statement": ("statement", "quote", "big", "impact", "cita", "mensaje", "number", "kpi"),
    "two_column": ("two", "2 col", "2col", "dos"),
    "three_column": ("three", "3 col", "3col", "tres"),
    "comparison": ("comparison", "compare", "versus", "vs", "comparación"),
    "image_split": ("picture", "image", "photo", "imagen", "foto", "split"),
    "chart": ("chart", "graph", "gráfico", "grafico"),
    "table": ("table", "tabla"),
    "timeline": ("timeline", "cronología", "milestone", "hitos"),
    "roadmap": ("roadmap", "phases", "fases", "plan"),
    "process": ("process", "steps", "proceso", "pasos", "chevron"),
    "matrix": ("matrix", "2x2", "matriz", "quadrant"),
    "content": ("title only", "content", "title and content", "blank", "contenido", "solo título", "solo titulo"),
    "conclusion": ("conclusion", "summary", "takeaway", "conclusión", "resumen", "agenda", "index", "índice"),
}
OOXML_HINTS = {"title": "cover", "secHead": "section", "titleOnly": "content", "obj": "one_column", "tx": "one_column", "twoObj": "two_column",
               "twoTxTwoObj": "comparison", "picTx": "image_split", "chart": "chart", "tbl": "table", "chartAndTx": "chart_commentary", "txAndChart": "chart_commentary",
               "blank": "content", "objTx": "chart_commentary", "txAndObj": "chart_commentary", "fourObj": "matrix", "objOnly": "one_column"}


def geometry_scores(f: dict) -> dict:
    """Evidence from geometry, placeholders, background and names → score per class (0–1)."""
    s = collections.defaultdict(float)
    tpos, n_content, cols = f["title_position"], f["content_placeholders"], f["columns"]
    bg = f["background"]
    if f["ooxml_type"] in OOXML_HINTS and f["ooxml_type"] != "cust":
        s[OOXML_HINTS[f["ooxml_type"]]] += 0.45
    name = f["layout_name"].lower()
    for cls, hints in NAME_HINTS.items():
        if any(re.search(rf"(^|[^a-z]){re.escape(h)}([^a-z]|$)", name) for h in hints):
            s[cls] += 0.35
    ctr = "CENTER_TITLE" in f["placeholder_types"]
    if ctr:
        s["cover"] += 0.35
    if tpos in ("middle", "bottom") and n_content == 0:
        s["section"] += 0.3
        s["cover"] += 0.2 if f["subtitle"] else 0.1
        s["statement"] += 0.15
        s["closing"] += 0.1
    if (bg["coloured"] or bg["dark"] or bg["picture"]) and n_content == 0:
        s["cover"] += 0.15
        s["section"] += 0.15
        s["closing"] += 0.1
    if tpos == "top":
        if n_content == 0:
            s["content"] += 0.55
        elif n_content == 1 and f["image_placeholders"] == 0:
            ph = f["body_geometry"]
            if f["chart_placeholders"]:
                s["chart"] += 0.6
            elif f["table_placeholders"]:
                s["table"] += 0.6
            elif ph and ph["nw"] > 0.6:
                s["one_column"] += 0.5
                s["content"] += 0.2
            else:
                s["one_column"] += 0.3
                s["special"] += 0.1
        if f["chart_placeholders"] and n_content >= 2:
            s["chart_commentary"] += 0.55
        if f["image_placeholders"] and n_content - f["image_placeholders"] >= 1 and f["image_placeholders"] <= 2:
            s["image_split"] += 0.5
        if f["image_placeholders"] >= 3:
            s["special"] += 0.4
        big = n_content - f["column_headers"]
        if cols == 2 and f["image_placeholders"] == 0:
            s["two_column"] += 0.5
            s["comparison"] += 0.25 + (0.25 if f["column_headers"] >= 2 else 0) + (0.1 if f["symmetric"] else 0)
        if cols == 3:
            s["three_column"] += 0.5
            s["comparison"] += 0.15
        if cols >= 4 and f["rows"] <= 1:
            s["process"] += 0.4
            s["timeline"] += 0.25
            s["roadmap"] += 0.2
        if f["orientation"] == "grid" and big == 4:
            s["matrix"] += 0.5
    if tpos is None:
        if n_content == 0:
            s["special"] += 0.3
            if bg["coloured"] or bg["dark"] or bg["picture"]:
                s["cover"] += 0.1
                s["closing"] += 0.1
        else:
            s["special"] += 0.35
    if bg["picture"] and n_content == 0:
        s["cover"] += 0.1
    return {k: round(min(1.0, v), 3) for k, v in s.items() if v > 0}


# ── example slides: how the system is really used ──────────────────────────────

def slide_observation(slide, sw, shh, theme, index: int, total: int) -> dict:
    """Observed archetype of an example slide from its actual content."""
    bgc = slide_background(slide, theme, sw, shh)
    coloured_bg = bgc == "picture" or (bgc is not None and (lum(bgc) < 0.45 or (not is_neutral(bgc) and lum(bgc) >= 0.25)))
    kinds = collections.Counter()
    text_boxes, run_sizes, shapes = [], [], []
    prsts = collections.Counter()
    for sh in slide.shapes:
        k = shape_kind(sh)
        if k == "placeholder":
            try:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    k = "text"
                elif getattr(sh, "has_chart", False) and sh.has_chart:
                    k = "chart"
                elif getattr(sh, "has_table", False) and sh.has_table:
                    k = "table"
                elif sh._element.tag == f"{P}pic":
                    k = "picture"
                else:
                    continue
            except Exception:
                continue
        kinds[k] += 1
        b = box(sh, sw, shh)
        shapes.append({**b, "kind": k})
        pg = prst(sh)
        if pg:
            prsts[pg] += 1
        if k == "text":
            text_boxes.append(b)
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.text.strip():
                        run_sizes.append(r.font.size.pt)
    big_number = any(sz >= 40 for sz in run_sizes)
    in_row = [s for s in shapes if s["kind"] in ("shape", "text") and 0.08 < s["nw"] < 0.3 and s["nh"] < 0.5]
    row_groups = collections.Counter(round(s["ny"], 1) for s in in_row)
    seq = max(row_groups.values()) if row_groups else 0
    chevrons = prsts.get("chevron", 0) + prsts.get("homePlate", 0)
    text_cols = _x_clusters([b for b in text_boxes if b["nh"] > 0.12 and b["ny"] > 0.18]) if text_boxes else 0
    arch = "text"
    if index == 0:
        arch = "cover"
    elif kinds["chart"]:
        arch = "chart_commentary" if kinds["text"] >= 3 else "chart"
    elif kinds["table"]:
        arch = "table"
    elif chevrons >= 3 or (seq >= 4 and prsts.get("rightArrow", 0) + chevrons >= 1):
        arch = "process"
    elif kinds["picture"] >= 3:
        arch = "special"
    elif kinds["picture"] >= 1 and kinds["text"] >= 1:
        arch = "image_split"
    elif big_number and len(text_boxes) <= 4:
        arch = "statement"
    elif seq >= 4:
        arch = "process"
    elif text_cols >= 3:
        arch = "three_column"
    elif text_cols == 2:
        arch = "two_column"
    elif len(text_boxes) <= 2 and coloured_bg and not kinds["picture"]:
        arch = "section"
    elif len(text_boxes) <= 2 and sum(len(sh.text_frame.text) for sh in slide.shapes if sh.has_text_frame) < 120:
        arch = "statement"
    if index == total - 1 and total > 3 and len(text_boxes) <= 3 and not kinds["chart"] and not kinds["table"]:
        arch = "closing"
    return {"archetype": arch, "kinds": dict(kinds), "prsts": dict(prsts), "text_boxes": len(text_boxes), "big_number": big_number, "background": bgc}


OBS_TO_CLASS = {"cover": "cover", "section": "section", "chart": "chart", "chart_commentary": "chart_commentary", "table": "table", "process": "process", "special": "special",
                "image_split": "image_split", "statement": "statement", "three_column": "three_column", "two_column": "two_column", "text": "one_column", "closing": "closing"}


def classify(geom: dict, usage: dict) -> list[dict]:
    """Combine geometry evidence and observed usage into ranked labels with confidences."""
    n = sum(usage.values())
    scores = collections.defaultdict(float)
    for k, v in geom.items():
        scores[k] += v * (0.6 if n else 1.0)
    if n:
        w = min(0.4, 0.15 + 0.05 * n)  # more example slides → more weight on usage
        for obs, c in usage.items():
            scores[OBS_TO_CLASS.get(obs, "unknown")] += w * c / n * 1.6
    if not scores:
        return [{"type": "unknown", "confidence": 0.2}]
    ranked = sorted(scores.items(), key=lambda kv: -kv[1])
    top = ranked[0][1]
    out = [{"type": k, "confidence": round(min(0.97, v), 2)} for k, v in ranked if v >= max(0.2, top * 0.6)][:3]
    return out


# ── the model ───────────────────────────────────────────────────────────────────

def analyse(prs, installed_fonts: set[str] | None = None) -> dict:
    sw, shh = _in(prs.slide_width), _in(prs.slide_height)
    masters, layouts, all_art = [], [], []
    layout_by_part: dict = {}
    themes = []
    for mi, master in enumerate(prs.slide_masters):
        th = theme_of_master(master)
        themes.append(th)
        mid = f"m{mi + 1}"
        mart, mph = [], []
        for sh in master.shapes:
            k = shape_kind(sh)
            b = box(sh, sw, shh)
            if k == "placeholder":
                mph.append({"type": ph_type(sh), "name": sh.name, **b})
            else:
                a = {"kind": k, "name": sh.name, "prst": prst(sh), "image": image_hash(sh) if k == "picture" else None, **b}
                mart.append(a)
        mbg = background(master._element, th)
        bleed = _full_bleed_fill(master.shapes, sw, shh, th)
        m = {"master_id": mid, "name": master.name, "theme": th, "background": mbg, "bleed": bleed, "placeholders": mph, "artwork": mart,
             "title_style": _txstyle(master, "titleStyle", th), "body_style": _txstyle(master, "bodyStyle", th), "layouts": []}
        for li_, lay in enumerate(master.slide_layouts):
            phs, art = [], []
            for sh in lay.shapes:
                k = shape_kind(sh)
                b = box(sh, sw, shh)
                if k == "placeholder":
                    phs.append({"type": ph_type(sh), "idx": sh.placeholder_format.idx, "name": sh.name, **b})
                else:
                    art.append({"kind": k, "name": sh.name, "prst": prst(sh), "image": image_hash(sh) if k == "picture" else None, **b})
            li = LayoutInfo(master_id=mid, layout_id=f"{mid}.l{li_ + 1}", name=lay.name, ooxml_type=lay._element.get("type"),
                            placeholders=phs, artwork=art, background=background(lay._element, th),
                            shows_master_artwork=lay._element.get("showMasterSp", "1") != "0")
            li.features["_bleed"] = _full_bleed_fill(lay.shapes, sw, shh, th)
            layout_by_part[lay.part.partname] = li
            m["layouts"].append(li.layout_id)
            layouts.append((li, lay, mart, mbg or ({"kind": "solid", "color": bleed} if bleed and bleed != "picture" else None), th))
            all_art += [("layout", li.layout_id, a) for a in art]
        all_art += [("master", mid, a) for a in mart]
        masters.append(m)
    # logos: the same image placed on several layouts / masters, small, near a corner (or named so)
    img_places = collections.defaultdict(set)
    for scope, owner, a in all_art:
        if a["kind"] == "picture" and a.get("image"):
            img_places[a["image"]].add(owner)
    for scope, owner, a in all_art:
        corner = (a["nx"] < 0.2 or a["nx"] + a["nw"] > 0.8) and (a["ny"] < 0.2 or a["ny"] + a["nh"] > 0.8)
        small = a["nw"] < 0.25 and a["nh"] < 0.25
        named = "logo" in a["name"].lower()
        a["logo"] = bool(a["kind"] in ("picture", "group", "shape") and small and corner and (named or (a.get("image") and len(img_places[a["image"]]) >= 2) or scope == "master" and a["kind"] == "picture"))
    # example slides
    slides_obs = []
    usage = collections.defaultdict(collections.Counter)
    total = len(prs.slides)
    for si, slide in enumerate(prs.slides):
        li = layout_by_part.get(slide.slide_layout.part.partname)
        th = themes[[m["master_id"] for m in masters].index(li.master_id)] if li else themes[0]
        obs = slide_observation(slide, sw, shh, th, si, total)
        obs["layout_id"] = li.layout_id if li else None
        slides_obs.append(obs)
        if li:
            usage[li.layout_id][obs["archetype"]] += 1
    catalog = []
    for li, lay, mart, mbg, th in layouts:
        li.features = layout_features(li, sw, shh, mbg, mart, th)
        li.geometry_scores = geometry_scores(li.features)
        li.usage = dict(usage.get(li.layout_id, {}))
        li.classification = classify(li.geometry_scores, li.usage)
        catalog.append({"master_id": li.master_id, "layout_id": li.layout_id, "layout": li.name, "classification": li.classification,
                        "usage": li.usage, "evidence": {"geometry": li.geometry_scores, "example_slides": sum(li.usage.values())},
                        "features": li.features})
    typography = infer_typography(prs, masters, themes, layouts, installed_fonts or set())
    theme_of_layout = {lay.part.partname: th for li, lay, _, _, th in layouts}
    palette = infer_palette(prs, masters, themes, sw, shh, theme_of_layout)
    grid = infer_grid(prs, catalog, sw, shh)
    assets = infer_assets(all_art, prs, sw, shh)
    rules = infer_rules(prs, catalog, slides_obs, palette, sw, shh, theme_of_layout)
    families = collections.Counter(c["classification"][0]["type"] for c in catalog)
    return {
        "slide_size": {"width_in": sw, "height_in": shh, "aspect": round(sw / shh, 4) if shh else None},
        "masters": [{k: v for k, v in m.items() if k != "layouts"} | {"layouts": len(m["layouts"])} for m in masters],
        "layouts": catalog, "layout_families": dict(families.most_common()),
        "example_slides": {"count": total, "observed_archetypes": dict(collections.Counter(o["archetype"] for o in slides_obs).most_common())},
        "typography": typography, "palette": palette, "grid": grid, "assets": assets, "rules": rules,
    }


def _txstyle(master, tag: str, th: dict) -> dict:
    el = master._element.find(f"{P}txStyles/{P}{tag}")
    if el is None:
        return {}
    lvl = el.find(f"{A}lvl1pPr")
    if lvl is None:
        return {}
    d = lvl.find(f"{A}defRPr")
    if d is None:
        return {}
    lat = d.find(f"{A}latin")
    return {"font": _resolve_font(lat.get("typeface") if lat is not None else None, th), "size_pt": int(d.get("sz")) / 100 if d.get("sz") else None,
            "bold": d.get("b") == "1", "color": _color_of(d, th), "caps": d.get("cap")}


# ── typography ───────────────────────────────────────────────────────────────────

GUIDE_WORDS = re.compile(r"\b(typeface|typography|font|fonts|tipograf[ií]a|fuente|schrift|police|typo)\b", re.I)


def infer_typography(prs, masters, themes, layouts, installed: set[str]) -> dict:
    ev = collections.defaultdict(lambda: collections.Counter())
    for th in themes:
        for role, key in (("heading", "majorFont"), ("body", "minorFont")):
            if th["fonts"].get(key):
                ev["theme"][th["fonts"][key]] += 1
    for m in masters:
        for st in ("title_style", "body_style"):
            if m[st].get("font"):
                ev["master"][m[st]["font"]] += 1
    for li, lay, _, _, th in layouts:
        for el in lay._element.iter(f"{A}latin"):
            f = _resolve_font(el.get("typeface"), th)
            if f:
                ev["layouts"][f] += 1
    chars_direct = collections.Counter()
    chars_inherited = collections.Counter()
    by_role = {"heading": collections.Counter(), "body": collections.Counter()}
    guide_mentions = collections.Counter()
    known = set()
    for th in themes:
        known |= {v for v in th["fonts"].values() if v}
    master_of_layout = {}
    for li, lay, _, _, th in layouts:
        master_of_layout[lay.part.partname] = th
    body_default = next((m["body_style"].get("font") for m in masters if m["body_style"].get("font")), None) or (themes[0]["fonts"].get("minorFont") if themes else None)
    head_default = next((m["title_style"].get("font") for m in masters if m["title_style"].get("font")), None) or (themes[0]["fonts"].get("majorFont") if themes else None)
    slide_texts = []
    for slide in prs.slides:
        th = master_of_layout.get(slide.slide_layout.part.partname) or themes[0]
        txt = []
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            is_title = sh.is_placeholder and ph_type(sh) in ("TITLE", "CENTER_TITLE")
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    t = r.text.strip()
                    if not t:
                        continue
                    txt.append(r.text)
                    lat = r._r.find(f"{A}rPr/{A}latin")
                    face = _resolve_font(lat.get("typeface") if lat is not None else None, th)
                    if face:
                        chars_direct[face] += len(t)
                        known.add(face)
                    else:
                        face = head_default if is_title else body_default
                        if face:
                            chars_inherited[face] += len(t)
                    if face:
                        by_role["heading" if is_title else "body"][face] += len(t)
        slide_texts.append(" ".join(txt))
    for t in slide_texts:
        if GUIDE_WORDS.search(t):
            for f in known | {x.title() for x in installed}:
                if f and len(f) > 2 and re.search(rf"\b{re.escape(f)}\b", t):
                    guide_mentions[f] += 1
    usage = chars_direct + chars_inherited

    def share(c: collections.Counter) -> dict:
        tot = sum(c.values())
        return {k: v / tot for k, v in c.items()} if tot else {}

    weights = {"theme": 0.15, "master": 0.10, "layouts": 0.10, "usage": 0.45, "guide": 0.20}
    comb = collections.Counter()
    srcs = {"theme": share(ev["theme"]), "master": share(ev["master"]), "layouts": share(ev["layouts"]), "usage": share(usage), "guide": share(guide_mentions)}
    active = {k: w for k, w in weights.items() if srcs[k]}
    norm = sum(active.values()) or 1
    for k, w in active.items():
        for f, v in srcs[k].items():
            comb[f] += w / norm * v
    ranked = comb.most_common()
    primary = ranked[0][0] if ranked else None
    declared = {"heading": [th["fonts"].get("majorFont") for th in themes], "body": [th["fonts"].get("minorFont") for th in themes]}
    declared_main = collections.Counter([f for v in declared.values() for f in v if f]).most_common(1)
    declared_main = declared_main[0][0] if declared_main else None
    direct_share = sum(chars_direct.values()) / max(1, sum(usage.values()))
    conflict = None
    top_used = usage.most_common(1)[0][0] if usage else None
    if declared_main and top_used and top_used.lower() != declared_main.lower():
        conflict = {"declared_theme_font": declared_main, "observed_primary_font": top_used,
                    "observed_share": round(usage[top_used] / max(1, sum(usage.values())), 3),
                    "direct_formatting_share": round(direct_share, 3), "style_guide_mentions": dict(guide_mentions),
                    "reason": f"The theme declares {declared_main}, but {round(100 * usage[top_used] / max(1, sum(usage.values())))}% of the text on the example slides is set in {top_used}"
                              + (" by direct formatting" if chars_direct.get(top_used) else "") + (f"; style-guide slides name {', '.join(guide_mentions)}" if guide_mentions else "") + "."}
    conf = round(ranked[0][1], 3) if ranked else 0.0
    fonts_report = []
    for f, sc in ranked[:6]:
        fonts_report.append({"font": f, "score": round(sc, 3), "declared": f in [x for v in declared.values() for x in v], "observed_chars": usage.get(f, 0),
                             "direct_chars": chars_direct.get(f, 0), "guide_mentions": guide_mentions.get(f, 0), "installed": f.lower() in installed})
    roles = {}
    for role, key in (("heading", "majorFont"), ("body", "minorFont")):
        c = by_role[role]
        dec = collections.Counter(th["fonts"].get(key) for th in themes if th["fonts"].get(key)).most_common(1)
        dec = dec[0][0] if dec else None
        if sum(c.values()) >= 40:
            f, n = c.most_common(1)[0]
            roles[role] = {"font": f, "source": "observed", "share": round(n / sum(c.values()), 3), "declared": dec, "conflict": bool(dec and f.lower() != dec.lower())}
        else:
            roles[role] = {"font": dec or primary, "source": "declared (too little example text to observe)", "share": None, "declared": dec, "conflict": False}
    return {"primary": primary, "confidence": conf, "declared": declared, "conflict": conflict, "candidates": fonts_report, "roles": roles,
            "evidence_weights": {k: round(w / norm, 3) for k, w in active.items()},
            "master_styles": [{"master": m["master_id"], "title": m["title_style"], "body": m["body_style"]} for m in masters],
            "sizes": _size_profile(prs)}


def _size_profile(prs) -> dict:
    sizes = collections.Counter()
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.text.strip():
                            sizes[round(r.font.size.pt)] += len(r.text.strip())
    return {"most_used_pt": [s for s, _ in sizes.most_common(6)]}


# ── palette ─────────────────────────────────────────────────────────────────────

def infer_palette(prs, masters, themes, sw, shh, theme_of_layout: dict) -> dict:
    fills, texts, bgs = collections.Counter(), collections.Counter(), collections.Counter()
    th0 = themes[0] if themes else {"colors": {}, "fonts": {}}
    for slide in prs.slides:
        th = theme_of_layout.get(slide.slide_layout.part.partname, th0)
        col = slide_background(slide, th, sw, shh)
        if col and col != "picture":
            bgs[col] += 1
        for sh in slide.shapes:
            try:
                c = _fill(sh, th)
                if c and not (_in(sh.width) >= 0.85 * sw and _in(sh.height) >= 0.85 * shh):
                    fills[c] += _in(sh.width) * _in(sh.height)
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            if not r.text.strip():
                                continue
                            rp = r._r.find(f"{A}rPr")
                            tc = _color_of(rp.find(f"{A}solidFill"), th) if rp is not None and rp.find(f"{A}solidFill") is not None else None
                            texts[tc or th["colors"].get("dk1", "000000")] += len(r.text.strip())
            except Exception:
                continue
    for mm in prs.slide_masters:  # template artwork (bars, marks, logos) is brand colour too
        thm = theme_of_master(mm)
        for shapes in [mm.shapes] + [lay.shapes for lay in mm.slide_layouts]:
            for sh in shapes:
                try:
                    c = _fill(sh, thm)
                    if c and not (_in(sh.width) >= 0.85 * sw and _in(sh.height) >= 0.85 * shh):
                        fills[c] += _in(sh.width) * _in(sh.height)
                except Exception:
                    continue
    total_bg = sum(bgs.values()) or 1
    brand_bg = [(c, n) for c, n in bgs.most_common() if not is_neutral(c) and sat(c) > 0.35 and lum(c) > 0.2]
    vivid = [(c, a) for c, a in fills.most_common() if sat(c) > 0.35 and lum(c) > 0.2]
    primary_vivid = (brand_bg[0][0] if brand_bg else (vivid[0][0] if vivid else None))
    text = texts.most_common(1)[0][0] if texts else th0["colors"].get("dk1")
    supporting = []
    for c, _ in fills.most_common(30):
        if c != primary_vivid and not is_neutral(c) and c not in supporting and all(_dist(c, s) > 40 for s in supporting + [primary_vivid or "000000"]):
            supporting.append(c)
    neutrals = [c for c, _ in fills.most_common(20) if is_neutral(c) and lum(c) > 0.8][:4]
    theme_accents = [th0["colors"].get(f"accent{i}") for i in range(1, 7) if th0["colors"].get(f"accent{i}")]
    all_accents = [t["colors"].get(f"accent{i}") for t in themes for i in (1, 2) if t["colors"].get(f"accent{i}")]
    conflict = None
    if primary_vivid and all_accents and all(_dist(primary_vivid, a) > 30 for a in all_accents):
        conflict = {"theme_accent1": theme_accents[0], "observed_primary": primary_vivid, "reason": "The most used brand colour on the slides is not the theme's first accent."}
    return {
        "primary": primary_vivid, "text": text, "supporting": supporting[:6], "neutrals": neutrals,
        "theme_scheme": th0["colors"], "theme_schemes_per_master": {m["master_id"]: m["theme"]["colors"] for m in masters},
        "brand_background_share": round(sum(n for c, n in brand_bg) / total_bg, 3) if bgs else None,
        "backgrounds": dict(bgs.most_common(6)), "text_colors": dict(texts.most_common(5)),
        "fills_by_area": {c: round(a, 1) for c, a in fills.most_common(10)}, "conflict": conflict,
        "confidence": round(min(0.95, 0.4 + 0.1 * len(prs.slides) ** 0.5), 2) if len(prs.slides) else 0.4,
    }


def _dist(a: str, b: str) -> float:
    return sum(abs(x - y) for x, y in zip(rgb(a), rgb(b)))


# ── grid ────────────────────────────────────────────────────────────────────────

def infer_grid(prs, catalog, sw, shh) -> dict:
    lefts, rights, edges = [], [], []
    for c in catalog:
        for p in [c["features"]["title_geometry"]] if c["features"]["title_geometry"] else []:
            lefts.append(p["nx"])
            rights.append(p["nx"] + p["nw"])
    for slide in prs.slides:
        for sh in slide.shapes:
            try:
                if _in(sh.width) >= 0.9 * sw or _in(sh.width) < 0.02 * sw:
                    continue
                x0, x1 = _in(sh.left) / sw, (_in(sh.left) + _in(sh.width)) / sw
                if 0 <= x0 <= 1 and 0 <= x1 <= 1:
                    edges += [x0, x1]
                    if shape_kind(sh) in ("text", "placeholder"):
                        lefts.append(x0)
                        rights.append(x1)
            except Exception:
                continue
    if not edges and not lefts:
        return {"columns": None, "confidence": 0.0}

    def mode(vals, lo, hi):
        vals = [round(v, 3) for v in vals if lo <= v <= hi]
        if not vals:
            return None
        return collections.Counter(round(v * 200) / 200 for v in vals).most_common(1)[0][0]

    ml = mode(lefts, 0.0, 0.2) or 0.05
    mr = 1 - (mode(rights, 0.8, 1.0) or 0.95)
    content = 1 - ml - mr
    best = None
    pts = [e for e in edges if ml - 0.005 <= e <= 1 - mr + 0.005] or [ml, 1 - mr]
    tol = 0.006
    for n in (4, 6, 8, 10, 12, 16):
        for g1000 in range(0, 41, 2):
            g = g1000 / 1000
            cw = (content - g * (n - 1)) / n
            if cw <= 0:
                continue
            bounds = []
            for i in range(n):
                x0 = ml + i * (cw + g)
                bounds += [x0, x0 + cw]
            hit = sum(1 for e in pts if min(abs(e - b) for b in bounds) <= tol) / len(pts)
            # prefer simpler systems unless a finer grid explains clearly more edges
            sc = hit - 0.004 * n
            if best is None or sc > best[0]:
                best = (sc, n, g, hit)
    _, n, g, hit = best
    return {"margin_left_in": round(ml * sw, 3), "margin_right_in": round(mr * sw, 3), "columns": n, "gutter_in": round(g * sw, 3),
            "edges_explained": round(hit, 3), "confidence": round(min(0.95, hit), 2), "edges_sampled": len(pts)}


# ── assets ──────────────────────────────────────────────────────────────────────

def infer_assets(all_art, prs, sw, shh) -> dict:
    logos = [dict(scope=s, owner=o, **{k: a[k] for k in ("name", "kind", "x", "y", "w", "h")}) for s, o, a in all_art if a.get("logo")]
    uniq_logo = {}
    for lg in logos:
        uniq_logo.setdefault((lg["name"], round(lg["x"], 1), round(lg["y"], 1)), lg)
    reserved = [dict(scope=s, owner=o, **{k: a[k] for k in ("name", "kind", "x", "y", "w", "h")}) for s, o, a in all_art if not a.get("logo") and a["nw"] * a["nh"] < 0.5]
    icons, pictures, vectors = set(), 0, 0
    for slide in prs.slides:
        for sh in slide.shapes:
            k = shape_kind(sh)
            if k == "picture":
                pictures += 1
                if _in(sh.width) < 0.12 * sw and _in(sh.height) < 0.2 * shh:
                    h = image_hash(sh)
                    if h:
                        icons.add(h)
            elif k == "group":
                vectors += 1
    corners = collections.Counter()
    for lg in uniq_logo.values():
        corners[("top" if lg["y"] < shh / 2 else "bottom") + "-" + ("left" if lg["x"] < sw / 2 else "right")] += 1
    return {"logos": list(uniq_logo.values())[:12], "logo_positions": dict(corners), "icons_distinct": len(icons), "pictures_on_slides": pictures,
            "vector_groups_on_slides": vectors, "reserved_artwork": len(reserved), "reserved_examples": reserved[:12]}


# ── brand rules ─────────────────────────────────────────────────────────────────

def _case_of(t: str) -> str | None:
    words = [w for w in re.findall(r"[A-Za-zÁÉÍÓÚÑáéíóúñ][\w'’-]*", t) if len(w) > 3]
    if len(words) < 3:
        return None
    rest = words[1:]
    caps = sum(1 for w in rest if w[0].isupper() and not w.isupper())
    if t.isupper():
        return "upper"
    return "title" if caps / len(rest) > 0.6 else "sentence"


def infer_rules(prs, catalog, slides_obs, palette, sw, shh, theme_of_layout: dict) -> dict:
    cases = collections.Counter()
    align = collections.Counter()
    shapes = collections.Counter()
    title_tops = []
    for slide in prs.slides:
        for sh in slide.shapes:
            pg = prst(sh)
            if pg:
                shapes[pg] += 1
            if sh.is_placeholder and ph_type(sh) in ("TITLE", "CENTER_TITLE") and sh.has_text_frame:
                c = _case_of(sh.text_frame.text)
                if c:
                    cases[c] += 1
                title_tops.append(_in(sh.top) / shh)
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    if p.text.strip():
                        a = p._p.find(f"{A}pPr")
                        align[(a.get("algn") if a is not None and a.get("algn") else "l")] += len(p.text)
    rects = shapes.get("rect", 0)
    rounds = shapes.get("roundRect", 0) + shapes.get("round2SameRect", 0) + shapes.get("snipRoundRect", 0)
    primary = palette.get("primary")
    first_last = []
    for si in (0, len(prs.slides) - 1) if len(prs.slides) else ():
        slide = prs.slides[si]
        th = theme_of_layout.get(slide.slide_layout.part.partname) or next(iter(theme_of_layout.values()))
        first_last.append(slide_background(slide, th, sw, shh))
    bookend = bool(primary) and len(first_last) == 2 and all(c and c != "picture" and _dist(c, primary) < 40 for c in first_last)
    same_ends = len(first_last) == 2 and first_last[0] == first_last[1] and first_last[0] not in (None, "picture") and not is_neutral(first_last[0])
    tot_case = sum(cases.values()) or 1
    tot_al = sum(align.values()) or 1
    content_titles = [c["features"]["title_geometry"] for c in catalog if c["classification"][0]["type"] in ("content", "one_column", "two_column", "chart", "table") and c["features"]["title_geometry"]]
    return {
        "headline_case": {"dominant": cases.most_common(1)[0][0] if cases else None, "shares": {k: round(v / tot_case, 2) for k, v in cases.items()}, "titles_sampled": sum(cases.values())},
        "bookend": {"first_and_last_in_brand_colour": bookend, "first_and_last_share_a_colour": same_ends, "first_last_backgrounds": first_last, "brand_colour": primary},
        "brand_colour_slide_share": palette.get("brand_background_share"),
        "text_alignment": {k: round(v / tot_al, 2) for k, v in align.most_common()},
        "headline_position": {"content_title_top_in": round(sorted(t["y"] for t in content_titles)[len(content_titles) // 2], 3) if content_titles else None,
                              "content_title_width_share": round(sorted(t["nw"] for t in content_titles)[len(content_titles) // 2], 3) if content_titles else None},
        "shapes": {"rounded_share_of_rectangles": round(rounds / (rounds + rects), 2) if rounds + rects else None, "chevrons": shapes.get("chevron", 0) + shapes.get("homePlate", 0),
                   "most_used": dict(shapes.most_common(6))},
    }
