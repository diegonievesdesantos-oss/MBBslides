"""Native, editable tables that behave like exhibits — not Excel dumps.

Supports: typed columns with number formats (right-aligned numbers, reduced
precision), header rule instead of fills, hairline row separators, subtotal /
total rows, hierarchy (indent), highlighted rows / columns, heatmap columns
(sequential or diverging, auto-contrast text), delta columns (sign-coloured),
Harvey-ball columns (0-4) and RAG status columns (drawn as editable shapes
centred on the cell).

Spec:
{
  "type": "table" | "heatmap" | "harvey_table" | "scorecard",
  "title": "...", "unit": "...",
  "columns": [{"label": "Segment", "align": "left", "width": 2.2},
              {"label": "2025", "format": {"decimals": 0, "suffix": "M"}, "kind": "number"},
              {"label": "Δ", "kind": "delta", "format": {"decimals": 1, "suffix": "pp"}},
              {"label": "Fit", "kind": "harvey"}, {"label": "Status", "kind": "rag"}],
  "rows": [["Grocery", 120, -1.8, 3, "R"], {"cells": [...], "style": "subtotal|total|highlight", "indent": 1}],
  "heatmap": {"columns": [1, 2], "mode": "sequential|diverging", "domain": [lo, hi]},
  "highlight_columns": [1]
}
"""
from __future__ import annotations

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from ..charts.numfmt import fmt
from ..design import text_metrics as tm
from ..design.tokens import FONT_FLOOR, LINES, best_text_on, interpolate, legible_fill, sequential_color
from ..layout.engine import Box
from ..pptx.painter import E, Painter, rgb
from ..pptx.text_components import exhibit_header

NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
CELL_PAD_X = 0.08
CELL_PAD_Y = 0.045
RAG = {"R": "negative", "A": "warning", "G": "positive", "red": "negative", "amber": "warning", "green": "positive", "grey": "neutral"}


def _norm_rows(ex: dict) -> list[dict]:
    out = []
    for r in ex.get("rows") or []:
        if isinstance(r, dict):
            out.append({"cells": list(r.get("cells", [])), "style": r.get("style"), "indent": r.get("indent", 0)})
        else:
            out.append({"cells": list(r), "style": None, "indent": 0})
    return out


def _norm_cols(ex: dict, rows: list[dict]) -> list[dict]:
    cols = ex.get("columns")
    if not cols:
        header = ex.get("header") or [f"Col {i + 1}" for i in range(len(rows[0]["cells"]))]
        cols = [{"label": h} for h in header]
    cols = [dict(c) if isinstance(c, dict) else {"label": str(c)} for c in cols]
    ex_kind = ex.get("type")
    for j, c in enumerate(cols):
        if "kind" not in c:
            vals = [r["cells"][j] for r in rows if j < len(r["cells"])]
            numeric = all(isinstance(v, (int, float)) or v is None for v in vals) and any(isinstance(v, (int, float)) for v in vals)
            c["kind"] = "harvey" if (ex_kind == "harvey_table" and numeric and j > 0) else ("number" if numeric else "text")
        c.setdefault("align", "left" if c["kind"] == "text" else ("center" if c["kind"] in ("harvey", "rag") else "right"))
        if c["kind"] in ("number", "delta") and "decimals" not in (c.get("format") or {}):
            vals = [r["cells"][j] for r in rows if j < len(r["cells"])]
            c["format"] = {**(c.get("format") or {}), "decimals": data_decimals(vals)}
    return cols


def data_decimals(vals) -> int:
    """Decimals the data actually carries (max 2): 1.9 must not print as "2"."""
    d = 0
    for v in vals:
        if isinstance(v, float) and not v.is_integer():
            txt = repr(v)
            d = max(d, len(txt.split(".")[1]) if "." in txt and "e" not in txt else 2)
    return min(d, 2)


def _cell_text(v, col: dict) -> str:
    if v is None:
        return "–"
    k = col["kind"]
    if k in ("harvey", "rag"):
        return ""
    if k in ("number", "delta") and isinstance(v, (int, float)):
        f = {"decimals": 0, "prefix": "", "suffix": "", "percent": False, "thousands": True, **(col.get("format") or {})}
        return fmt(float(v), f, plus=(k == "delta"))
    return str(v)


def _tc_borders(cell, top=None, bottom=None, left=None, right=None, fill: str | None = None) -> None:
    """Rebuild <a:tcPr> in schema order: lnL, lnR, lnT, lnB, fill."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for ch in list(tcPr):
        tcPr.remove(ch)

    def ln(tag, spec):
        el = etree.SubElement(tcPr, qn(tag))
        if spec is None:
            el.set("w", "0")
            etree.SubElement(el, qn("a:noFill"))
        else:
            color, width = spec
            el.set("w", str(int(width * 12700)))
            el.set("cap", "flat")
            sf = etree.SubElement(el, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), {"val": color})
            etree.SubElement(el, qn("a:prstDash"), {"val": "solid"})

    ln("a:lnL", left)
    ln("a:lnR", right)
    ln("a:lnT", top)
    ln("a:lnB", bottom)
    if fill:
        sf = etree.SubElement(tcPr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), {"val": fill})
    else:
        etree.SubElement(tcPr, qn("a:noFill"))


def _col_widths(p: Painter, cols, rows, total_w: float, size: float) -> list[float]:
    """Needed widths from real metrics; text columns capped, value columns equalised
    (a heatmap / numeric grid reads best with equal cells), slack to the label column."""
    need = []
    for j, c in enumerate(cols):
        head = max((tm.text_width_in(w, size, True) for w in str(c["label"]).split()), default=0) + 2 * CELL_PAD_X + 0.04
        cells = [_cell_text(r["cells"][j] if j < len(r["cells"]) else None, c) for r in rows]
        body = max((tm.text_width_in(t, size, r["style"] in ("subtotal", "total")) for t, r in zip(cells, rows)), default=0) + 2 * CELL_PAD_X + 0.04
        if c["kind"] in ("harvey", "rag"):
            body = 0.6
        need.append(max(0.55, head, body if c["kind"] != "text" else min(body, total_w * 0.42)))
    widths = [float(c["width"]) if c.get("width") else need[j] for j, c in enumerate(cols)]
    free = [j for j, c in enumerate(cols) if not c.get("width")]
    slack = total_w - sum(widths)
    if slack < 0:
        scale = total_w / sum(widths)
        return [w * scale for w in widths]
    value_cols = [j for j in free if cols[j]["kind"] != "text"]
    text_cols = [j for j in free if cols[j]["kind"] == "text"]
    if value_cols:
        # equalise value columns up to a comfortable width, keep the rest for labels
        target = max(widths[j] for j in value_cols)
        comfy = max(target, min(1.5, (sum(widths[j] for j in value_cols) + slack * (0.6 if text_cols else 1.0)) / len(value_cols)))
        for j in value_cols:
            slack -= comfy - widths[j]
            widths[j] = comfy
    if slack > 0:
        tgt = text_cols or free or [0]
        for j in tgt:
            widths[j] += slack / len(tgt)
    return widths


def render(p: Painter, box: Box, ex: dict) -> dict:
    plot = exhibit_header(p, box, ex)
    if ex.get("type") == "harvey_table":
        plot = Box(plot.x, plot.y, plot.w, plot.h - 0.4)  # room for the Harvey legend
    rows = _norm_rows(ex)
    cols = _norm_cols(ex, rows)
    theme = p.theme
    nrows, ncols = len(rows) + 1, len(cols)
    base_size = p.style("table_body")["size"]
    floor = FONT_FLOOR["table_body"]
    size = base_size
    widths = _col_widths(p, cols, rows, plot.w, size)

    def heights(sz):
        hs = []
        head_lines = max(len(tm.wrap_lines(c["label"], max(0.2, w - 2 * CELL_PAD_X), sz, True)) for c, w in zip(cols, widths))
        hs.append(head_lines * tm.line_height_in(sz) + 2 * CELL_PAD_Y + 0.06)
        for r in rows:
            lines = 1
            for j, (c, w) in enumerate(zip(cols, widths)):
                t = _cell_text(r["cells"][j] if j < len(r["cells"]) else None, c)
                ind = 0.18 * r["indent"] if j == 0 else 0
                lines = max(lines, len(tm.wrap_lines(t, max(0.2, w - 2 * CELL_PAD_X - ind), sz, r["style"] in ("subtotal", "total"))))
            hs.append(max(0.3, lines * tm.line_height_in(sz) + 2 * CELL_PAD_Y + 0.04))
        return hs

    hs = heights(size)
    while sum(hs) > plot.h and size > floor:
        size -= 0.5
        widths = _col_widths(p, cols, rows, plot.w, size)
        hs = heights(size)
    overflow = sum(hs) > plot.h + 1e-3
    # distribute spare height (tables should not float in a sea of white)
    spare = plot.h - sum(hs)
    if spare > 0 and ex.get("stretch", True):
        per = min(spare / len(rows), p.profile.get("table_stretch", 0.22))
        hs = [hs[0]] + [h + per for h in hs[1:]]
    p.manifest.fits.append({"zone": p.zone, "role": "table", "what": ex.get("title", "table"), "base_pt": base_size, "chosen_pt": size, "need_h": round(sum(hs), 3), "box_h": round(plot.h, 3), "lines": nrows, "overflow": overflow})

    gf = p.slide.shapes.add_table(nrows, ncols, E(plot.x), E(plot.y), E(sum(widths)), E(sum(hs)))
    gf.name = p._name("table")
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE
    for j, w in enumerate(widths):
        tbl.columns[j].width = E(w)
    for i, h in enumerate(hs):
        tbl.rows[i].height = E(h)

    hl_cols = set(ex.get("highlight_columns") or [])
    heat = ex.get("heatmap") or ({"columns": list(range(1, ncols))} if ex.get("type") == "heatmap" else None)
    heat_cols = set(heat.get("columns", [])) if heat else set()
    heat_vals = [r["cells"][j] for r in rows for j in heat_cols if j < len(r["cells"]) and isinstance(r["cells"][j], (int, float)) and r["style"] not in ("total", "subtotal")]
    if heat and heat_vals:
        dlo, dhi = heat.get("domain") or (min(heat_vals), max(heat_vals))
    rule = theme.c("rule")
    strong = theme.c("primary")
    grid = theme.c("gridline")

    def set_text(cell, text, bold, color, align, indent=0.0, sz=size):
        cell.margin_left = E(CELL_PAD_X + indent)
        cell.margin_right = E(CELL_PAD_X)
        cell.margin_top = E(CELL_PAD_Y)
        cell.margin_bottom = E(CELL_PAD_Y)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        para = tf.paragraphs[0]
        para.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
        for r_ in list(para.runs):
            r_._r.getparent().remove(r_._r)
        run = para.add_run()
        run.text = text
        run.font.name = theme.font_latin
        run.font.size = Pt(sz)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)

    # header
    for j, c in enumerate(cols):
        cell = tbl.cell(0, j)
        _tc_borders(cell, top=None, bottom=(strong, LINES["rule"] * 1.3))
        set_text(cell, c["label"], True, theme.c("primary") if j in hl_cols else theme.c("text"), c["align"] if c["kind"] != "text" else "left")
    shapes_after = []
    y = plot.y + hs[0]
    for i, r in enumerate(rows, start=1):
        st = r["style"]
        row_fill = theme.c("faint") if st == "highlight" else None
        bold_row = st in ("subtotal", "total", "highlight")
        for j, c in enumerate(cols):
            v = r["cells"][j] if j < len(r["cells"]) else None
            cell = tbl.cell(i, j)
            fill = row_fill
            color = theme.c("text")
            if j in hl_cols and not fill:
                fill = theme.c("faint")
            if heat and j in heat_cols and isinstance(v, (int, float)) and st not in ("total", "subtotal"):
                t = 0.0 if dhi == dlo else (v - dlo) / (dhi - dlo)
                if heat.get("mode") == "diverging":
                    stops = theme.diverging
                    pos = max(0.0, min(1.0, t)) * (len(stops) - 1)
                    k = min(int(pos), len(stops) - 2)
                    fill = interpolate(stops[k], stops[k + 1], pos - k)
                else:
                    fill = sequential_color(theme, 0.08 + 0.85 * max(0.0, min(1.0, t)))
                fill = legible_fill(fill, theme)
                color = best_text_on(fill, theme)
            if c["kind"] == "delta" and isinstance(v, (int, float)):
                good = c.get("good", "up")
                color = theme.c("positive") if (v > 0) == (good == "up") and v != 0 else (theme.c("negative") if v != 0 else theme.c("text_muted"))
            top = None
            if st == "subtotal":
                top = (rule, LINES["rule"])
            elif st == "total":
                top = (strong, LINES["rule"] * 1.3)
            bottom = (grid, LINES["hairline"]) if i < len(rows) else (rule, LINES["rule"])
            _tc_borders(cell, top=top, bottom=bottom, fill=fill)
            set_text(cell, _cell_text(v, c), bold_row or (j == 0 and st is None and ex.get("bold_first_column", False)), color, c["align"], indent=0.18 * r["indent"] if j == 0 else 0.0)
            if c["kind"] in ("harvey", "rag") and v is not None:
                shapes_after.append((c["kind"], v, sum(widths[:j]) + plot.x, y, widths[j], hs[i], j))
        y += hs[i]
    # overlays: Harvey balls / RAG dots (editable shapes centred on cells)
    for kind, v, cx, cy, cw, ch, j in shapes_after:
        d = min(0.26, ch - 0.08)
        bx = Box(cx + cw / 2 - d / 2, cy + ch / 2 - d / 2, d, d)
        if kind == "rag":
            p.oval(bx, fill=RAG.get(str(v), "neutral"), kind="marker")
        else:
            _harvey(p, bx, float(v), "primary" if j in hl_cols else "secondary")
    if ex.get("type") == "harvey_table":
        _harvey_legend(p, Box(plot.x, plot.y + sum(hs) + 0.1, plot.w, 0.28))
    return {"type": ex.get("type", "table"), "rows": len(rows), "cols": ncols, "font_pt": size, "overflow": overflow}


def _harvey(p: Painter, bx: Box, v: float, color: str = "secondary") -> None:
    """Harvey ball in the secondary colour: the focus colour stays free for the highlighted column."""
    q = max(0, min(4, int(round(v))))
    p.oval(bx, fill="background", line=color, line_w=LINES["rule"], kind="marker")
    if q == 0:
        return
    if q == 4:
        p.oval(bx, fill=color, line=color, kind="marker")
        return
    sp = p.shape(MSO_SHAPE.PIE, bx, fill=color, kind="marker")
    # adj1 = start angle, adj2 = end angle (60000ths of a degree), clockwise from 12 o'clock
    start = 270.0
    end = (start + 90.0 * q) % 360
    prstGeom = sp._element.spPr.find(qn("a:prstGeom"))
    av = prstGeom.find(qn("a:avLst"))
    for ch in list(av):
        av.remove(ch)
    etree.SubElement(av, qn("a:gd"), {"name": "adj1", "fmla": f"val {int(start * 60000)}"})
    etree.SubElement(av, qn("a:gd"), {"name": "adj2", "fmla": f"val {int(end * 60000)}"})


def _harvey_legend(p: Painter, box: Box) -> None:
    x = box.x
    for q, lab in ((0, "None"), (2, "Partial"), (4, "Full")):
        _harvey(p, Box(x, box.y + 0.04, 0.18, 0.18), q, "secondary")
        p.text(Box(x + 0.24, box.y, 0.8, 0.26), lab, role="chart_axis", fit=False, kind="label", anchor="middle")
        x += 1.0
