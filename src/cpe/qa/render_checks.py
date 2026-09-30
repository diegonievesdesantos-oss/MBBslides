"""Render-based QA: inspects what LibreOffice actually drew.

The PDF exported by LibreOffice carries every rendered text span with its
bounding box. Comparing those spans with the shapes of the .pptx detects the
defects a model-based check can miss (renderer wraps differently, chart labels
collide, text spills out of its box). The PNG is used for ink coverage.

Codes:
  RENDER_TEXT_SPILL (error)      rendered text lies outside every text box / table / chart frame
  RENDER_OFF_SLIDE (error)       rendered text beyond the page
  RENDER_OUTSIDE_SAFE (error)    rendered text in the margins
  RENDER_TEXT_COLLISION (error)  two rendered spans overlap
  RENDER_HEADLINE_LINES (error)  headline renders on more than 2 lines
  RENDER_HEADLINE_WIDOW (warn)   one short word alone on the last headline line
  RENDER_SMALL_TEXT (error)      rendered text below 7.5 pt
  RENDER_LABEL_TRUNCATED (error) chart axis labels cut with an ellipsis by the renderer
  RENDER_LABEL_ROTATED (error)   chart axis labels rotated because they do not fit
  RENDER_TOO_EMPTY (warning)     body ink coverage very low
  RENDER_UNBALANCED (info)       large empty region next to dense content
"""
from __future__ import annotations

from pptx import Presentation
from pptx.util import Emu

from ..design.tokens import GRID, SLIDE_H, SLIDE_W
from ..layout.engine import Box
from ..spec import issue
from .geometry import _parse_name


def _spans(page) -> list[dict]:
    out = []
    d = page.get_text("dict")
    for block in d.get("blocks", []):
        for line in block.get("lines", []):
            for sp in line.get("spans", []):
                t = sp.get("text", "")
                if not t.strip():
                    continue
                x0, y0, x1, y1 = sp["bbox"]
                out.append({"text": t, "box": Box(x0 / 72, y0 / 72, (x1 - x0) / 72, (y1 - y0) / 72), "size": sp.get("size", 0), "origin_y": sp.get("origin", (0, y1))[1] / 72, "dir": line.get("dir", (1, 0))})
    return out


def _containers(slide) -> list[tuple[Box, str, str]]:
    out = []
    for sh in slide.shapes:
        zone, kind = _parse_name(sh.name)
        b = Box(Emu(sh.left).inches, Emu(sh.top).inches, Emu(sh.width).inches, Emu(sh.height).inches)
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append((b, "text", sh.name))
        elif getattr(sh, "has_chart", False) and sh.has_chart:
            out.append((b, "chart", sh.name))
        elif getattr(sh, "has_table", False) and sh.has_table:
            out.append((b, "table", sh.name))
    return out


def _ink_stats(png_path: str, band: Box) -> dict:
    from PIL import Image

    im = Image.open(png_path).convert("L")
    W, H = im.size
    sx, sy = W / SLIDE_W, H / SLIDE_H
    crop = im.crop((int(band.x * sx), int(band.y * sy), int(band.r * sx), int(band.b * sy)))
    cw, ch = crop.size
    px = crop.load()
    gx, gy = 8, 4
    cells = []
    total_ink = 0
    total = 0
    step = 3
    for j in range(gy):
        for i in range(gx):
            ink = n = 0
            for y in range(int(j * ch / gy), int((j + 1) * ch / gy), step):
                for x in range(int(i * cw / gx), int((i + 1) * cw / gx), step):
                    n += 1
                    if px[x, y] < 235:
                        ink += 1
            cells.append(ink / max(1, n))
            total_ink += ink
            total += n
    return {"coverage": total_ink / max(1, total), "cells": cells, "grid": (gx, gy)}


def check(pdf_path: str, pptx_path: str, manifests: list[dict], pngs: list[str] | None = None) -> tuple[list[dict], dict]:
    import pymupdf

    doc = pymupdf.open(pdf_path)
    prs = Presentation(pptx_path)
    out: list[dict] = []
    metrics: dict = {}
    for idx, (page, slide) in enumerate(zip(doc, prs.slides), start=1):
        man = manifests[idx - 1] if idx - 1 < len(manifests) else {}
        sid = man.get("slide_id") or f"#{idx}"
        spans = _spans(page)
        conts = _containers(slide)
        full_bleed = man.get("layout") in ("divider",)
        for sp in spans:
            b = sp["box"]
            if b.x < -0.02 or b.y < -0.02 or b.r > SLIDE_W + 0.02 or b.b > SLIDE_H + 0.02:
                out.append(issue("error", "RENDER_OFF_SLIDE", f"'{sp['text'][:30]}' rendered beyond the slide", sid))
                continue
            if not full_bleed and (b.x < GRID.margin_l - 0.05 or b.r > SLIDE_W - GRID.margin_r + 0.05 or b.b > GRID.footer_y + GRID.footer_h + 0.08):
                out.append(issue("error", "RENDER_OUTSIDE_SAFE", f"'{sp['text'][:30]}' rendered in the margin", sid))
            # glyph boxes include ascent/descent padding: shrink before testing containment
            core = b.inset(0.01, b.h * 0.18, 0.01, b.h * 0.18)
            if not any(c.contains(core, 0.05) for c, _, _ in conts):
                out.append(issue("error", "RENDER_TEXT_SPILL", f"'{sp['text'][:40]}' is rendered outside its box (overflow)", sid))
            if sp["size"] and sp["size"] < 7.5:
                out.append(issue("error", "RENDER_SMALL_TEXT", f"'{sp['text'][:30]}' rendered at {sp['size']:.1f} pt", sid))
            in_chart = any(k == "chart" and c.contains(core, 0.05) for c, k, _ in conts)
            if in_chart and sp["text"].rstrip().endswith(("...", "…")):
                out.append(issue("error", "RENDER_LABEL_TRUNCATED", f"Chart label truncated by the renderer: '{sp['text'][:30]}'", sid))
            if in_chart and abs(sp["dir"][1]) > 0.05:
                out.append(issue("error", "RENDER_LABEL_ROTATED", f"Chart label rotated by the renderer (not enough room): '{sp['text'][:30]}'", sid))
        # collisions between spans on different lines
        seen = set()
        for i in range(len(spans)):
            a = spans[i]
            ca = a["box"].inset(0.005, a["box"].h * 0.22, 0.005, a["box"].h * 0.22)
            for j in range(i + 1, len(spans)):
                c = spans[j]
                if abs(a["origin_y"] - c["origin_y"]) < 0.01 and (a["box"].r <= c["box"].x + 0.02 or c["box"].r <= a["box"].x + 0.02):
                    continue  # consecutive runs on the same line
                cc = c["box"].inset(0.005, c["box"].h * 0.22, 0.005, c["box"].h * 0.22)
                inter = ca.intersection(cc)
                if inter > 0.0015 and inter / max(1e-6, min(ca.w * ca.h, cc.w * cc.h)) > 0.12:
                    key = (a["text"][:20], c["text"][:20])
                    if key in seen:
                        continue
                    seen.add(key)
                    out.append(issue("error", "RENDER_TEXT_COLLISION", f"Rendered text '{a['text'][:25]}' collides with '{c['text'][:25]}'", sid))
        # headline lines from the render
        # include the area below the box: an overflowing third line lands there
        head_zone = Box(GRID.margin_l - 0.05, GRID.headline_y - 0.05, GRID.content_w + 0.1, GRID.headline_h + 0.7)
        head_spans = [sp for sp in spans if head_zone.contains(sp["box"], 0.02) and sp["size"] >= 18]
        if head_spans:  # keep only the headline's own size (not an 18 pt number just below it)
            hs = max(sp["size"] for sp in head_spans if sp["box"].y < GRID.headline_y + 0.5) if any(sp["box"].y < GRID.headline_y + 0.5 for sp in head_spans) else 0
            head_spans = [sp for sp in head_spans if abs(sp["size"] - hs) < 0.6]
        if head_spans and man.get("layout") not in ("cover", "divider"):
            ys = sorted({round(sp["origin_y"], 2) for sp in head_spans})
            if len(ys) > 2:
                out.append(issue("error", "RENDER_HEADLINE_LINES", f"Headline renders on {len(ys)} lines", sid))
            if len(ys) >= 2:
                last = " ".join(sp["text"] for sp in head_spans if round(sp["origin_y"], 2) == ys[-1]).strip()
                if len(last.split()) == 1 and len(last) < 12:
                    out.append(issue("warning", "RENDER_HEADLINE_WIDOW", f"Headline's last rendered line is only '{last}'", sid))
        # ink coverage of the body
        if pngs and idx - 1 < len(pngs) and man.get("layout") not in ("cover", "divider", "agenda", "statement"):
            band = Box(GRID.margin_l, GRID.body_y, GRID.content_w, GRID.body_h)
            st = _ink_stats(pngs[idx - 1], band)
            metrics[sid] = {"ink_coverage": round(st["coverage"], 3)}
            if st["coverage"] < 0.015:
                out.append(issue("warning", "RENDER_TOO_EMPTY", f"Body ink coverage {st['coverage']:.1%}: the slide looks empty", sid))
            empty = sum(1 for c in st["cells"] if c < 0.004)
            if empty / len(st["cells"]) > 0.45 and st["coverage"] > 0.03:
                out.append(issue("info", "RENDER_UNBALANCED", f"{empty}/{len(st['cells'])} body cells empty next to dense content: check balance", sid))
    doc.close()
    return out, metrics
