"""Existing-deck ingestion and update (v1.8): old client deck + new data + new question → updated deck.

    cpe deck ingest old.pptx -o work/         → work/old_deck.json      what the old deck says and how
                                                work/old_deck_spec.json a starting deck.json: same storyline,
                                                                         roles and exhibits rebuilt natively
                                                work/old_ghost.md       the old argument, headlines only
    cpe deck stale work/                       → work/update_plan.json/.md  every number in the old deck
                                                checked against the NEW fact model: current / outdated
                                                (with the new value and its fact id) / untraced

The old deck is a source of structure and conventions, never of facts: an old number is reused only
once the new fact model confirms it. Nothing is silently carried over.
"""
from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

from ..qa.proof import headline_quantities

ROLE_RULES = [
    ("agenda", r"\b(agenda|índice|indice|contents|contenido)\b"),
    ("exec_summary", r"\b(executive summary|resumen ejecutivo|resumen|summary|key messages|mensajes clave)\b"),
    ("next_steps", r"\b(next steps|próximos pasos|proximos pasos|decisions?|decisiones|asks?|la petición)\b"),
    ("appendix", r"\b(appendix|anexo|apéndice|backup)\b"),
]
CHART_TYPE = {"BAR_CLUSTERED": "bar", "COLUMN_CLUSTERED": "column", "BAR_STACKED": "stacked_bar", "COLUMN_STACKED": "stacked_column",
              "BAR_STACKED_100": "stacked_100", "COLUMN_STACKED_100": "stacked_100", "LINE": "line", "LINE_MARKERS": "line", "PIE": "pie",
              "DOUGHNUT": "donut", "AREA": "area", "XY_SCATTER": "scatter", "BUBBLE": "bubble"}


def _shape_texts(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            size = max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0)
            is_title = bool(getattr(sh, "is_placeholder", False) and sh.placeholder_format.type is not None
                            and "TITLE" in str(sh.placeholder_format.type))
            out.append({"text": " ".join(sh.text_frame.text.split()), "size": size, "top": sh.top or 0, "title": is_title,
                        "fonts": [r.font.name for p in sh.text_frame.paragraphs for r in p.runs if r.font.name]})
    return out


def ingest_deck(path: str | Path) -> dict:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides, fonts, n = [], Counter(), len(prs.slides)
    for i, slide in enumerate(prs.slides, start=1):
        texts = _shape_texts(slide)
        for t in texts:
            fonts.update(t["fonts"])
        # the action title: a title placeholder, else the top-most SENTENCE in a large size (a KPI hero
        # "€13.6bn" is large but is not the headline), else the largest text
        sentences = [t for t in texts if len(t["text"].split()) >= 4 and t["size"] >= 16]
        title = (next((t["text"] for t in texts if t["title"]), None)
                 or (min(sentences, key=lambda t: t["top"])["text"] if sentences else None)
                 or (max(texts, key=lambda t: (t["size"], -t["top"]))["text"] if texts else ""))
        body = [t["text"] for t in texts if t["text"] != title]
        exhibits = []
        for sh in slide.shapes:
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                try:
                    plot = ch.plots[0]
                    exhibits.append({"type": CHART_TYPE.get(str(ch.chart_type).split(".")[-1].split(" ")[0], "column"),
                                     "title": ch.chart_title.text_frame.text if ch.has_title else "",
                                     "categories": [str(c) for c in plot.categories],
                                     "series": [{"name": s.name, "values": [v for v in s.values]} for s in plot.series]})
                except Exception:  # an unusual chart: keep its presence, not its data
                    exhibits.append({"type": "chart", "unreadable": True})
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text for c in r.cells] for r in sh.table.rows]
                exhibits.append({"type": "table", "header": rows[0] if rows else [], "rows": rows[1:]})
        words = sum(len(t["text"].split()) for t in texts)
        labels = " ".join([title] + [b for b in body if len(b.split()) <= 4])  # trackers / section labels ("Executive summary")
        low = labels.lower()
        role = next((r for r, rx in ROLE_RULES if re.search(rx, low)), None)
        if i == 1:
            role = "title"
        elif not role and not exhibits and words <= 12:
            role = "divider"
        role = role or "content"
        numbers = []
        for where, text in [("title", title)] + [(f"body[{k}]", b) for k, b in enumerate(body)]:
            for q in headline_quantities(text):
                numbers.append({"where": where, "raw": q["raw"], "value": q["value"], "kind": q["kind"], "context": text[:160]})
        for e_i, ex in enumerate(exhibits):
            for s in ex.get("series") or []:
                for k, v in enumerate(s["values"]):
                    if isinstance(v, (int, float)):
                        cat = (ex.get("categories") or [""] * (k + 1))[k] if k < len(ex.get("categories") or []) else ""
                        numbers.append({"where": f"exhibit[{e_i}].{s['name']}[{cat}]", "raw": f"{v:g}", "value": float(v), "kind": "data",
                                        "context": f"{ex.get('title') or ''} {s['name']} {cat}".strip()})
            for r_i, row in enumerate(ex.get("rows") or []):
                for c_i, cell in enumerate(row[1:], start=1):
                    for q in headline_quantities(cell):
                        hdr = (ex.get("header") or [""] * (c_i + 1))[c_i] if c_i < len(ex.get("header") or []) else ""
                        numbers.append({"where": f"exhibit[{e_i}].rows[{r_i}][{c_i}]", "raw": q["raw"], "value": q["value"], "kind": q["kind"],
                                        "context": f"{row[0]} {hdr}".strip()})
        slides.append({"n": i, "role": role, "layout": slide.slide_layout.name, "headline": title, "body": body, "exhibits": exhibits,
                       "numbers": numbers, "words": words, "is_last": i == n})
    return {"source": Path(path).name, "slides": slides,
            "conventions": {"fonts": [f for f, _ in fonts.most_common(4)], "slide_count": n,
                            "layouts": Counter(s["layout"] for s in slides).most_common(8)}}


def _waterfall_steps(ex: dict) -> list[dict]:
    """A waterfall drawn as stacked columns (an invisible Base + visible Total / Increase / Decrease
    series, as MBBslides and most consultants draw them) back to its steps."""
    ser = {x["name"]: [float(v or 0) for v in x["values"]] for x in ex["series"]}
    steps = []
    for k, cat in enumerate(ex["categories"]):
        tot = sum(abs(ser[n][k]) for n in ser if n.startswith("Total"))
        up = sum(abs(ser[n][k]) for n in ser if n.startswith("Increase"))
        down = sum(abs(ser[n][k]) for n in ser if n.startswith("Decrease"))
        if tot:
            neg = any(ser[n][k] < 0 for n in ser if n.startswith("Total"))
            steps.append({"label": cat, "value": -tot if neg else tot, "type": "total"})
        else:
            steps.append({"label": cat, "value": up - down})
    return steps


def to_spec(inv: dict, title: str = "") -> dict:
    """A starting deck.json that keeps the old storyline, roles and exhibits (rebuilt as native,
    editable visuals). Every number in it is still the OLD number until `update_plan` confirms it."""
    slides = []
    for s in inv["slides"]:
        sid = f"s{s['n']:02d}"
        if s["role"] == "title":
            slides.append({"id": sid, "kind": "cover", "title": s["headline"], "subtitle": " ".join(s["body"][:1])})
            continue
        if s["role"] == "divider":
            slides.append({"id": sid, "kind": "divider", "title": s["headline"]})
            continue
        sl = {"id": sid, "kind": "exec_summary" if s["role"] == "exec_summary" else "content", "headline": s["headline"],
              "purpose": f"(from the old deck, slide {s['n']}: {s['role']})", "_old_slide": s["n"]}
        ex = next((e for e in s["exhibits"] if not e.get("unreadable")), None)
        if ex and ex["type"] == "table":
            sl["visual"] = {"type": "table", "columns": [{"label": h} for h in ex["header"]], "rows": ex["rows"]}
        elif ex and {x["name"] for x in ex["series"]} >= {"Base"} and any(x["name"].startswith(("Increase", "Decrease", "Total")) for x in ex["series"]):
            sl["visual"] = {"type": "waterfall", "title": ex.get("title") or "", "data": {"steps": _waterfall_steps(ex)}}
        elif ex:
            sl["visual"] = {"type": ex["type"], "title": ex.get("title") or "", "data": {"categories": ex["categories"], "series": ex["series"]}}
        elif s["body"]:
            sl["commentary"] = {"points": s["body"][:5]}
        slides.append(sl)
    return {"meta": {"title": title or (inv["slides"][0]["headline"] if inv["slides"] else ""), "from_deck": inv["source"]},
            "storyline": {"governing_thought": next((s["headline"] for s in inv["slides"] if s["role"] == "exec_summary"), "")}, "slides": slides}


def _words(t: str) -> set[str]:
    from .conflicts import STOP
    return {w for w in re.findall(r"[a-záéíóúñü]{3,}", (t or "").lower())} - STOP - {"hasta", "con", "más", "than", "with", "from", "desde"}


UNIT_TEXT = {"EUR": "€", "EUR_K": "k€", "EUR_M": "M€", "EUR_BN": "bn€", "USD": "$", "USD_K": "k$", "USD_M": "M$", "USD_BN": "bn$",
             "GBP": "£", "GBP_K": "k£", "GBP_M": "M£", "PCT": "%", "PP": "pp", "BPS": "bps", "PLAIN_K": "k", "PLAIN_M": "M"}


def _with_unit(x: dict, by_fact: dict) -> str:
    for v in (by_fact.get(x["fact"]) or {}).get("values") or []:
        if f"{v['value']:g}" == x["raw"]:
            u = UNIT_TEXT.get(str(v.get("unit") or "").upper(), "")
            return f"{x['raw']}{u}" if u in ("%",) else (f"{x['raw']} {u}" if u else x["raw"])
    return x["raw"]


def update_plan(inv: dict, facts: list[dict]) -> dict:
    """Every number of the old deck against the new fact model. A number is CURRENT when a new fact
    holds it; OUTDATED when a new fact about the same thing (shared words with its context) holds a
    different value; UNTRACED when nothing in the new sources speaks to it."""
    from ..reasoning.grounding import fact_scalars

    scal = fact_scalars(facts)
    by_fact = {f["id"]: f for f in facts}
    out = []
    for s in inv["slides"]:
        items = []
        for q in s["numbers"]:
            tol = max(1e-9, 0.005 * abs(q["value"]))
            same = [x for x in scal if abs(x["value"] - q["value"]) <= tol and (q["kind"] in ("data", "plain") or x["kind"] == q["kind"])]
            ctx = _words(q["context"])
            if same:
                items.append({**q, "status": "current", "facts": sorted({x["fact"] for x in same})[:3]})
                continue
            cands = []
            for x in scal:
                if q["kind"] not in ("data", "plain") and x["kind"] != q["kind"]:
                    continue
                f = by_fact.get(x["fact"]) or {}
                overlap = len(ctx & _words(f.get("claim", "")))
                # a money or % number needs one shared measure word with a fact of the same kind;
                # a bare data value (a chart point, a table cell) needs two (its row and its series)
                if overlap >= (2 if q["kind"] in ("data", "plain") else 1):
                    cands.append((overlap, x))
            if cands:
                best = max(cands, key=lambda c: c[0])[1]
                items.append({**q, "status": "outdated", "new_value": _with_unit(best, by_fact), "fact": best["fact"],
                              "new_claim": (by_fact.get(best["fact"]) or {}).get("claim", "")[:140]})
            else:
                items.append({**q, "status": "untraced"})
        st = Counter(i["status"] for i in items)
        action = "keep" if not items or st["current"] == len(items) else ("update" if st["outdated"] else "review")
        out.append({"slide": s["n"], "role": s["role"], "headline": s["headline"], "action": action, "counts": dict(st), "numbers": items})
    tot = Counter(i["status"] for s in out for i in s["numbers"])
    return {"source": inv["source"], "slides": out, "totals": dict(tot),
            "rule": "an old number is reused only when the new fact model confirms it; outdated numbers are replaced from the cited new fact"}


def plan_markdown(plan: dict) -> str:
    L = [f"# Update plan — {plan['source']}", "", f"Numbers in the old deck: {plan['totals']}", "", "| slide | role | action | current | outdated | untraced | headline |",
         "|---|---|---|---|---|---|---|"]
    for s in plan["slides"]:
        c = s["counts"]
        L.append(f"| {s['slide']} | {s['role']} | **{s['action']}** | {c.get('current', 0)} | {c.get('outdated', 0)} | {c.get('untraced', 0)} | {s['headline'][:70]} |")
    L += ["", "## Outdated numbers", ""]
    for s in plan["slides"]:
        for q in s["numbers"]:
            if q["status"] == "outdated":
                L.append(f"- slide {s['slide']} {q['where']}: **{q['raw']} → {q['new_value']}** ({q['fact']}: {q['new_claim']})")
    return "\n".join(L) + "\n"


def write_ingest(pptx: str | Path, work: str | Path) -> dict:
    work = Path(work)
    work.mkdir(parents=True, exist_ok=True)
    inv = ingest_deck(pptx)
    (work / "old_deck.json").write_text(json.dumps(inv, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8", newline="\n")
    (work / "old_deck_spec.json").write_text(json.dumps(to_spec(inv), indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8", newline="\n")
    ghost = ["# Old deck — the argument from its headlines", ""] + [f"{s['n']}. [{s['role']}] {s['headline']}" for s in inv["slides"]]
    (work / "old_ghost.md").write_text("\n".join(ghost) + "\n", encoding="utf-8", newline="\n")
    return inv


def write_plan(work: str | Path) -> dict:
    work = Path(work)
    inv = json.loads((work / "old_deck.json").read_text(encoding="utf-8"))
    facts = json.loads((work / "facts.json").read_text(encoding="utf-8")).get("facts", [])
    plan = update_plan(inv, facts)
    (work / "update_plan.json").write_text(json.dumps(plan, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
    (work / "update_plan.md").write_text(plan_markdown(plan), encoding="utf-8", newline="\n")
    return plan
