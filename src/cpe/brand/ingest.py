"""Brand ingestion: corporate .pptx template → engine theme + compatibility report.

    cpe brand ingest template.pptx -o brands/acme

produces
    brands/acme/theme.json          engine theme (colours, fonts, grid, reserved areas, template)
    brands/acme/template.pptx       copy of the template (decks are generated on its masters)
    brands/acme/compatibility.md    human report   (+ compatibility.json)

The report makes every approximation explicit so the engine never silently
produces a different deck: fonts detected / available / missing and the
fallback used for measurement and rendering, theme colours and how they were
mapped to roles, slide size, master layouts and recognised placeholders, logos
and other master artwork (turned into reserved areas the QA protects), and
elements the engine does not use.
"""
from __future__ import annotations

import json
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from ..design import text_metrics as tm
from ..design.tokens import SLIDE_H, SLIDE_W, contrast_ratio, hex_to_rgb, interpolate, load_theme

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
SCHEME_SLOTS = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]
IN = 914400


def _in(v) -> float:
    return round(Emu(v).inches, 3)


def read_theme(prs) -> list[dict]:
    """Theme (colours, fonts) of EVERY slide master."""
    from .model import theme_of_master

    return [theme_of_master(m) for m in prs.slide_masters]


def installed_font_families() -> set[str]:
    try:
        out = subprocess.run(["fc-list", ":", "family"], capture_output=True, text=True, timeout=20, encoding="utf-8", errors="replace").stdout
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return set()
    fams = set()
    for line in out.splitlines():
        for f in line.split(","):
            fams.add(f.strip().lower())
    return fams


def fc_match(name: str) -> str:
    try:
        r = subprocess.run(["fc-match", "-f", "%{family}", name], capture_output=True, text=True, timeout=10, encoding="utf-8", errors="replace")
        return r.stdout.split(",")[0].strip() or "unknown"
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return "unknown (fontconfig not available)"


def font_report(name: str | None, installed: set[str]) -> dict:
    if not name:
        return {"font": None}
    low = name.lower()
    fam = tm.FONT_ALIASES.get(low)
    installed_ok = low in installed
    metric = low in tm.METRIC_COMPATIBLE
    measure = fam if fam and tm.family_available(fam) else "LiberationSans"
    if not installed_ok and not metric:
        # the deck will be RENDERED here with the fontconfig substitute: measure with it when we can
        sub = tm.FONT_ALIASES.get(fc_match(name).lower())
        if sub and tm.family_available(sub):
            measure = sub
    if installed_ok and tm.installed_file(name):  # the real font is here: measure with it
        measure = tm.family_for(name)
    exact_measure = (installed_ok and bool(tm.installed_file(name))) or (bool(fam) and tm.family_available(fam) and metric)
    # what the renderer will actually draw with when the font itself is missing
    render_fallback = None if installed_ok else fc_match(name)
    status = "available" if installed_ok else ("metric-compatible substitute" if metric and fam and tm.family_available(fam) else "missing")
    return {
        "font": name,
        "installed": installed_ok,
        "status": status,
        "measure_family": measure,
        "measurement": "exact" if exact_measure else "approximate",
        "render_fallback": render_fallback,
    }


def _saturation(h: str) -> float:
    r, g, b = [c / 255 for c in hex_to_rgb(h)]
    mx, mn = max(r, g, b), min(r, g, b)
    return 0 if mx == 0 else (mx - mn) / mx


def _lum(h: str) -> float:
    r, g, b = hex_to_rgb(h)
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255


def _hue(h: str) -> float:
    import colorsys

    r, g, b = [c / 255 for c in hex_to_rgb(h)]
    return colorsys.rgb_to_hsv(r, g, b)[0] * 360


def map_colors(scheme: dict, base) -> tuple[dict, dict, list[str]]:
    """Theme scheme → engine roles. Returns (colors, series, notes)."""
    notes = []
    text = scheme.get("dk1", base.c("text"))
    bg = scheme.get("lt1", "FFFFFF")
    accents = [scheme[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6") if k in scheme]
    dark_candidates = [c for c in [scheme.get("dk2")] + accents if c and _lum(c) < 0.45 and contrast_ratio(c, bg) >= 4.5]
    primary = dark_candidates[0] if dark_candidates else base.c("primary")
    if not dark_candidates:
        notes.append("No dark brand colour with ≥4.5:1 contrast on the background: primary kept from the default theme.")
    def signal(c):  # reds and greens are reserved for negative / positive
        h = _hue(c)
        return (h <= 15 or h >= 340) or (90 <= h <= 170)

    vivid = sorted([c for c in accents if c != primary and _saturation(c) > 0.35], key=lambda c: (signal(c), -_saturation(c)))
    highlight = vivid[0] if vivid else base.c("highlight")
    secondary = next((c for c in accents if c not in (primary, highlight) and 0.2 < _lum(c) < 0.65), interpolate(primary, "FFFFFF", 0.35))
    greens = [c for c in accents if 90 <= _hue(c) <= 170 and _saturation(c) > 0.3]
    reds = [c for c in accents if (_hue(c) <= 15 or _hue(c) >= 340) and _saturation(c) > 0.35]
    colors = dict(base.colors)
    colors.update({
        "primary": primary, "secondary": secondary, "highlight": highlight,
        "text": text if contrast_ratio(text, bg) >= 7 else base.c("text"),
        "background": bg,
        "positive": greens[0] if greens else base.c("positive"),
        "negative": reds[0] if reds else base.c("negative"),
        "muted": interpolate(text, bg, 0.78), "faint": interpolate(text, bg, 0.93), "surface": interpolate(text, bg, 0.95),
        "rule": interpolate(text, bg, 0.68), "gridline": interpolate(text, bg, 0.88), "neutral": interpolate(text, bg, 0.5),
    })
    tm_ = interpolate(text, bg, 0.35)
    for _ in range(30):  # bounded: with unconventional slots the text/background pair may never reach it
        if contrast_ratio(tm_, bg) >= 4.6:
            break
        tm_ = interpolate(tm_, text, 0.2)
    colors["text_muted"] = tm_
    if contrast_ratio(highlight, bg) < 3:
        notes.append(f"Highlight #{highlight} has contrast {contrast_ratio(highlight, bg):.1f}:1 — used for fills only; text in it may be flagged by QA.")
    series = [primary, secondary, interpolate(primary, bg, 0.55), colors["muted"], highlight, colors["neutral"]]
    return colors, series, notes


def _shape_kind(sh) -> str:
    if sh.is_placeholder:
        return "placeholder"
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if getattr(sh, "has_chart", False) and sh.has_chart:
        return "chart"
    if sh.shape_type in (MSO_SHAPE_TYPE.TABLE, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.MEDIA) or sh._element.tag.endswith("graphicFrame"):
        return "graphic_frame"
    if sh.has_text_frame and sh.text_frame.text.strip():
        return "text"
    return "shape"


def _box(sh) -> dict:
    return {"x": _in(sh.left or 0), "y": _in(sh.top or 0), "w": _in(sh.width or 0), "h": _in(sh.height or 0)}


def _bg_kind(el) -> str | None:
    bg = el.find(f"{P}cSld/{P}bg")
    if bg is None:
        return None
    if bg.find(f".//{A}blipFill") is not None:
        return "picture background"
    if bg.find(f".//{A}gradFill") is not None:
        return "gradient background"
    return "solid background"


def read_structure(prs) -> dict:
    """Masters, layouts, placeholders and artwork of EVERY master (layouts keep their master's artwork)."""
    out = {"masters": [], "layouts": [], "unsupported": []}
    for mi, master in enumerate(prs.slide_masters):
        _read_master(master, mi, out)
    out["master"] = out["masters"][0] if out["masters"] else {"artwork": [], "background": None, "placeholders": []}
    return out


def _read_master(master, mi: int, out: dict) -> None:
    mrec = {"id": f"m{mi + 1}", "name": master.name, "artwork": [], "background": _bg_kind(master._element), "placeholders": []}
    out["masters"].append(mrec)
    for sh in master.shapes:
        k = _shape_kind(sh)
        if k == "placeholder":
            mrec["placeholders"].append({"type": str(sh.placeholder_format.type).split(".")[-1].split(" ")[0], "name": sh.name, **_box(sh)})
        else:
            mrec["artwork"].append({"kind": k, "name": sh.name, **_box(sh)})
    if mrec["background"] in ("picture background", "gradient background"):
        out["unsupported"].append(f"Master '{master.name}' has a {mrec['background']}: kept from the template, but contrast QA assumes a plain background.")
    for li, lay in enumerate(master.slide_layouts):
        phs, art = [], []
        for sh in lay.shapes:
            if sh.is_placeholder:
                phs.append({"type": str(sh.placeholder_format.type).split(".")[-1].split(" ")[0], "idx": sh.placeholder_format.idx, "name": sh.name, **_box(sh)})
            else:
                art.append({"kind": _shape_kind(sh), "name": sh.name, **_box(sh)})
        bgk = _bg_kind(lay._element)
        out["layouts"].append({"name": lay.name, "id": f"{mrec['id']}.l{li + 1}", "master": mrec["id"], "master_artwork": mrec["artwork"],
                               "placeholders": phs, "artwork": art, "background": bgk,
                               "shows_master_artwork": lay._element.get("showMasterSp", "1") != "0"})
        if bgk in ("picture background", "gradient background"):
            out["unsupported"].append(f"Layout '{lay.name}' has a {bgk} (not used as the base layout).")
        for a in art:
            if a["kind"] in ("chart", "graphic_frame", "group"):
                out["unsupported"].append(f"Layout '{lay.name}' contains a {a['kind']} ('{a['name']}'): treated as artwork, not content.")


RECOGNISED = {"TITLE": "headline", "CENTER_TITLE": "cover title", "SUBTITLE": "cover subtitle", "BODY": "body text", "OBJECT": "content",
              "DATE": "date", "FOOTER": "footer", "SLIDE_NUMBER": "slide number", "PICTURE": "picture", "CHART": "chart", "TABLE": "table"}


def choose_base_layout(struct: dict, model: dict | None = None) -> dict:
    """Base layout for engine-drawn slides (fallback mode): across ALL masters, the emptiest layout
    that shows its master's artwork, on a light background, preferring a recognised content layout."""
    def content_phs(lay):
        return [p for p in lay["placeholders"] if p["type"] not in ("DATE", "FOOTER", "SLIDE_NUMBER")]

    feats = {c["layout_id"]: c for c in (model or {}).get("layouts", [])}

    def light(lay):
        f = (feats.get(lay.get("id")) or {}).get("features", {}).get("background", {})
        return 0 if not (f.get("dark") or f.get("picture")) else 1

    def content_rank(lay):
        cl = (feats.get(lay.get("id")) or {}).get("classification") or []
        return 0 if cl and cl[0]["type"] in ("content", "one_column") else 1

    cands = sorted(struct["layouts"], key=lambda l: (light(l), len(content_phs(l)) > 1, 0 if "blank" in l["name"].lower() else 1, content_rank(l),
                                                     len(content_phs(l)), 0 if l["shows_master_artwork"] else 1, 0 if not l["background"] else 1))
    return cands[0] if cands else {"name": None}


def derive_grid(struct: dict, sw: float, sh: float, model: dict | None = None, k: float = 1.0) -> tuple[dict, list[str]]:
    """Margins from the inferred grid when the example slides support it, else from the title placeholder."""
    notes = []
    grid = {}
    g = (model or {}).get("grid") or {}
    if g.get("confidence", 0) >= 0.6 and g.get("edges_sampled", 0) >= 40 and abs(sw / sh - SLIDE_W / SLIDE_H) < 0.02:
        ml, mr = min(1.2, max(0.25, g["margin_left_in"] * k)), min(1.2, max(0.25, g["margin_right_in"] * k))
        grid.update(margin_l=round(ml, 3), margin_r=round(mr, 3))
        notes.append(f"Margins from the inferred {g['columns']}-column grid (explains {round(100 * g['edges_explained'])}% of {g['edges_sampled']} shape edges): "
                     f"left {ml:.2f} in, right {mr:.2f} in (engine canvas).")
        return grid, notes
    title = None
    for lay in struct["layouts"]:
        t = next((p for p in lay["placeholders"] if p["type"] == "TITLE"), None)
        if t:
            title = t
            break
    if title is None:
        title = next((p for p in struct["master"]["placeholders"] if p["type"] == "TITLE"), None)
    if title and abs(sw / sh - SLIDE_W / SLIDE_H) < 0.02:
        ml = min(1.2, max(0.3, title["x"] * k))
        mr = min(1.2, max(0.3, (sw - title["x"] - title["w"]) * k))
        grid.update(margin_l=round(ml, 3), margin_r=round(mr, 3))
        notes.append(f"Margins taken from the title placeholder: left {ml:.2f} in, right {mr:.2f} in.")
    return grid, notes


def reserved_areas(struct: dict, base: dict, sw: float, sh: float) -> list[dict]:
    """Master/base-layout artwork the content must not cover (logos, bars, marks) — the base layout's OWN master."""
    out = []
    items = (base.get("master_artwork", struct["master"]["artwork"]) if base.get("shows_master_artwork", True) else []) + base.get("artwork", [])
    for a in items:
        full_width = a["w"] > sw * 0.8
        full_height = a["h"] > sh * 0.8
        if full_width and full_height:
            continue  # a full-slide backdrop, not a reserved area
        out.append({"name": a["name"], "kind": a["kind"], "x": a["x"], "y": a["y"], "w": a["w"], "h": a["h"]})
    return out


def lum_hex(h: str) -> float:
    r, g, b = hex_to_rgb(h)
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255


def _scaled(b: dict, k: float) -> dict:
    return {kk: round(v * k, 3) if kk in ("x", "y", "w", "h") else v for kk, v in b.items()}


def corporate_layouts(model: dict, k: float) -> list[dict]:
    """The layout catalogue in ENGINE canvas units, as the builder's matcher needs it."""
    out = []
    ids = {}
    for c in model["layouts"]:
        f = c["features"]
        mi = int(c["master_id"][1:]) - 1
        ids.setdefault(mi, 0)
        li = ids[mi]
        ids[mi] += 1
        out.append({
            "id": c["layout_id"], "name": c["layout"], "master": c["master_id"], "master_index": mi, "layout_index": li,
            "classification": c["classification"], "usage_n": c["evidence"]["example_slides"],
            "placeholders": [_scaled({"type": p["type"], "idx": p.get("idx"), "x": p["x"], "y": p["y"], "w": p["w"], "h": p["h"]}, k) for p in c["_placeholders"]],
            "reserved": [_scaled(a, k) for a in f["reserved_artwork"]],
            "dark_bg": f["background"]["dark"], "coloured_bg": f["background"]["coloured"], "picture_bg": f["background"]["picture"],
            "artwork_in_body": f["artwork_in_body"],
        })
    return out


def font_warning(role: str, font: str | None, installed: set[str], declared: str | None, observed: bool) -> dict:
    rep = font_report(font, installed)
    rep.update(role=role, declared=bool(declared and font and declared.lower() == font.lower()), observed=observed)
    renderable = rep.get("installed") or rep.get("status") == "metric-compatible substitute"
    rep["renderable"] = bool(renderable)
    if font and not rep.get("installed"):
        rep["warning"] = (f"FONT WARNING — corporate font {font} is not available in the authoring environment. "
                          f"LibreOffice renders it with {rep.get('render_fallback')}; text is measured with {rep.get('measure_family')} "
                          f"({rep.get('measurement')}). Expected risk: " + ("none for widths (metric-compatible)." if rep.get("measurement") == "exact"
                                                                            else "line wrapping and box fits may differ in the corporate environment where the font is installed."))
    return rep


def ingest(template: str | Path, out_dir: str | Path, name: str | None = None, base_theme: str = "meridian") -> dict:
    from .model import analyse
    from .rescale import rescale

    template = Path(template)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(template))
    sw, sh = _in(prs.slide_width), _in(prs.slide_height)
    base = load_theme(base_theme)
    installed = installed_font_families()
    model = analyse(prs, installed)
    # keep raw placeholder geometry for the matcher (template units)
    for c, (mi, lay) in zip(model["layouts"], [(mi, lay) for mi, m in enumerate(prs.slide_masters) for lay in m.slide_layouts]):
        c["_placeholders"] = [{"type": str(x.placeholder_format.type).split(".")[-1].split(" ")[0], "idx": x.placeholder_format.idx,
                               "x": _in(x.left), "y": _in(x.top), "w": _in(x.width), "h": _in(x.height)} for x in lay.placeholders]
    notes, unsupported = [], []
    same_size = abs(sw - SLIDE_W) < 0.05 and abs(sh - SLIDE_H) < 0.05
    same_ratio = sh > 0 and abs((sw / sh) / (SLIDE_W / SLIDE_H) - 1) < 0.01
    k = 1.0
    scale_info = None
    if same_size:
        rescale(str(template), str(out / "template.pptx"), target_w_in=sw)  # a copy without the example slides
        use_template = True
    elif same_ratio:
        scale_info = rescale(str(template), str(out / "template.pptx"), target_w_in=SLIDE_W)
        k = scale_info["factor"]
        use_template = True
        notes.append(f"Template canvas {sw}×{sh} in has the engine's 16:9 ratio: masters and layouts were rescaled ×{k} to {SLIDE_W}×{SLIDE_H} in "
                     "(positions, sizes, font sizes, spacing, line widths). PowerPoint's Slide Size dialog scales the output back without distortion.")
    else:
        use_template = False
        unsupported.append(f"Slide size {sw}×{sh} in (ratio {sw / sh if sh else 0:.2f}) is not 16:9: colours and fonts are applied, masters are NOT used.")
    typo = model["typography"]
    roles = typo.get("roles") or {}
    heading = (roles.get("heading") or {}).get("font") or base.font_latin
    body = (roles.get("body") or {}).get("font") or base.font_latin
    fonts = {r: font_warning(r, f, installed, (roles.get(r) or {}).get("declared"), (roles.get(r) or {}).get("source") == "observed")
             for r, f in (("heading", heading), ("body", body))}
    struct = read_structure(prs)
    unsupported += struct["unsupported"]
    base_layout = choose_base_layout(struct, model)
    # colours: theme scheme mapped to roles, corrected by what the slides actually use
    # the scheme of the master that carries engine-drawn slides (not blindly master 0)
    th0 = next((m["theme"] for m in model["masters"] if m["master_id"] == base_layout.get("master")), None) or (model["masters"][0]["theme"] if model["masters"] else {"colors": {}})
    colors, series, color_notes = map_colors(th0["colors"], base)
    pal = model["palette"]
    sources = {k_: "theme" for k_ in ("primary", "highlight", "text")}
    if pal.get("primary") and contrast_ratio(pal["primary"], colors["background"]) >= 1.5:
        if pal["primary"] != colors["highlight"]:
            color_notes.append(f"Highlight set to the colour the slides actually use most (#{pal['primary']}) instead of the theme mapping (#{colors['highlight']}).")
        colors["highlight"] = pal["primary"]
        sources["highlight"] = "observed usage"
        series = [colors["primary"], colors["secondary"], interpolate(colors["primary"], colors["background"], 0.55), colors["muted"], colors["highlight"], colors["neutral"]]
    page, text = pal.get("page"), pal.get("text")
    template_colors = {c for m in model["masters"] for c in m["theme"]["colors"].values()} | set(pal.get("supporting") or []) | set(pal.get("dark") or [])
    if page and text and contrast_ratio(text, page) >= 7 and model["example_slides"]["count"] >= 5:
        # v1.3: roles from what the template really draws (page, text), not from theme slot names —
        # some templates use dk1/lt1 unconventionally. Derived greys come from the real pair, so
        # every text/fill combination keeps its contrast by construction.
        if (page, text) != (colors["background"], colors["text"]):
            color_notes.append(f"Page #{page} and text #{text} taken from the template's content layouts and the text it actually sets "
                               f"(theme slots gave page #{colors['background']}, text #{colors['text']}).")
        colors["background"], colors["text"] = page, text
        sources.update(background="content layouts", text="observed usage (inheritance resolved)")
        dark = (pal.get("dark") or [None])[0]
        if dark and contrast_ratio(dark, page) >= 7:
            colors["primary"] = dark
        elif contrast_ratio(colors["primary"], page) < 4.5 or lum_hex(colors["primary"]) > 0.45 or colors["primary"] not in template_colors:
            colors["primary"] = text  # no dark brand colour of its own: data ink and headers in the text colour
        sources["primary"] = "observed usage"
        sec = next((c for c in pal.get("supporting") or [] if contrast_ratio(c, page) >= 3 and contrast_ratio(c, colors["highlight"]) >= 1.5), None)
        sec = sec or interpolate(colors["primary"], page, 0.45)
        for _ in range(30):  # data fills carry white labels
            if contrast_ratio("FFFFFF", sec) >= 4.6 or contrast_ratio(sec, page) >= 7:
                break
            sec = interpolate(sec, "1A1A1A", 0.15)
        colors["secondary"] = sec
        colors.update({
            "muted": interpolate(text, page, 0.78), "faint": interpolate(text, page, 0.93), "surface": interpolate(text, page, 0.95),
            "rule": interpolate(text, page, 0.68), "gridline": interpolate(text, page, 0.88), "neutral": interpolate(text, page, 0.5),
        })
        tm_ = interpolate(text, page, 0.35)
        for _ in range(30):
            if contrast_ratio(tm_, page) >= 4.6:
                break
            tm_ = interpolate(tm_, text, 0.2)
        colors["text_muted"] = tm_
        series = [colors["primary"], colors["secondary"], interpolate(colors["primary"], page, 0.55), colors["muted"], colors["highlight"], colors["neutral"]]
        if contrast_ratio(colors["highlight"], page) < 3:
            color_notes.append(f"Highlight #{colors['highlight']} has contrast {contrast_ratio(colors['highlight'], page):.1f}:1 on the page: used for fills and marks; "
                               "text set in it is darkened automatically.")
    elif text and contrast_ratio(text, colors["background"]) >= 7 and text != colors["text"]:
        color_notes.append(f"Text colour set to the one used on the slides (#{text}).")
        colors["text"] = text
        sources["text"] = "observed usage"
    pal.pop("_text_counts", None)
    notes += color_notes
    grid, grid_notes = derive_grid(struct, sw, sh, model, k)
    notes += grid_notes
    reserved = [_scaled(r, k) for r in reserved_areas(struct, base_layout, sw, sh)]
    from ..design.tokens import GRID

    ml = grid.get("margin_l", GRID.margin_l)
    for r in reserved:
        if r["y"] < GRID.body_y and r["y"] + r["h"] > GRID.tracker_y and r["x"] > SLIDE_W / 2:
            grid["headline_right_limit"] = round(min(grid.get("headline_right_limit") or SLIDE_W, r["x"] - 0.15), 3)
        if r["y"] + r["h"] > GRID.footer_y and r["x"] > SLIDE_W / 2:
            grid["footer_right_limit"] = round(min(grid.get("footer_right_limit") or SLIDE_W, r["x"] - 0.15), 3)
    if grid.get("headline_right_limit"):
        notes.append(f"Headline width limited to end at {grid['headline_right_limit']} in to keep clear of master artwork.")
    for r in reserved:
        if r["y"] < 6.78 and r["y"] + r["h"] > 1.62 and r["x"] < SLIDE_W - 0.4 and r["x"] + r["w"] > ml + 0.4:
            unsupported.append(f"Master artwork '{r['name']}' sits inside the content area: QA will flag any content that overlaps it.")
    if not use_template:
        pass
    elif not same_size and not scale_info:
        use_template = False
    brand_name = name or template.stem
    corp = corporate_layouts(model, k) if use_template else []
    theme = {
        "name": brand_name,
        "description": f"Brand theme ingested from {template.name}",
        "font_latin": body,
        "font_heading": heading if heading != body else None,
        "font_fallback_file": fonts["body"]["measure_family"],
        "colors": colors,
        "series": series,
        "sequential": [interpolate(colors["primary"], "FFFFFF", t) for t in (0.92, 0.7, 0.45, 0.2)] + [colors["primary"]],
        "diverging": [colors["negative"], interpolate(colors["negative"], "FFFFFF", 0.55), colors["surface"], interpolate(colors["positive"], "FFFFFF", 0.55), colors["positive"]],
        "extras": {
            "template": "template.pptx" if use_template else None,
            "base_layout": base_layout.get("name") if use_template else None,
            "base_layout_id": base_layout.get("id") if use_template else None,
            "grid": grid,
            "reserved": reserved if use_template else [],
            "measure_fonts": {f["font"]: f["measure_family"] for f in fonts.values() if f.get("font")},
            "corporate": {"layouts": corp, "scale": k},
            "brand_rules": {"bookend": bool(model["rules"]["bookend"]["first_and_last_in_brand_colour"]),
                            "brand_colour_slide_share": pal.get("brand_background_share"),
                            "headline_case": model["rules"]["headline_case"]["dominant"]},
        },
    }
    for c in model["layouts"]:
        c.pop("_placeholders", None)
    report = {
        "brand": brand_name,
        "template": template.name,
        "slide_size": {"width_in": sw, "height_in": sh, "supported": use_template, "rescaled": scale_info},
        "masters": [{"id": m["master_id"], "name": m["name"], "layouts": m["layouts"], "theme": m["theme"]["name"], "fonts": m["theme"]["fonts"]} for m in model["masters"]],
        "layout_count": len(model["layouts"]),
        "layout_families": model["layout_families"],
        "example_slides": model["example_slides"],
        "fonts": fonts,
        "typography": {k_: typo[k_] for k_ in ("primary", "confidence", "conflict", "roles", "candidates", "evidence_weights", "sizes")},
        "theme_colors": th0["colors"],
        "palette": {k_: pal[k_] for k_ in ("primary", "text", "supporting", "neutrals", "brand_background_share", "conflict", "confidence")},
        "color_mapping": {k_: colors[k_] for k_ in ("primary", "secondary", "highlight", "text", "text_muted", "background", "positive", "negative")},
        "color_sources": sources,
        "grid": model["grid"],
        "assets": model["assets"],
        "rules": model["rules"],
        "layouts": [{"id": c["layout_id"], "master": c["master_id"], "name": c["layout"], "classification": c["classification"], "usage": c["usage"]} for c in model["layouts"]],
        "base_layout": base_layout.get("name"),
        "reserved_areas": reserved,
        "grid_overrides": grid,
        "unsupported": unsupported,
        "notes": notes,
        "masters_used": use_template,
    }
    (out / "theme.json").write_text(json.dumps(theme, indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "brand_model.json").write_text(json.dumps(model, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    (out / "layout_catalog.json").write_text(json.dumps(model["layouts"], indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    (out / "compatibility.json").write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    (out / "compatibility.md").write_text(to_markdown(report), encoding="utf-8")
    return report


def _bar(conf) -> str:
    if conf is None:
        return "—"
    return "high" if conf >= 0.75 else "medium" if conf >= 0.5 else "low"


def to_markdown(r: dict) -> str:
    ty, pal, g, rules = r["typography"], r["palette"], r["grid"], r["rules"]
    L = [f"# Brand ingest report — {r['brand']}", "", f"Template: `{r['template']}`", ""]
    ok = r["masters_used"] and all(f.get("measurement") == "exact" for f in r["fonts"].values() if f.get("font"))
    L.append(f"**Verdict:** {'✅ fully compatible' if ok else '⚠️ compatible with approximations (see below)'}")
    sz = r["slide_size"]
    L += ["", "## Presentation", "", f"- Canvas: {sz['width_in']} × {sz['height_in']} in"
          + (f" — rescaled ×{sz['rescaled']['factor']} to the engine canvas {sz['rescaled']['to_in'][0]} × {sz['rescaled']['to_in'][1]} in" if sz.get("rescaled") else "")
          + ("" if sz["supported"] else " — masters NOT used"),
          f"- Masters: **{len(r['masters'])}** detected · Layouts: **{r['layout_count']}** · Example slides: **{r['example_slides']['count']}**"]
    for m in r["masters"]:
        L.append(f"  - {m['id']} `{m['name']}` — {m['layouts']} layouts, theme `{m['theme']}` (heading {m['fonts'].get('majorFont')}, body {m['fonts'].get('minorFont')})")
    L += ["", "## Typography", "", f"- Primary observed: **{ty['primary']}** (confidence {ty['confidence']}, {_bar(ty['confidence'])})"]
    for role, x in (ty.get("roles") or {}).items():
        L.append(f"- {role}: **{x['font']}** — {x['source']}" + (f", {round(100 * x['share'])}% of {role} text" if x.get("share") else "") + f"; theme declares {x['declared']}"
                 + (" → **conflict**" if x.get("conflict") else ""))
    if ty.get("conflict"):
        c = ty["conflict"]
        L += ["", f"> **Conflict detected:** declared theme font `{c['declared_theme_font']}` vs observed `{c['observed_primary_font']}`. {c['reason']}"]
    L += ["", "| candidate | score | declared | observed chars | direct formatting | style-guide mentions | installed |", "|---|---|---|---|---|---|---|"]
    for c in ty["candidates"]:
        L.append(f"| {c['font']} | {c['score']} | {'yes' if c['declared'] else 'no'} | {c['observed_chars']} | {c['direct_chars']} | {c['guide_mentions']} | {'yes' if c['installed'] else 'no'} |")
    L += ["", f"Evidence weights: `{json.dumps(ty['evidence_weights'])}` · most used sizes: {ty['sizes']['most_used_pt']} pt", ""]
    L += ["## Fonts in this environment", "", "| role | font | declared | observed | installed | renderable | measured with | measurement | LibreOffice fallback |", "|---|---|---|---|---|---|---|---|---|"]
    for role, f in r["fonts"].items():
        if f.get("font"):
            L.append(f"| {role} | {f['font']} | {'yes' if f.get('declared') else 'no'} | {'yes' if f.get('observed') else 'no'} | {'yes' if f['installed'] else 'no'} | "
                     f"{'yes' if f.get('renderable') else 'substituted'} | {f['measure_family']} | {f['measurement']} | {f['render_fallback'] or '—'} |")
    for w in dict.fromkeys(f["warning"] for f in r["fonts"].values() if f.get("warning")):
        L += ["", f"> ⚠️ {w}"]
    L += ["", "## Colours", "", f"- Primary brand colour (observed): `#{pal['primary']}` · text: `#{pal['text']}` · supporting: "
          + (", ".join(f"`#{c}`" for c in pal["supporting"]) or "—") + f" · neutrals: {', '.join(f'`#{c}`' for c in pal['neutrals']) or '—'}",
          f"- Share of example slides on a brand-colour background: {pal['brand_background_share']}"]
    if pal.get("conflict"):
        L.append(f"- **Conflict:** {pal['conflict']['reason']} (theme accent1 `#{pal['conflict']['theme_accent1']}`, observed `#{pal['conflict']['observed_primary']}`)")
    L += ["", "| engine role | colour | source |", "|---|---|---|"] + [f"| {k} | `#{v}` | {r['color_sources'].get(k, 'theme mapping')} |" for k, v in r["color_mapping"].items()]
    L += ["", "## Grid", "", f"- Likely **{g.get('columns')}-column** system, margins {g.get('margin_left_in')} / {g.get('margin_right_in')} in, gutter {g.get('gutter_in')} in "
          f"(explains {g.get('edges_explained')} of {g.get('edges_sampled')} shape edges; confidence {g.get('confidence')}, {_bar(g.get('confidence'))})",
          f"- Engine grid overrides: `{json.dumps(r['grid_overrides'])}`", ""]
    L += ["## Layout families", ""] + [f"- {k}: {v}" for k, v in r["layout_families"].items()]
    L += ["", "| id | layout | classification (confidence) | observed use on example slides |", "|---|---|---|---|"]
    for l in r["layouts"]:
        L.append(f"| {l['id']} | {l['name']} | " + ", ".join(f"{c['type']} ({c['confidence']})" for c in l["classification"]) + f" | {', '.join(f'{k}×{v}' for k, v in l['usage'].items()) or '—'} |")
    a = r["assets"]
    L += ["", "## Assets", "", f"- Logos: {len(a['logos'])} ({', '.join(f'{k}×{v}' for k, v in a['logo_positions'].items()) or '—'})",
          f"- Icons (distinct small images on slides): {a['icons_distinct']} · pictures on slides: {a['pictures_on_slides']} · vector groups: {a['vector_groups_on_slides']}",
          f"- Reserved artwork on masters/layouts: {a['reserved_artwork']}", f"- Base layout for engine-drawn slides: **{r['base_layout']}**", ""]
    hc = rules["headline_case"]
    L += ["## Inferred brand rules", "",
          f"- Headline case: **{hc['dominant']}** ({json.dumps(hc['shares'])}, {hc['titles_sampled']} titles)",
          f"- Bookend: first and last slides in the brand colour: {rules['bookend']['first_and_last_in_brand_colour']} (same colour: {rules['bookend'].get('first_and_last_share_a_colour')}; {rules['bookend']['first_last_backgrounds']})",
          f"- Slides on a brand-colour background: {rules['brand_colour_slide_share']}",
          f"- Text alignment: {json.dumps(rules['text_alignment'])}",
          f"- Content headline: top at {rules['headline_position']['content_title_top_in']} in, width {rules['headline_position']['content_title_width_share']} of the slide",
          f"- Shapes: rounded share of rectangles {rules['shapes']['rounded_share_of_rectangles']}, chevrons {rules['shapes']['chevrons']}, most used {json.dumps(rules['shapes']['most_used'])}", ""]
    L += ["## Reserved areas (protected by QA)", ""] + ([f"- {x['kind']} `{x['name']}` at ({x['x']}, {x['y']}) {x['w']}×{x['h']} in" for x in r["reserved_areas"]] or ["- none"])
    L += ["", "## Unsupported / not used", ""] + ([f"- {u}" for u in r["unsupported"]] or ["- none"])
    L += ["", "## Notes", ""] + ([f"- {n}" for n in r["notes"]] or ["- none"])
    return "\n".join(L) + "\n"
