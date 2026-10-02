"""Corporate usage and brand fidelity of a generated deck (v1.8).

Two questions are reported separately, and no number here is blended into a single score:

1. **How the template was used, slide by slide** (`corporate_usage`):
   - native: the corporate layout's own placeholders carry the slide;
   - adaptive: a corporate layout gives the frame (background, artwork, title, footer) and the
     engine composes the body inside its free area;
   - engine_fallback: no corporate layout could carry the slide, so the engine layout was used
     on the template's base, with the corporate theme;
   - with the reason for each choice and the layouts that were rejected.
2. **How faithful the result is to the brand** (`brand_fidelity`), one metric per dimension:
   - typography: characters set in the brand fonts, or inheriting them;
   - palette: explicit colours inside the brand palette or neutral greys;
   - grid: shapes inside the brand margins;
   - artwork protection: shapes covering reserved template artwork, such as a logo;
   - structural slides: cover, dividers and closing built on corporate layouts;
   - chart styling: chart series coloured from the brand palette;
   - layout-family appropriateness: needs a human (is this the layout a brand designer
     would pick?), so it is listed as pending, never computed.

Each metric is {value: share in [0, 1] or a count, n: what was measured, misses: examples}.
"""
from __future__ import annotations

import colorsys
import json
from collections import Counter
from pathlib import Path

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
EMU = 914400
STRUCTURAL = ("cover", "divider", "appendix_divider", "closing")


def corporate_usage(manifests: list[dict], resolved: dict | None = None) -> dict:
    kinds = {s.get("id"): s.get("kind", "content") for s in (resolved or {}).get("slides") or []}
    rows = []
    for m in manifests:
        d = m.get("corporate")
        if d is None:
            mode = "no_template"
        else:
            mode = {"native": "native", "adaptive": "adaptive"}.get(d.get("mode"), "engine_fallback")
        rows.append({"slide": m.get("slide_id"), "kind": kinds.get(m.get("slide_id"), ""), "mode": mode,
                     "corporate_layout": (d or {}).get("layout"), "master": (d or {}).get("master"), "engine_layout": m.get("layout"),
                     "why": (d or {}).get("why", ""), "rejected": ((d or {}).get("rejected") or [])[:4]})
    cnt = Counter(r["mode"] for r in rows)
    return {"slides": rows, "counts": dict(cnt), "n": len(rows),
            "corporate_share": round((cnt["native"] + cnt["adaptive"]) / len(rows), 3) if rows else None}


def usage_markdown(u: dict) -> str:
    L = ["# Corporate template usage", "", f"{u['n']} slides: " + ", ".join(f"{k} {v}" for k, v in sorted(u["counts"].items())), "",
         "| slide | kind | mode | corporate layout | engine layout | why |", "|---|---|---|---|---|---|"]
    for r in u["slides"]:
        L.append(f"| {r['slide']} | {r['kind']} | **{r['mode']}** | {r['corporate_layout'] or '–'} | {r['engine_layout']} | {(r['why'] or '')[:110]} |")
    return "\n".join(L) + "\n"


def _neutral(h: str) -> bool:
    """White, black and greys (no hue) are allowed in any brand."""
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    _, _, sat = colorsys.rgb_to_hls(r, g, b)
    return sat < 0.08 or max(r, g, b) - min(r, g, b) < 0.04


def _close(h: str, palette: set[str], tol: int = 10) -> bool:
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    for p in palette:
        pr, pg, pb = (int(p[i:i + 2], 16) for i in (0, 2, 4))
        if abs(r - pr) <= tol and abs(g - pg) <= tol and abs(b - pb) <= tol:
            return True
    return False


def _metric(ok: int, n: int, misses: list) -> dict:
    return {"value": round(ok / n, 3) if n else None, "n": n, "misses": misses[:8]}


def brand_fidelity(pptx: str | Path, theme, manifests: list[dict], issues: list[dict] | None = None, resolved: dict | None = None) -> dict:
    from pptx import Presentation

    prs = Presentation(str(pptx))
    sw = prs.slide_width / EMU
    fonts = {f.lower() for f in theme.fonts() if f}
    palette = {p.upper() for p in theme.palette()}
    grid = (theme.extras.get("grid") or {})
    ml, mr = grid.get("margin_l", 0.55), grid.get("margin_r", 0.55)
    ids = [m.get("slide_id") for m in manifests]

    chars = ok_chars = 0
    font_miss: Counter = Counter()
    cols = ok_cols = 0
    col_miss: Counter = Counter()
    shp = ok_shp = 0
    grid_miss = []
    ser = ok_ser = 0
    ser_miss: Counter = Counter()
    for k, slide in enumerate(prs.slides):
        sid = ids[k] if k < len(ids) else str(k + 1)
        root = slide._element
        for r in root.iter(f"{A}r"):
            t = "".join(x.text or "" for x in r.iter(f"{A}t"))
            n = len(t.strip())
            if not n:
                continue
            latin = r.find(f"{A}rPr/{A}latin")
            face = latin.get("typeface") if latin is not None else None
            chars += n
            if face is None or face.startswith("+") or face.lower() in fonts:
                ok_chars += n
            else:
                font_miss[face] += n
        for tag in (f"{A}srgbClr",):
            for c in root.iter(tag):
                h = (c.get("val") or "").upper()
                if len(h) != 6:
                    continue
                cols += 1
                if h in palette or _close(h, palette) or _neutral(h):
                    ok_cols += 1
                else:
                    col_miss[h] += 1
        for s in slide.shapes:
            if s.left is None or s.width is None or s.is_placeholder:
                continue
            x, w = s.left / EMU, s.width / EMU
            if w >= 0.9 * sw:  # full-bleed backgrounds and bands
                continue
            shp += 1
            if x >= ml - 0.06 and x + w <= sw - mr + 0.06:
                ok_shp += 1
            else:
                grid_miss.append(f"{sid}: {s.name} ({x:.2f}–{x + w:.2f} in)")
            if s.has_chart:
                cx = s.chart.part._element
                for sr in cx.iter(f"{C}ser"):
                    fill = sr.find(f"{C}spPr/{A}solidFill/{A}srgbClr")
                    if fill is None:
                        continue
                    h = fill.get("val").upper()
                    ser += 1
                    if h in palette or _close(h, palette) or _neutral(h):
                        ok_ser += 1
                    else:
                        ser_miss[h] += 1
    usage = corporate_usage(manifests, resolved)
    struct = [r for r in usage["slides"] if r["kind"] in STRUCTURAL]
    struct_ok = [r for r in struct if r["mode"] in ("native", "adaptive")]
    overlaps = [i for i in issues or [] if i.get("code") == "BRAND_RESERVED_OVERLAP"]
    return {
        "typography": _metric(ok_chars, chars, [f"{f} ({n} chars)" for f, n in font_miss.most_common()]),
        "palette": _metric(ok_cols, cols, [f"#{h} ×{n}" for h, n in col_miss.most_common()]),
        "grid": _metric(ok_shp, shp, grid_miss),
        "artwork_protection": {"value": len(overlaps), "n": len(manifests), "misses": [f"{i.get('slide')}: {i.get('message')}" for i in overlaps][:8],
                               "note": "shapes covering reserved template artwork (logos, bands); 0 is the target"},
        "structural_slides": _metric(len(struct_ok), len(struct), [f"{r['slide']} ({r['kind']}): {r['mode']}" for r in struct if r not in struct_ok]),
        "chart_styling": _metric(ok_ser, ser, [f"#{h} ×{n}" for h, n in ser_miss.most_common()]),
        "layout_family_appropriateness": {"value": None, "status": "needs human judgement",
                                          "note": "whether each slide uses the layout a brand designer would pick; see corporate_usage for the choices"},
        "usage": usage["counts"],
        "rule": "metrics are reported separately; none is blended into the deck score",
    }


def fidelity_markdown(f: dict) -> str:
    L = ["# Brand fidelity", "", "| dimension | value | measured | examples of misses |", "|---|---|---|---|"]
    for k in ("typography", "palette", "grid", "artwork_protection", "structural_slides", "chart_styling", "layout_family_appropriateness"):
        m = f[k]
        v = m.get("value")
        vs = "human" if k == "layout_family_appropriateness" else ("–" if v is None else (f"{v:.0%}" if k != "artwork_protection" else f"{v} overlaps"))
        L.append(f"| {k.replace('_', ' ')} | {vs} | {m.get('n', '–')} | {'; '.join(map(str, m.get('misses') or []))[:160]} |")
    L += ["", f"Template usage: {f['usage']}", "", f"_{f['rule']}_"]
    return "\n".join(L) + "\n"


def write_run_reports(out: str | Path, pptx: str | Path, theme, manifests: list[dict], issues: list[dict], resolved: dict) -> dict | None:
    """Written next to the QA report when the deck uses a corporate template."""
    if not theme.extras.get("template"):
        return None
    out = Path(out)
    u = corporate_usage(manifests, resolved)
    f = brand_fidelity(pptx, theme, manifests, issues, resolved)
    (out / "corporate_usage.json").write_text(json.dumps(u, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (out / "corporate_usage.md").write_text(usage_markdown(u), encoding="utf-8")
    (out / "brand_fidelity.json").write_text(json.dumps(f, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (out / "brand_fidelity.md").write_text(fidelity_markdown(f), encoding="utf-8")
    return f


def from_run_dir(run: str | Path, brand: str | None = None) -> dict:
    """Recompute both reports from a finished run folder (deck.pptx, build_manifest.json, resolved.json, qa_report.json)."""
    from ..design.tokens import theme_for

    run = Path(run)
    resolved = json.loads((run / "resolved.json").read_text(encoding="utf-8"))
    manifests = json.loads((run / "build_manifest.json").read_text(encoding="utf-8"))
    manifests = manifests.get("slides", manifests) if isinstance(manifests, dict) else manifests
    qa = json.loads((run / "qa_report.json").read_text(encoding="utf-8")) if (run / "qa_report.json").exists() else {}
    issues = qa.get("issues") or []
    pptx = next(iter(sorted(run.glob("*.pptx"))))
    meta = dict(resolved.get("meta") or {})
    if brand:
        meta["brand"] = brand
    elif meta.get("brand") and not Path(meta["brand"]).exists() and str(meta["brand"]).startswith("/work/"):
        meta["brand"] = str(Path(__file__).resolve().parents[3] / meta["brand"][len("/work/"):])  # rendered in the container
    theme = theme_for(meta)
    return write_run_reports(run, pptx, theme, manifests, issues, resolved) or {"note": "the deck does not use a corporate template"}
