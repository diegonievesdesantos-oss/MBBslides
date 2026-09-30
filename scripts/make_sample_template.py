"""Build a fictitious corporate template (Kestrel Capital) to exercise `cpe brand ingest`.

It deliberately contains what real corporate templates contain: custom theme
colours, a serif heading font that is usually NOT installed (Georgia) with a
Calibri body, a logo and a bottom bar on the master, custom title positions and
a gradient background on one layout.

    python scripts/make_sample_template.py examples/brand/kestrel_template.pptx
"""
import copy
import sys
from io import BytesIO
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
COLORS = {"dk1": "222222", "lt1": "FFFFFF", "dk2": "1B3A2F", "lt2": "EEF2EE", "accent1": "2E6B4F", "accent2": "C8A24A",
          "accent3": "7FA88F", "accent4": "5B6770", "accent5": "A33A2B", "accent6": "3F8F5F", "hlink": "2E6B4F", "folHlink": "5B6770"}


def logo_png() -> bytes:
    im = Image.new("RGBA", (700, 170), (255, 255, 255, 0))
    d = ImageDraw.Draw(im)
    d.polygon([(10, 140), (80, 20), (150, 140)], fill=(27, 58, 47, 255))
    d.polygon([(60, 140), (110, 60), (160, 140)], fill=(200, 162, 74, 255))
    try:
        f = ImageFont.truetype("/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf", 96)
    except OSError:
        f = ImageFont.load_default()
    d.text((180, 22), "KESTREL", fill=(27, 58, 47, 255), font=f)
    buf = BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def main(out: str) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    master = prs.slide_masters[0]
    theme_part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    cs = root.find(f".//{{{A}}}clrScheme")
    cs.set("name", "Kestrel")
    for slot, val in COLORS.items():
        el = cs.find(f"{{{A}}}{slot}")
        for ch in list(el):
            el.remove(ch)
        etree.SubElement(el, f"{{{A}}}srgbClr", {"val": val})
    fs = root.find(f".//{{{A}}}fontScheme")
    fs.find(f"{{{A}}}majorFont/{{{A}}}latin").set("typeface", "Georgia")
    fs.find(f"{{{A}}}minorFont/{{{A}}}latin").set("typeface", "Calibri")
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    # title placeholders on the widescreen canvas: 0.7 in margins
    for container in [master] + list(master.slide_layouts):
        for sh in container.placeholders:
            if sh.placeholder_format.type in (1, 3):  # TITLE, CENTER_TITLE
                sh.left, sh.width = Inches(0.7), Inches(13.333 - 1.4)
            elif sh.placeholder_format.type in (2, 7):  # BODY, OBJECT
                sh.left, sh.width = Inches(0.7), Inches(13.333 - 1.4)
    # master artwork (python-pptx cannot add shapes to masters directly): draw on a scratch
    # slide, then move the XML into the master and re-link the image to the master part.
    scratch = prs.slides.add_slide(master.slide_layouts[6])
    pic = scratch.shapes.add_picture(BytesIO(logo_png()), Inches(11.2), Inches(0.26), Inches(1.45), Inches(0.35))
    pic.name = "Kestrel logo"
    bar = scratch.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.40), prs.slide_width, Inches(0.10))
    bar.name = "Brand bar"
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string("1B3A2F")
    bar.line.fill.background()
    _, rId = master.part.get_or_add_image_part(BytesIO(logo_png()))
    tree = master.shapes._spTree
    pic_el = copy.deepcopy(pic._element)
    pic_el.find(".//{%s}blip" % A).set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rId)
    tree.append(pic_el)
    tree.append(copy.deepcopy(bar._element))
    sldIdLst = prs.slides._sldIdLst
    sid = sldIdLst[-1]
    prs.part.drop_rel(sid.rId)
    sldIdLst.remove(sid)
    # a gradient background on the title layout (common in corporate templates)
    title_layout = master.slide_layouts[0]
    cSld = title_layout._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}cSld")
    bg = etree.fromstring(f'<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{A}"><p:bgPr><a:gradFill><a:gsLst>'
                          f'<a:gs pos="0"><a:srgbClr val="1B3A2F"/></a:gs><a:gs pos="100000"><a:srgbClr val="2E6B4F"/></a:gs></a:gsLst></a:gradFill><a:effectLst/></p:bgPr></p:bg>')
    cSld.insert(0, bg)
    # one sample slide (templates usually ship with one); the engine drops it
    s = prs.slides.add_slide(master.slide_layouts[0])
    s.shapes.title.text = "Kestrel Capital template"
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"template → {out}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "examples/brand/kestrel_template.pptx")
