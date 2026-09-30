"""Builds real decks, re-opens them, and checks that QA catches real defects."""
import json

from pptx import Presentation
from pptx.util import Inches, Pt

from cpe.core.planner import plan
from cpe.design.tokens import load_profile, load_theme
from cpe.pipeline import run
from cpe.pptx.builder import build
from cpe.qa import geometry, render_checks
from cpe.render import renderer
from cpe.spec import VISUAL_TYPES

from conftest import content_slide, mini_spec, needs_render


def _codes(issues):
    return {i["code"] for i in issues}


def test_gallery_builds_every_exhibit_family_and_is_editable(gallery, tmp_path):
    res, _ = plan(gallery)
    out = tmp_path / "g.pptx"
    manifests = build(res, out)
    prs = Presentation(str(out))
    assert len(prs.slides) == len(manifests) == len(res["slides"])
    charts = tables = 0
    for slide in prs.slides:
        for sh in slide.shapes:
            charts += bool(getattr(sh, "has_chart", False) and sh.has_chart)
            tables += bool(getattr(sh, "has_table", False) and sh.has_table)
            assert sh.name.startswith("cpe|"), sh.name  # every shape is traceable
    assert charts >= 8 and tables >= 2  # native, data-editable exhibits
    used = {v.get("chosen") for s in res["slides"] for v in (s.get("_plan") or {}).get("visuals", [])}
    assert {"bar", "line", "scatter", "donut", "bridge", "funnel", "tile_map", "org_chart", "driver_tree", "journey", "architecture", "harvey_table", "scorecard", "timeline", "process"} <= used


def test_no_theme_shadows_and_valid_reopen(alvora, tmp_path):
    res, _ = plan(alvora)
    out = tmp_path / "a.pptx"
    build(res, out)
    prs = Presentation(str(out))
    for slide in prs.slides:
        xml = slide._element.xml
        assert "<p:style>" not in xml  # no inherited theme shadows / outlines
        for sh in slide.shapes:
            if getattr(sh, "has_chart", False) and sh.has_chart:
                assert sh.chart.plots[0].categories is not None


def test_geometry_detects_overflow_collision_offslide_placeholder(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    def tb(name, x, y, w, h, text, size=12):
        t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        t.name = name
        t.text_frame.word_wrap = True
        t.text_frame.text = text
        for p in t.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)
                r.font.name = "Arial"
        return t

    tb("cpe|body|text|1", 1, 2, 2, 0.3, "This sentence is far too long to fit in a small box of two inches wide and a third high")
    tb("cpe|body|text|2", 5, 2, 3, 0.4, "Overlapping text A")
    tb("cpe|body|text|3", 5.2, 2.05, 3, 0.4, "Overlapping text B")
    tb("cpe|body|text|4", 12.8, 3, 2, 0.4, "Off the slide")
    tb("cpe|body|text|5", 1, 4, 3, 0.4, "TODO add source")
    tb("cpe|body|text|6", 1, 5, 3, 0.4, "Tiny", size=6)
    p = tmp_path / "bad.pptx"
    prs.save(str(p))
    man = [{"slide_id": "bad", "layout": "x", "zones": {"body": {"role": "exhibit", "x": 0.55, "y": 1.6, "w": 12.2, "h": 5.2}}}]
    codes = _codes(geometry.check(str(p), man, load_theme(), load_profile("standard")))
    assert {"TEXT_OVERFLOW", "TEXT_COLLISION", "OFF_SLIDE", "PLACEHOLDER_TEXT", "FONT_TOO_SMALL"} <= codes


@needs_render
def test_render_qa_detects_real_spill_and_long_headline(tmp_path):
    long_head = "This headline keeps adding words and qualifiers and caveats about revenue, margins, customers, channels, regions and competitors until it cannot possibly fit on two lines of a slide anymore"
    s = content_slide(headline=long_head, layout="exhibit_commentary_right", commentary={"points": ["Very long bullet with many words " * 8] * 6})
    rep = run(mini_spec([s]), tmp_path, max_iter=1)
    codes = _codes(rep["issues"])
    assert not rep["passed"]
    assert "RENDER_HEADLINE_LINES" in codes
    assert codes & {"RENDER_TEXT_SPILL", "TEXT_OVERFLOW", "CONTENT_OVER_CAPACITY"}
    assert (tmp_path / "qa_report.md").exists() and (tmp_path / "renders").exists()


@needs_render
def test_self_correction_loop_fixes_visual_encoding(tmp_path):
    cats = ["Northern European distribution centres", "Southern European distribution centres", "Central European hubs", "Eastern European hubs", "UK and Ireland stores"]
    s = content_slide(headline="Northern distribution centres cost 40% more than the other regions", message_type="ranking",
                      evidence=[{"claim": "cost", "values": [140, 100, 104, 98, 101]}],
                      visual={"type": "column", "title": "Cost index", "unit": "index", "data": {"categories": cats, "series": [{"name": "Index", "values": [140, 100, 104, 98, 101]}]}})
    rep = run(mini_spec([s]), tmp_path, max_iter=3)
    assert len(rep["iterations"]) >= 2
    fixed = json.loads((tmp_path / "deck.autofixed.json").read_text())
    assert fixed["slides"][0]["visual"]["type"] == "bar"


@needs_render
def test_demo_deck_passes_the_gate(alvora, tmp_path):
    rep = run(alvora, tmp_path, max_iter=2)
    assert rep["passed"], [i for i in rep["issues"] if i["level"] == "error"]
    assert rep["deck_score"] >= 95
    assert len(list((tmp_path / "renders").glob("*.png"))) == 12
