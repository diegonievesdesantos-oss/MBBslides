"""v1.9 — DEBT_V18 fixes (tool findings from the real-project run), on synthetic data."""


def test_u4_u5_week_and_date_references_are_periods():
    from cpe.core.headline import numbers_in
    from cpe.ingest.readers import mask_dates
    assert numbers_in("En régimen (semanas 31 a 40) rinde 1,89x; la mejor semana (S37) 85,7") == [(1.89, "x"), (85.7, "")]
    assert numbers_in("Decidir la firma el 27 de noviembre") == []
    assert "12" not in mask_dates("weeks 12-15 and W40, KW 12")
    assert numbers_in("durante 8 semanas por encima de 2,2x") == [(2.2, "x")]  # a duration before the word is not masked, and is exempt


def test_u6_unit_comes_from_header_or_stated_label_not_quoted_figures():
    from cpe.reasoning.facts import _label_unit_text, detect_unit
    assert detect_unit(_label_unit_text("Real + mermas de 120 k€ (sin fuente)")) == ""
    assert detect_unit(_label_unit_text("FY2027 orgánico (+3,5% sin Marea)")) == ""
    assert detect_unit(_label_unit_text("Ventas netas (M€)")) == "EUR_M"
    assert detect_unit(_label_unit_text("Margen (%)")) == "PCT"


def test_u8_statement_text_is_fact_checked():
    from cpe.reasoning.checks import factcheck_deck
    facts = {"F1": {"id": "F1", "claim": "Ventas 2025: 1.460 M€", "values": [{"value": 1460, "unit": "EUR_M"}]}}
    deck = {"slides": [{"id": "s02", "kind": "exec_summary", "headline": "Las ventas crecen pero el margen cae", "evidence": [{"fact": "F1"}],
                        "visual": {"type": "statements", "data": {"items": [{"title": "Volumen", "text": "Las ventas llegan a 1.460 M€ y el margen a 23%"}]}}}]}
    bad = [i for i in factcheck_deck(deck, facts) if i["code"] == "UNSUPPORTED_NUMBER"]
    assert len(bad) == 1 and "23%" in bad[0]["message"] and "items[0].text" in bad[0]["ref"]


def test_u9_scaled_numbers_ground_against_unscaled_facts():
    from cpe.reasoning.grounding import ground_numbers
    f = [{"id": "F1", "claim": "Líneas FY2027 (miles)", "values": [{"value": 4588, "unit": ""}]},
         {"id": "F2", "claim": "Capacidad (millones de líneas)", "values": [{"value": 4.6, "unit": ""}]}]
    st = {g["number"]: g["status"] for g in ground_numbers("4,59 M de líneas frente a 4,6 M de capacidad", f)}
    assert st == {"4,59 M": "grounded", "4,6 M": "grounded"}
    assert ground_numbers("unos 7,3 M de líneas", f)[0]["status"] == "unsupported"


def test_u2_banner_line_and_repeated_two_level_header(tmp_path):
    import json
    import subprocess
    import sys
    src = tmp_path / "src"
    src.mkdir()
    (src / "presupuesto.csv").write_text(
        "Presupuesto almacén · versión 2 (borrador 25/09/2026);;;;;\n"
        "Concepto;FY2025;FY2026;FY2026;FY2027;Comentario\n"
        ";Real;Presupuesto;Real;Presupuesto;\n"
        "Volumen;;;;;\n"
        "Líneas preparadas (miles);4.396;4.660;4.373;4.723;+8% s/ real FY2026\n"
        "Ventas netas (M€);94,5;100,1;94,0;101,6;\n"
        "Mantenimiento;n/d;75;33;133;Contrato\n", encoding="utf-8")
    subprocess.run([sys.executable, "-m", "cpe", "reason", "facts", str(src), "-o", str(tmp_path / "w")], check=True, capture_output=True)
    fs = json.loads((tmp_path / "w" / "facts.json").read_text(encoding="utf-8"))["facts"]
    by = {(f["values"][0]["label"].split(" · ")[0], f["values"][0]["period"], f["values"][0]["basis"]): f["values"][0] for f in fs if f.get("fact_type") == "table_value"}
    assert by[("Líneas preparadas (miles)", "FY2025", "actual")]["value"] == 4396 and by[("Líneas preparadas (miles)", "FY2025", "actual")]["unit"] == ""
    assert by[("Líneas preparadas (miles)", "FY2027", "budget")]["value"] == 4723
    assert by[("Ventas netas (M€)", "FY2025", "actual")] == {**by[("Ventas netas (M€)", "FY2025", "actual")], "value": 94.5, "unit": "EUR_M"}
    assert by[("Mantenimiento", "FY2026", "actual")]["value"] == 33


def test_u3_conflicts_read_each_number_with_its_own_words():
    from cpe.reasoning.conflicts import detect_conflicts

    def tx(i, file, claim, value, unit, period=None):
        return {"id": i, "fact_type": "text_statement", "claim": claim, "source": {"file": file, "loc": "line 1"},
                "values": [{"value": value, "unit": unit, "period": period}]}

    def tv(i, file, label, col, value, unit, period=None):
        return {"id": i, "fact_type": "table_value", "claim": f"{label} — {col}", "source": {"file": file, "loc": "sheet A"},
                "values": [{"value": value, "unit": unit, "period": period, "label": label, "column": col}]}

    real = [tx("F1", "informe.md", "La productividad de la zona automatizada ya es 2,1 veces la productividad manual.", 2.1, "X"),
            tx("F2", "oferta.md", "Productividad de preparación en puestos GTP: 2,4 veces la productividad manual de referencia.", 2.4, "X"),
            tx("F3", "ops.md", "El coste de cierre de Valencia (indemnizaciones, penalización del contrato de alquiler y desmantelamiento) asciende a 3,2 millones de euros.", 3.2, "EUR_M"),
            tx("F4", "propuesta.md", "El coste de cierre estimado es de 1,5 M€.", 1.5, "EUR_M")]
    assert len(detect_conflicts(real)) == 2
    noise = [  # a comparison base is not the number's period; phases, zones and mixed manual/automated wording are not one quantity
        tx("F1", "informe.md", "FY2026 cierra con 4,37 millones de líneas preparadas, un 0,5% menos que en FY2025.", 0.5, "PCT", "FY2026"),
        tv("F2", "presupuesto.csv", "Líneas preparadas (miles)", "Comentario: +8% s/ real FY2026", 8.0, "PCT", "FY2026"),
        tv("F3", "modelo.xlsx", "resultado · Ahorro neto anual · k€/año", "fase_2", 643, "EUR_K"),
        tv("F4", "analisis.csv", "Ahorro neto anual de la fase 1 en el business case", "importe (k€)", 732, "EUR_K"),
        tv("F5", "deck.pptx", "Tasa de error manual / automatizada", "Valor: 0,9% / 0,2%", 0.2, "PCT"),
        tv("F6", "wms.csv", "Tasa de error de la zona frío manual (%)", "valor", 0.91, "PCT"),
        tv("F7", "ventas.csv", "Líneas preparadas (miles) · FY2025 · Real", "valor", 4396, ""),
        tv("F8", "wms.csv", "Líneas de la zona frío en FY2026 (miles)", "valor", 1927, "")]
    assert detect_conflicts(noise) == []


def test_u1_stale_matches_on_words_unit_and_period_not_on_value():
    from cpe.reasoning.deck_update import _kpi_labels_before, _pair_label, update_plan

    def tx(i, file, claim, value, unit, period=None):
        return {"id": i, "fact_type": "text_statement", "claim": claim, "source": {"file": file}, "values": [{"value": value, "unit": unit, "period": period}]}

    def tv(i, file, label, col, value, unit=""):
        return {"id": i, "fact_type": "table_value", "claim": f"{label} — {col}", "source": {"file": file},
                "values": [{"value": value, "unit": unit, "label": label, "column": col}]}

    def num(raw, value, kind, context, text=None, unit_after=""):
        return {"where": "body[0]", "raw": raw, "value": value, "kind": kind, "context": context, "text": text or context, "unit_after": unit_after}

    inv = {"source": "old.pptx", "slides": [{"n": 2, "role": "content", "headline": "Business case", "numbers": [
        num("26%", 26.0, "pct", "26% TIR a 10 años"),                                     # equal value elsewhere (a cost overrun %): not "current"
        num("732", 732.0, "plain", "Ahorro neto anual Fase 1 · Ambiente"),                 # the business case repeats it; the analysis does not
        num("2,4x", 2.4, "x", "2,4x productividad de picking"),
        num("€38.500", 38500.0, "money", "38.500 € Coste empresa por FTE"),
        num("1.640", 1640.0, "plain", "1.640 h/FTE; coste por", unit_after="h"),           # no new fact about hours per FTE
        num("118", 118.0, "plain", "oferta K-ES-2025-118 rev. B"),                         # a document code
        num("2", 2.0, "plain", "2", text="2")]}]}                                          # the page number
    facts = [tv("F1", "analysis/desviacion_fase1.csv", "Coste completo frente a los 3.200 k€ aprobados", "desviacion_pct", 26, "PCT"),
             tv("F2", "analysis/ahorro_fase1.csv", "Ahorro neto anual de la fase 1 en el business case", "importe (k€)", 732, "EUR_K"),
             tv("F3", "analysis/ahorro_fase1.csv", "Ahorro neto anual a la productividad real", "importe (k€)", 451, "EUR_K"),
             tv("F4", "analysis/productividad.csv", "Ambiente automatizada en régimen", "productividad_veces", 1.89),
             tx("F5", "nota_rrhh.md", "El coste empresa medio por FTE de almacén queda en 41.400 € en 2026.", 41400, "EUR", "2026"),
             tx("F6", "mayor.csv", "Asiento 26078 · Kinetra — importe", 2.4, "")]
    st = {q["raw"]: q for q in update_plan(inv, facts)["slides"][0]["numbers"]}
    assert st["26%"]["status"] == "untraced"
    assert st["732"]["status"] == "outdated" and st["732"]["fact"] == "F3"
    assert st["2,4x"]["status"] == "outdated" and st["2,4x"]["fact"] == "F4"
    assert st["€38.500"]["status"] == "outdated" and st["€38.500"]["fact"] == "F5"
    assert st["1.640"]["status"] == "untraced"
    assert st["118"]["status"] == "ignored" and st["2"]["status"] == "ignored"
    assert _pair_label("Tasa de error manual / automatizada", 1) == "Tasa de error automatizada"
    assert _kpi_labels_before(["Plantilla almacén", "142 FTE"]) and not _kpi_labels_before(["5,6 M€", "Inversión total"])


def _textbox_deck(path):
    """A deck built with free text boxes on placeholder-less layouts (like many real client decks)."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(s, x, y, w, h, t, size, bold=False):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.text_frame.text = t
        r = tb.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, RGBColor(0x1F, 0x3A, 0x5F)

    s = prs.slides.add_slide(blank)
    text(s, 0.6, 2.4, 11, 1.0, "Plan de expansión de la red de tiendas", 40, True)
    text(s, 0.6, 3.4, 11, 0.6, "Comité de dirección", 22)
    for i in range(4):
        s = prs.slides.add_slide(blank)
        text(s, 0.6, 0.35, 12.1, 0.9, f"Las ventas por tienda crecieron un {i + 3}% en el último año", 26, True)
        text(s, 0.6, 1.8, 6.0, 4.0, "Detalle del análisis por región y formato de tienda", 14)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.3), Inches(12.1), Inches(0.05))  # repeated on every page
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(0xE0, 0x7A, 0x1F)
        text(s, 0.6, 7.0, 6, 0.3, "Comité de dirección · 12/03/2025 · Confidencial", 9)
    prs.save(str(path))


def test_d2_layouts_learned_from_slide_geometry(tmp_path):
    from pptx import Presentation

    from cpe.brand.ingest import ingest
    from cpe.brand.learned import needs_learning

    _textbox_deck(tmp_path / "deck.pptx")
    assert needs_learning(Presentation(str(tmp_path / "deck.pptx")))
    rep = ingest(tmp_path / "deck.pptx", tmp_path / "brand")
    learned = {x["kind"]: x for x in rep["learned_layouts"]}
    assert set(learned) == {"cover", "content"}
    assert learned["content"]["title"]["y"] == 0.35 and learned["content"]["title"]["size"] == 26.0 and learned["content"]["repeated_shapes"] >= 2
    assert "subtitle" in learned["cover"] and learned["content"]["slides"] == [2, 3, 4, 5]
    cls = {x["name"]: x["classification"][0]["type"] for x in rep["layouts"]}
    assert cls["Learned · Cover"] == "cover" and cls["Learned · Title and content"] in ("one_column", "content")
    assert any("12/03/2025" in u for u in rep["unsupported"])  # a dated footer on the layout is flagged
    # a template that already has title placeholders is left alone
    assert not needs_learning(Presentation())


def test_d2_footer_artwork_moves_the_source_line_up():
    from cpe.brand.matching import band_conflict, footer_shift, limits
    from cpe.design.tokens import GRID

    lay = {"name": "L", "reserved": [{"name": "footer text", "x": 0.6, "y": 7.0, "w": 6.0, "h": 0.3}, {"name": "logo", "x": 11.6, "y": 7.0, "w": 1.3, "h": 0.3}]}
    shift = footer_shift(lay, GRID)
    assert shift is not None and shift + GRID.footer_h <= 7.0 and band_conflict(lay, GRID) is None
    lim = limits(lay, GRID)
    assert lim["footer_y"] == shift and lim["body_bottom"] < shift and "footer_right_limit" not in lim  # the logo is now below the source line
    high = {"name": "H", "reserved": [{"name": "band", "x": 0.6, "y": 5.4, "w": 12.0, "h": 2.0}]}  # would leave too little body
    assert footer_shift(high, GRID) is None and band_conflict(high, GRID)


def test_d1_patch_changes_only_approved_numbers_in_the_original_pptx(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    from cpe.reasoning.deck_patch import apply_edits, proposed_edits
    from cpe.reasoning.deck_update import format_like, ingest_deck, update_plan

    prs = Presentation()
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tb.text_frame.text = "La inversión de la fase 1 es de 3,2 M€ y el ahorro neto anual de 732 k€"
    r = tb.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold = Pt(24), True
    rows = s.shapes.add_table(3, 2, Inches(0.5), Inches(2), Inches(6), Inches(1.2)).table
    for i, (a, b) in enumerate([("Concepto", "Valor"), ("Mantenimiento anual (k€)", "-150"), ("Rango de temperatura", "2-4 °C")]):
        rows.cell(i, 0).text, rows.cell(i, 1).text = a, b
    cd = CategoryChartData()
    cd.categories = ["2025", "2026"]
    cd.add_series("Líneas (M)", (4.4, 4.6))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.8), Inches(2), Inches(3), Inches(3), cd)
    prs.slides.add_slide(blank).shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "Caja acumulada"
    prs.save(str(tmp_path / "old.pptx"))

    inv = ingest_deck(tmp_path / "old.pptx")
    facts = [{"id": "F1", "fact_type": "table_value", "claim": "x", "source": {"file": "analysis/coste_fase1.csv"},
              "values": [{"value": 4034, "unit": "EUR_K", "label": "Inversión completa de la fase 1", "column": "importe (k€)"}]}]
    e = proposed_edits(update_plan(inv, facts))
    prop = next(x for x in e["edits"] if x["old"] == "€3,2 M")
    assert prop["replace"] == "4,03" and prop["approved"] is False
    prop["approved"] = True
    e["edits"] += [{"op": "number", "slide": 1, "where": "exhibit[0].rows[0][1]", "find": "150", "replace": "-210", "approved": True},
                   {"op": "number", "slide": 1, "where": "exhibit[0].rows[1][1]", "find": "4", "replace": "5", "approved": False},  # not approved
                   {"op": "number", "slide": 1, "where": "exhibit[1].Líneas (M)[2026]", "find": "4.6", "replace": "4,37", "approved": True},
                   {"op": "number", "slide": 1, "where": "body[9]", "find": "1", "replace": "2", "approved": True},  # no such text
                   {"op": "delete_slide", "slide": 2}]
    rep = apply_edits(tmp_path / "old.pptx", e, tmp_path / "new.pptx", mark=True)
    assert len(rep["applied"]) == 4 and len(rep["failed"]) == 1 and rep["skipped_unapproved"] == 1 and rep["slides_deleted"] == [2]
    new = Presentation(str(tmp_path / "new.pptx"))
    assert len(new.slides) == 1
    s = new.slides[0]
    title = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("La inversión"))
    assert title.text_frame.text == "La inversión de la fase 1 es de 4,03 M€ y el ahorro neto anual de 732 k€"
    runs = title.text_frame.paragraphs[0].runs
    assert [x.text for x in runs] == ["La inversión de la fase 1 es de ", "4,03", " M€ y el ahorro neto anual de 732 k€"]
    assert all(x.font.bold and x.font.size == Pt(24) for x in runs)  # formatting kept on every piece
    assert runs[1]._r.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}highlight") is not None
    tbl = next(sh for sh in s.shapes if sh.has_table).table
    assert tbl.cell(1, 1).text == "-210" and tbl.cell(2, 1).text == "2-4 °C"
    ch = next(sh for sh in s.shapes if sh.has_chart).chart
    assert list(ch.plots[0].series[0].values) == [4.4, 4.37]
    assert "3,2 M€ → 4,03" in s.notes_slide.notes_text_frame.text or "€3,2 M → 4,03" in s.notes_slide.notes_text_frame.text
    assert format_like("38.500", 41400) == "41.400" and format_like("1,374.5", 1200.25).startswith("1,200")
