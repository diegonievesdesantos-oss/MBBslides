"""Command line: `python -m cpe <command>` (or `scripts/cpe <command>`).

Thinking commands (no drawing):
  ingest     files → inventory.json + inventory.md (blocks, tables, facts with context)
  scaffold   deck type → skeleton deck.json with storyline and slide intents to fill
  outline    deck.json → ghost deck (headline-only read-through)
  lint       deck.json → structure, storyline, headline, visual and density issues
  recommend  message type (+ exhibit) → ranked visual encodings with reasons
  plan       deck.json → resolved.json (visual + layout + fit decisions, with rationale)
Rendering commands:
  build      deck.json → deck.pptx (+ build_manifest.json)
  render     deck.pptx → PDF, PNGs, contact sheet
  qa         deck.pptx (+ manifest) → qa_report.json/.md  (exit 1 if not passed)
  run        the full loop: plan → build → render → QA → autofix → … (exit 1 if not passed)
  patch      deck.json + patches.json → patched deck.json
  review     score an agent-filled review.json (semantic visual QA)
  brand      `brand ingest template.pptx -o brands/acme` → brand theme + compatibility report
  eval       visual-quality evals: --suite regression (gate vs baseline) | holdout_v2 (sealed, release candidates only) |
             holdout_v1 (retired, history) | examples;
             --docker runs inside the pinned visual environment; --record writes evals/results/latest.json
  robustness metamorphic robustness: small content perturbations of development seeds; quality drop, catastrophic changes
  repro      render a spec twice and compare (slides, dimensions, PNG hashes, pixels, spans, metrics, QA)
  human      blind A/B preference rounds: build | serve | import | report (docs/EVALS.md)
  holdout    `holdout private`: corporate templates in .private/holdouts; `holdout external`: deck specs in
             .private/holdouts/decks (never in the repo; skipped if absent; sanitized aggregates only)
  quality    eval_report.json → quality profile (macro archetype, P10, weakest, coverage, absolute gates); --check for CI
  update     (v2.0) existing deck + new sources → reviewed in-place update of the original pptx
  reason     source-to-deck reasoning (v1.7): facts | check | ghost | trace | eval  (docs/REASONING_PROTOCOL.md)
  results    `results readme [--check]`: regenerate the README metrics block from evals/results/latest.json
Reference:
  catalog    layout library (markdown)        themes    available themes
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def _p(*a):
    print(*a, flush=True)


def cmd_ingest(a):
    from .ingest.readers import ingest, summary_markdown

    inv = ingest(a.files)
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(inv, indent=2, ensure_ascii=False, default=str), encoding="utf-8", newline="\n")
    out.with_suffix(".md").write_text(summary_markdown(inv), encoding="utf-8", newline="\n")
    _p(f"inventory: {len(inv['blocks'])} blocks, {len(inv['tables'])} tables, {len(inv['facts'])} facts → {out} (+ .md)")


def cmd_scaffold(a):
    from .core.storyline import scaffold
    from .spec import save_spec

    spec = scaffold(a.deck_type, a.title or "", a.governing_thought or "", a.theme)
    save_spec(spec, a.out)
    _p(f"skeleton with {len(spec['slides'])} slides ({spec['storyline']['framework']}) → {a.out}")


def cmd_outline(a):
    from .core.storyline import ghost_deck
    from .spec import load_spec

    txt = ghost_deck(load_spec(a.spec))
    if a.out:
        Path(a.out).write_text(txt, encoding="utf-8", newline="\n")
    _p(txt)


def cmd_lint(a):
    from .core.planner import plan
    from .spec import load_spec

    _, issues = plan(load_spec(a.spec))
    levels = {"error": 0, "warning": 0, "info": 0}
    for i in issues:
        levels[i["level"]] += 1
        if i["level"] != "info" or a.verbose:
            _p(f"{i['level']:7} {i['code']:28} {i.get('slide', '-'):6} {i['message']}")
    _p(f"\n{levels['error']} errors, {levels['warning']} warnings, {levels['info']} info")
    return 1 if levels["error"] else 0


def cmd_recommend(a):
    from .core.visual_reasoning import recommend

    ex = json.loads(Path(a.exhibit).read_text(encoding="utf-8")) if a.exhibit else {}
    for r in recommend(a.message_type, ex)[:6]:
        _p(f"{r['score']:>3}  {r['visual']:16} {'; '.join(r['reasons'])}")


def cmd_plan(a):
    from .core.planner import plan
    from .spec import load_spec

    res, issues = plan(load_spec(a.spec))
    Path(a.out).write_text(json.dumps(res, indent=2, ensure_ascii=False, default=str), encoding="utf-8", newline="\n")
    for s in res["slides"]:
        lay = (s.get("_plan") or {}).get("layout") or {}
        vis = ", ".join(v.get("chosen", "") for v in (s.get("_plan") or {}).get("visuals", []))
        _p(f"{s.get('id', ''):6} {s.get('kind', ''):12} {lay.get('id', ''):32} {vis}")
    _p(f"→ {a.out} ({sum(1 for i in issues if i['level'] == 'error')} errors, {sum(1 for i in issues if i['level'] == 'warning')} warnings)")


def cmd_build(a):
    from .core.planner import plan
    from .pptx.builder import build, save_manifest
    from .spec import load_spec

    res, _ = plan(load_spec(a.spec))
    out = Path(a.out)
    m = build(res, out)
    save_manifest(m, out.with_name("build_manifest.json"))
    _p(f"built {len(m)} slides → {out}")


def cmd_render(a):
    from .render.renderer import render

    info = render(a.pptx, a.out or Path(a.pptx).parent, dpi=a.dpi)
    _p(f"PDF {info['pdf']}\n{len(info['pngs'])} PNGs\ncontact sheet {info['contact_sheet']}")


def cmd_qa(a):
    from .design.tokens import load_profile, load_theme
    from .qa import geometry, render_checks
    from .qa import report as rep
    from .render import renderer

    pptx = Path(a.pptx)
    out = Path(a.out or pptx.parent)
    manifests = json.loads(Path(a.manifest).read_text(encoding="utf-8")) if a.manifest else []
    theme = load_theme(a.theme)  # a built-in name, a theme JSON or a brand directory
    profile = load_profile(a.profile)
    issues = geometry.check(str(pptx), manifests, theme, profile)
    metrics = {}
    if not a.no_render:
        info = renderer.render(pptx, out, dpi=a.dpi)
        ri, metrics = render_checks.check(info["pdf"], str(pptx), manifests, info["pngs"])
        issues += ri
    from pptx import Presentation

    ids = [m.get("slide_id") for m in manifests] or [f"#{i}" for i in range(1, len(Presentation(str(pptx)).slides) + 1)]
    r = rep.summarize(issues, ids, metrics=metrics)
    rep.write(r, out)
    _p(f"{'PASSED' if r['passed'] else 'FAILED'} score {r['deck_score']} errors {r['counts']['error']} warnings {r['counts']['warning']} → {out / 'qa_report.md'}")
    return 0 if r["passed"] else 1


def cmd_run(a):
    from .pipeline import run
    from .spec import load_spec

    r = run(load_spec(a.spec), a.out, max_iter=a.max_iter, do_render=not a.no_render, dpi=a.dpi, name=a.name, compose=not a.no_compose)
    _p(f"\n{'PASSED' if r['passed'] else 'FAILED'} · score {r['deck_score']} · errors {r['counts']['error']} · warnings {r['counts']['warning']}")
    from .design.tokens import theme_for
    from .pipeline import _template_not_used

    for w in _template_not_used(load_spec(a.spec).get("meta", {}), theme_for(load_spec(a.spec).get("meta", {}))):
        _p(f"WARNING {w['code']}: {w['message']}")  # v3.0: a template that was not used is said on screen, not only in the report
    if r.get("pending_actions"):
        _p("Pending author actions:")
        for x in r["pending_actions"]:
            _p(f"  - {x['slide']} {x['code']}: {x['action']}")
    _p(f"Artifacts in {a.out}: {a.name}.pptx, renders/, contact_sheet.png, qa_report.md, review.md, ghost_deck.md")
    return 0 if r["passed"] else 1


def cmd_patch(a):
    from .spec import apply_patches, load_spec, save_spec

    spec = load_spec(a.spec)
    patches = json.loads(Path(a.patches).read_text(encoding="utf-8"))
    if isinstance(patches, dict):  # review.json style {slide: {patches: [...]}}
        flat = []
        for sid, r in patches.items():
            for p in r.get("patches", []):
                flat.append({"slide": sid, **p} if "slide" not in p else p)
        patches = flat
    new, log = apply_patches(spec, patches)
    save_spec(new, a.out or a.spec)
    for line in log:
        _p(line)


def cmd_review(a):
    from .qa.report import evaluate_review

    r = evaluate_review(json.loads(Path(a.review).read_text(encoding="utf-8")))
    for sid, v in r["slides"].items():
        _p(f"{sid:6} {v['total']}/{v['max']} {'ok' if v['passed'] else 'REVISE'}")
    _p("REVIEW PASSED" if r["passed"] else "REVIEW: revise the flagged slides")
    return 0 if r["passed"] else 1


def _in_docker(argv: list[str]) -> int:
    """Re-run this command inside the reproducible visual environment (scripts/cpe-docker)."""
    import os
    import subprocess

    wrapper = Path(__file__).resolve().parents[2] / "scripts" / "cpe-docker"
    if os.environ.get("CPE_CONTAINER_IMAGE"):
        return -1  # already inside the container
    return subprocess.call([str(wrapper), *[x for x in argv if x != "--docker"]])


def cmd_eval(a):
    if a.docker:
        rc = _in_docker(sys.argv[1:])
        if rc >= 0:
            return rc
    from .evals import run_suite

    out = a.out or f"out/evals/{a.suite}"
    from .evals import DirtyEvaluation, HoldoutGuard

    try:
        r = run_suite(a.cases, out, a.baseline, update_baseline=a.update_baseline, compose=not a.no_compose, tolerance=a.tolerance, suite=a.suite, record=a.record, match=a.match,
                      release_candidate=a.release_candidate, allow_dirty=a.allow_dirty, rerun_reason=a.rerun_reason)
    except (DirtyEvaluation, HoldoutGuard) as e:
        _p(f"REFUSED     {e}")
        return 2
    for x in r["environment_notes"]:
        _p(f"NOTE        {x}")
    for x in r["regressions"]:
        _p(f"{'REGRESSION' if not a.suite.startswith('holdout') else 'weak      '}  {x}")
    for x in r["improvements"]:
        _p(f"improved    {x}")
    _p(f"\n{'PASSED' if r['passed'] else 'FAILED'} · {a.suite} · composition (archetype fitness) {r['suite_composition']} · "
       f"env {r['environment']['fingerprint']} · report {out}/eval_report.md{' · recorded in evals/results/latest.json' if a.record else ''}")
    return 0 if r["passed"] else 1


def cmd_robustness(a):
    if a.docker:
        rc = _in_docker(sys.argv[1:])
        if rc >= 0:
            return rc
    from . import robustness as rb

    out = a.out or "out/robustness"
    s = rb.run(out)
    problems = []
    if a.update_baseline:
        keep = ("variants", "seeds", "catastrophic", "catastrophic_rate", "new_visual_errors", "median_drop", "p90_drop", "p95_drop", "max_drop",
                "large_drop_rate", "font_drop_rate", "layout_change_rate", "layout_change_with_quality_drop_rate")
        rb.BASELINE.write_text(json.dumps({k: s.get(k) for k in keep}, indent=2) + "\n", encoding="utf-8", newline="\n")
    elif rb.BASELINE.exists():
        base = json.loads(rb.BASELINE.read_text(encoding="utf-8"))
        problems = rb.compare(s, base)
        for x in rb.provisional_breaches(s, base):
            _p(f"NOTE (provisional gate)  {x}")
    if a.record:
        from .evals import record_result

        record_result("robustness", {k: v for k, v in s.items() if k != "rows"})
    for c in s["catastrophic_cases"]:
        _p(f"catastrophic  {c}")
    for x in problems:
        _p(f"REGRESSION  {x}")
    _p(f"\n{'FAILED' if problems else 'PASSED'} · robustness · {s['variants']} variants of {s['seeds']} seeds · median drop {s['median_drop']} · "
       f"P90 drop {s['p90_drop']} · catastrophic {s['catastrophic']} ({s['catastrophic_rate']}) · report {out}/robustness_report.md")
    return 1 if problems else 0


def cmd_repro(a):
    if a.docker:
        rc = _in_docker(sys.argv[1:])
        if rc >= 0:
            return rc
    from .reproducibility import check
    from .spec import load_spec

    r = check(load_spec(a.spec), a.out, dpi=a.dpi)
    for k, v in r["checks"].items():
        _p(f"{'ok  ' if v else 'FAIL'} {k}")
    _p(f"worst pixel difference {r['worst_pixel_diff_share']:.6f} (tolerance {r['tolerances']['pixel_diff_share']}) · env {r['environment_fingerprint']}")
    _p("REPRODUCIBLE" if r["passed"] else "NOT REPRODUCIBLE")
    return 0 if r["passed"] else 1


def cmd_human(a):
    from . import human

    if a.human_cmd == "build":
        specs = []
        for pr in a.pair:
            name, _, rest = pr.partition("=")
            base, _, chal = rest.partition(",")
            if not (name and base and chal):
                raise SystemExit("--pair NAME=BASELINE_ROOT,CHALLENGER_ROOT")
            specs.append((name, base, chal))
        quotas = {q.split("=")[0]: int(q.split("=")[1]) for q in (a.quota or [])}
        r = human.build_round(a.out, specs, n=a.n, repeats=a.repeats, seed=a.seed, quotas=quotas or None, purpose=a.purpose or "",
                              key_out=a.key_out, private_key=not a.key_in_round, identical_controls=a.identical_controls)
        _p(f"round {a.out}: {r['pairs']} pairs {r['by_comparison']} · key {r['key']} · serve it with `cpe human serve {a.out}`")
    elif a.human_cmd == "serve":
        human.serve(a.round, port=a.port, host=a.host)
    elif a.human_cmd == "import":
        _p(f"imported {human.import_votes(a.round, a.votes)} votes")
    elif a.human_cmd == "report":
        r = human.report(a.round, a.key)
        md = human.to_markdown(r)
        # the report names versions and scores: while the key is private, so is the report
        dest = Path(a.round) if (Path(a.round) / "key.json").exists() else human.key_path(a.round, a.key).parent
        (dest / "report.md").write_text(md, encoding="utf-8", newline="\n")
        (dest / "report.json").write_text(json.dumps(r, indent=2), encoding="utf-8", newline="\n")
        (dest / "human_score_disagreements.md").write_text(human.disagreements_markdown(r.get("verdicts") or [], r["round"]), encoding="utf-8", newline="\n")
        _p(md)
        _p(f"report written to {dest}")
        if a.record:
            human.record(r) if r.get("comparisons") else human.record_status(a.round)
            _p("recorded in evals/results/latest.json (human_reference) — an independent signal, not part of the automatic score")
    elif a.human_cmd == "status":
        st = human.read_status(Path(a.round))
        _p(json.dumps(st, indent=2))
        if a.record:
            human.record_status(a.round)
    elif a.human_cmd == "build-text":
        specs = []
        for pr in a.pair:
            name, _, rest = pr.partition("=")
            base, _, chal = rest.partition(",")
            specs.append((name, base, chal))
        r = human.build_text_round(a.out, specs, context_dir=a.context, repeats=a.repeats, kind=a.kind, purpose=a.purpose or "")
        _p(f"text round {a.out}: {r['pairs']} pairs · key {r['key']} · distribute with `cpe human package`")
    elif a.human_cmd == "package":
        r = human.package_round(a.round, a.out)
        _p(f"voting package {r['zip']}: {r['images']} images, no key — send it to evaluators instead of the repository")
    elif a.human_cmd == "close":
        st = human.close_round(a.round, a.key)
        _p(f"{a.round}: voting closed, key revealed ({st['status']})")
    elif a.human_cmd == "mark-used":
        st = human.mark_used_for_calibration(a.round, a.change)
        _p(f"{a.round}: now development data ({st['status']})")
    return 0


def cmd_holdout(a):
    from . import private_holdout

    if a.which == "intake":
        for name, st in (("external", private_holdout.external_status()), ("corporate", private_holdout.corporate_status())):
            _p(f"{st['status']}" + "".join(f"\n    {k}: {v}" for k, v in st.items() if k != "status"))
        return 0
    if a.which == "external-seal":
        seal = private_holdout.seal_external()
        _p(f"sealed {len(seal['files'])} decks at {seal['sealed_at']} — do not render or open them until the frozen run")
        return 0
    if a.which in ("external-run", "corporate-run"):
        r = (private_holdout.run_external_holdout if a.which == "external-run" else private_holdout.run_corporate_unseen)(a.out, record=a.record)
        if not r.get("ran"):
            _p(r["status"] + "".join(f"\n    {k}: {v}" for k, v in r.items() if k not in ("status", "ran")))
            return 0
        _p(json.dumps({k: v for k, v in r.items() if k not in ("archetypes", "provenance")}, default=str))
        return 0
    if a.which == "external":
        r = private_holdout.run_external(a.root, a.out, record=a.record)
        _p(f"external holdout skipped: {r['reason']}" if r["status"] == "skipped" else json.dumps({k: v for k, v in r.items() if k != "archetypes"}))
        return 0
    r = private_holdout.run(a.root, a.out, record=a.record)
    if r["status"] == "skipped":
        _p(f"private holdout skipped: {r['reason']}")
        return 0
    for name, s in r["holdouts"].items():
        _p(f"{name}: " + json.dumps(s))
    _p(f"private results in {a.out or 'private_results/'} (git-ignored){' · sanitized summary recorded in evals/results/latest.json' if a.record else ''}")
    return 0


def cmd_deck(a):
    """Existing-deck ingestion, update plan (v1.8), in-place edits and patch (v1.9)."""
    from .reasoning import deck_update

    if a.deck_cmd == "ingest":
        inv = deck_update.write_ingest(a.pptx, a.out)
        _p(f"{len(inv['slides'])} slides read from {a.pptx} → {a.out}/old_deck.json, old_deck_spec.json, old_ghost.md")
        return 0
    if a.deck_cmd in ("edits", "patch"):
        from .reasoning import deck_patch

        if a.deck_cmd == "edits":
            e = deck_patch.write_edits(a.work, rederive=a.derive)
            d = e.get("derived") or {}
            _p(f"{len(e['edits'])} edits ({sum(1 for x in e['edits'] if x.get('approved'))} approved), {len(e['review'])} numbers to review; "
               f"derived figures: {d.get('added', 0)} added, {d.get('updated', 0)} updated, {d.get('unchanged', 0)} unchanged, {d.get('dropped', 0)} withdrawn → {a.work}/edits.json, edits.md")
            return 0
        r = deck_patch.write_patch(a.pptx, a.edits, a.out, mark=a.mark, accept_proposed=a.accept_proposed)
        _p(f"applied {len(r['applied'])} · failed {len(r['failed'])} · not approved {r['skipped_unapproved']} → {a.out} (patch_report.md beside it)")
        return 1 if r["failed"] else 0
    plan = deck_update.write_plan(a.work)
    _p(deck_update.plan_markdown(plan))
    return 0


def cmd_update(a):
    """(v2.0) Update an existing deck: prepare the edits, then apply the approved ones to the original."""
    from .reasoning import update

    if a.sheet or a.read_sheet:  # v2.1: the review as a spreadsheet
        from .reasoning import review_sheet

        if a.sheet:
            _p(f"review sheet → {review_sheet.write_sheet(a.sheet)} (fill the coloured columns, then --read-sheet or --apply)")
            return 0
        st = review_sheet.read_sheet(a.read_sheet)
        _p(f"approved {st['approved']} · corrected {st['corrected']} · rejected {st['rejected']} · added {st['added']} · headlines {st['headlines']} · "
           f"conflicts decided {st['conflicts']} · slides {st['slides']}" + (f" · rebuild: --rebuild {a.read_sheet} --slides {','.join(map(str, st['rebuild']))}" if st["rebuild"] else ""))
        for p in st["problems"]:
            _p(f"  ! {p}")
        return 1 if st["problems"] else 0
    if a.rebuild:
        r = update.rebuild(a.rebuild, [int(x) for x in a.slides.split(",")] if a.slides else None, render=not a.no_render)
        if not r["slides"]:
            _p(r["why"])
            return 0
        _p(f"slides {r['slides']} rebuilt in the old deck's style → {r['built']} (QA {r['qa']}); spec to edit: {r['spec']}; "
           f"approve the replace_slide edits in edits.json, then --apply")
        return 0
    if a.apply:
        if not a.out:
            _p("--apply needs -o NEW.pptx")
            return 2
        r = update.apply(a.apply, a.out, mark=a.mark, accept_derived=a.accept_derived)
        _p(f"applied {len(r['applied'])} · failed {len(r['failed'])} · not approved {r['skipped_unapproved']} · "
           f"headlines to check {len(r['rewrite_headlines'])} · left for review {len(r['left_unchanged'])} → {a.out} (update_report.md beside it)")
        return 1 if r["failed"] else 0
    if not (a.pptx and a.sources and a.out):
        _p("usage: cpe update OLD.pptx SOURCES -o WORK  |  cpe update --rebuild WORK  |  cpe update --apply WORK -o NEW.pptx")
        return 2
    m = update.prepare(a.pptx, a.sources, a.out)
    _p(f"{m['slides']} slides, {m['facts']} facts; numbers {m['plan']}; {m['edits']} edits in {a.out}/edits.json "
       f"({m['approved_kept']} approvals kept).\nNext: review {a.out}/review.xlsx (approve, correct, dismiss), then "
       f"`cpe update --apply {a.out} -o new.pptx --mark`")
    return 0


def cmd_reason(a):
    """Source-to-deck reasoning artifacts (v1.7): facts, checks, ghost deck, trace, benchmark."""
    from .reasoning import benchmark, checks, facts, ghost, graph

    if a.reason_cmd == "conflicts":  # v2.1: the ranked list; dismiss or resolve one
        from .reasoning.conflicts import review_conflict

        if a.dismiss or a.resolve:
            c = review_conflict(a.work, a.dismiss or a.resolve, dismiss=bool(a.dismiss), why=a.why or "", use=a.use)
            _p(f"{c['id']}: {'dismissed' if a.dismiss else 'resolved'} — {c.get('dismissed') or c.get('resolution')}")
            return 0
        cs = json.loads(Path(a.work, "fact_conflicts.json").read_text(encoding="utf-8")).get("conflicts") or []
        for c in cs:
            state = "dismissed" if c.get("dismissed") else "resolved" if c.get("resolution") else "open"
            vers = " / ".join(f"{x['value']:g} ({x['source']})" for x in c["facts"])
            _p(f"{c['id']} [{c.get('priority', '-')}] {state:9} {c['type']}: {vers}")
        return 0
    if a.reason_cmd == "facts":
        fm = facts.write_fact_model(a.sources, a.out)
        _p(f"{fm['stats']['facts']} facts from {fm['stats']['sources']} sources ({fm['stats']['by_type']}) → {a.out}/facts.json")
        return 0
    if a.reason_cmd == "check":
        if a.enrich:
            ghost.enrich_evidence(a.work)
        r = checks.check_work(a.work, a.sources)
        Path(a.work, "reasoning_report.json").write_text(json.dumps(r, indent=2, ensure_ascii=False), encoding="utf-8", newline="\n")
        Path(a.work, "reasoning_report.md").write_text(checks.report_markdown(r), encoding="utf-8", newline="\n")
        graph.save_graph(a.work)
        ghost.write_ghost(a.work)
        _p(checks.report_markdown(r))
        return 1 if r["status"] == "blocked" else 0
    if a.reason_cmd == "analyze":
        from .reasoning.analysis import run_analysis

        e = run_analysis(a.work, a.script, a.sources)
        _p(f"analysis {e['name']}: {len(e['outputs'])} tables ({', '.join(e['outputs'])}), reproducible {e['reproducible']} → re-run `cpe reason facts` to cite them")
        return 0 if e["reproducible"] else 1
    if a.reason_cmd == "ghost":
        _p(ghost.write_ghost(a.work).read_text(encoding="utf-8"))
        return 0
    if a.reason_cmd == "trace":
        t = graph.trace(a.work, a.text)
        _p(json.dumps(t, indent=2, ensure_ascii=False) if a.json else graph.trace_markdown(t))
        return 0 if t.get("match") else 1
    if a.reason_cmd == "eval-text":
        r = benchmark.evaluate_text(a.case, a.storyline)
        _p(json.dumps(r, indent=2, ensure_ascii=False))
        return 0
    if a.reason_cmd == "eval":
        from . import __version__
        from .reasoning import PROTOCOL_VERSION

        meta = {"engine": __version__, "reasoning_protocol": PROTOCOL_VERSION, "model": a.model, "skill": a.skill, "author": a.author}
        r = benchmark.evaluate(a.case, a.work, meta)
        Path(a.work, "s2d_eval.json").write_text(json.dumps(r, indent=2, ensure_ascii=False), encoding="utf-8", newline="\n")
        Path(a.work, "s2d_eval.md").write_text(benchmark.to_markdown(r), encoding="utf-8", newline="\n")
        _p(benchmark.to_markdown(r))
        return 1 if r["factuality"]["status"] == "FAIL" else 0
    return 2


def cmd_quality(a):
    """Quality profile of an eval report: distribution, archetype health, absolute gates."""
    from . import quality

    rep = json.loads(Path(a.report).read_text(encoding="utf-8"))
    prof = quality.profile(rep["cases"], deck_mean=rep.get("suite_composition"))
    for line in quality.profile_markdown(prof):
        _p(line)
    fails = quality.gate_failures(prof)
    for f in fails:
        _p(f"GATE FAILED  {f}")
    nh = prof["not_healthy"]
    _p(f"{'FAILED' if fails else 'PASSED'} · macro {prof['macro_archetype_score']} · P10 {prof['distribution']['p10']} · weakest {prof['weakest_archetype']} "
       f"{prof['weakest_archetype_score']} · not healthy (provisional): {', '.join(nh) or 'none'}")
    return 1 if (fails and a.check) else 0


def cmd_results(a):
    from .results_report import update_readme, verify_provenance

    if a.what == "verify":
        problems = verify_provenance()
        for x in problems:
            _p(f"NOT RELEASE TRUTH  {x}")
        _p("provenance " + ("FAILED" if problems else "ok: every release signal was evaluated on a clean commit with the current engine inputs"))
        return 1 if problems else 0
    changed = update_readme(check=a.check)
    if a.check and changed:
        _p("README metrics block is stale: run `scripts/cpe results readme` and commit (numbers come from evals/results/latest.json)")
        return 1
    _p("README metrics block " + ("updated" if changed else "up to date"))
    return 0


def cmd_brand_fidelity(a):
    from .brand.fidelity import fidelity_markdown, from_run_dir

    f = from_run_dir(a.run, a.brand)
    if "typography" not in f:
        _p(f["note"])
        return
    print(fidelity_markdown(f))
    _p(f"→ {a.run}/brand_fidelity.json, {a.run}/corporate_usage.md")


def cmd_brand(a):
    from .brand.ingest import ingest

    r = ingest(a.template, a.out, name=a.name, base_theme=a.base_theme)
    f = r["fonts"]
    _p(f"brand '{r['brand']}' → {a.out}/theme.json  (use it with \"meta\": {{\"brand\": \"{a.out}\"}})")
    _p(f"slide size {r['slide_size']['width_in']}×{r['slide_size']['height_in']} in · masters used: {r['masters_used']} · base layout: {r['base_layout']}")
    for role, x in f.items():
        if x.get("font"):
            _p(f"font {role:7} {x['font']:24} {x['status']:28} measured with {x['measure_family']} ({x['measurement']})")
    for u in r["unsupported"]:
        _p(f"unsupported: {u}")
    _p(f"report: {a.out}/compatibility.md")


def cmd_measure(a):
    """Composition metrics of an existing run directory (any engine version)."""
    from .design.tokens import theme_for
    from .qa.composition import measure_deck

    d = Path(a.run_dir)
    resolved = json.loads((d / "resolved.json").read_text(encoding="utf-8"))
    manifests = json.loads((d / "build_manifest.json").read_text(encoding="utf-8"))
    pngs = sorted(str(p) for p in (d / "renders").glob("slide-*.png"))
    comps = measure_deck(str(d / f"{a.name}.pdf"), pngs, resolved, manifests, theme_for(resolved.get("meta", {})))
    # integrity from the CURRENT checks on the stored render (not the run's own, possibly older, QA)
    from .design.tokens import theme_for as _tf
    from .qa import geometry, render_checks
    from .qa.composition import apply_integrity

    issues = geometry.check(str(d / f"{a.name}.pptx"), manifests, _tf(resolved.get("meta", {})), resolved.get("_profile") or {})
    issues += render_checks.check(str(d / f"{a.name}.pdf"), str(d / f"{a.name}.pptx"), manifests, pngs)[0]
    apply_integrity(comps, issues, manifests)
    out = {"deck_score": round(sum(c.score for c in comps) / len(comps), 1) if comps else None, "slides": [c.to_dict() for c in comps]}
    (d / "composition_measure.json").write_text(json.dumps(out, indent=2), encoding="utf-8", newline="\n")
    for c in comps:
        _p(f"{c.slide_id:8} {c.score:5.1f} {', '.join(c.flags)}")
    _p(f"deck composition {out['deck_score']}")


def cmd_catalog(a):
    from .layout.engine import catalog_markdown

    txt = catalog_markdown()
    if a.out:
        Path(a.out).write_text(txt, encoding="utf-8", newline="\n")
    _p(txt)


def cmd_themes(a):
    from .design.tokens import available_themes

    _p("\n".join(available_themes()))


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):  # a Windows console or a redirect may not be UTF-8: never crash on "→" or "€"
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(prog="cpe", description="Consulting Presentation Engine")
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("ingest"); s.add_argument("files", nargs="+"); s.add_argument("-o", "--out", default="inventory.json"); s.set_defaults(f=cmd_ingest)
    s = sub.add_parser("scaffold"); s.add_argument("--deck-type", required=True); s.add_argument("--title"); s.add_argument("--governing-thought"); s.add_argument("--theme", default="meridian"); s.add_argument("-o", "--out", default="deck.json"); s.set_defaults(f=cmd_scaffold)
    s = sub.add_parser("outline"); s.add_argument("spec"); s.add_argument("-o", "--out"); s.set_defaults(f=cmd_outline)
    s = sub.add_parser("lint"); s.add_argument("spec"); s.add_argument("-v", "--verbose", action="store_true"); s.set_defaults(f=cmd_lint)
    s = sub.add_parser("recommend"); s.add_argument("message_type"); s.add_argument("--exhibit"); s.set_defaults(f=cmd_recommend)
    s = sub.add_parser("plan"); s.add_argument("spec"); s.add_argument("-o", "--out", default="resolved.json"); s.set_defaults(f=cmd_plan)
    s = sub.add_parser("build"); s.add_argument("spec"); s.add_argument("-o", "--out", default="out/deck.pptx"); s.set_defaults(f=cmd_build)
    s = sub.add_parser("render"); s.add_argument("pptx"); s.add_argument("-o", "--out"); s.add_argument("--dpi", type=int, default=110); s.set_defaults(f=cmd_render)
    s = sub.add_parser("qa"); s.add_argument("pptx"); s.add_argument("--manifest"); s.add_argument("--theme", default="meridian"); s.add_argument("--profile", default="standard"); s.add_argument("-o", "--out"); s.add_argument("--no-render", action="store_true"); s.add_argument("--dpi", type=int, default=110); s.set_defaults(f=cmd_qa)
    s = sub.add_parser("run"); s.add_argument("spec"); s.add_argument("-o", "--out", default="out"); s.add_argument("--max-iter", type=int, default=3); s.add_argument("--no-render", action="store_true"); s.add_argument("--dpi", type=int, default=110); s.add_argument("--name", default="deck"); s.add_argument("--no-compose", action="store_true", help="skip the composition engine"); s.set_defaults(f=cmd_run)
    s = sub.add_parser("patch"); s.add_argument("spec"); s.add_argument("patches"); s.add_argument("-o", "--out"); s.set_defaults(f=cmd_patch)
    s = sub.add_parser("review"); s.add_argument("review"); s.set_defaults(f=cmd_review)
    s = sub.add_parser("eval"); s.add_argument("cases", nargs="?", default=None, help="case directory (default: the suite's)"); s.add_argument("--suite", default="regression", choices=["regression", "holdout_v1", "holdout_v2", "examples"])
    s.add_argument("-o", "--out"); s.add_argument("--baseline", default=None); s.add_argument("--update-baseline", action="store_true"); s.add_argument("--no-compose", action="store_true"); s.add_argument("--tolerance", type=float, default=2.0)
    s.add_argument("--match", help="only cases whose name contains this text (development; cannot record)")
    s.add_argument("--release-candidate", action="store_true", help="required for the sealed holdout v2 (frozen, clean engine)")
    s.add_argument("--allow-dirty", action="store_true", help="record a development result from a dirty tree (never release truth)")
    s.add_argument("--rerun-reason", help="holdout v2 only: why a second run on another engine commit is not tuning")
    s.add_argument("--docker", action="store_true", help="run inside the pinned visual environment (scripts/cpe-docker)"); s.add_argument("--record", action="store_true", help="write the result into evals/results/latest.json"); s.set_defaults(f=cmd_eval)
    s = sub.add_parser("robustness"); s.add_argument("-o", "--out"); s.add_argument("--update-baseline", action="store_true"); s.add_argument("--record", action="store_true")
    s.add_argument("--docker", action="store_true"); s.set_defaults(f=cmd_robustness)
    s = sub.add_parser("repro"); s.add_argument("spec"); s.add_argument("-o", "--out"); s.add_argument("--dpi", type=int, default=80); s.add_argument("--docker", action="store_true"); s.set_defaults(f=cmd_repro)
    s = sub.add_parser("human"); hs = s.add_subparsers(dest="human_cmd", required=True)
    h = hs.add_parser("build"); h.add_argument("-o", "--out", required=True); h.add_argument("--pair", action="append", required=True, help="NAME=BASELINE_ROOT,CHALLENGER_ROOT (run dirs matched by deck + slide id)")
    h.add_argument("--n", type=int, default=40); h.add_argument("--repeats", type=int, default=2); h.add_argument("--seed", type=int, default=7)
    h.add_argument("--quota", action="append", help="ARCHETYPE=K: sample K pairs of this archetype first (the rest are controls)"); h.add_argument("--purpose")
    h.add_argument("--key-out", help="where the private key goes (default .private/human_reference/keys/<round>/key.json)")
    h.add_argument("--key-in-round", action="store_true", help="legacy (≤ v1.4): write key.json inside the round")
    h.add_argument("--identical-controls", type=int, default=0, help="add K identical-image control pairs (evaluator noise)")
    h.set_defaults(f=cmd_human)
    h = hs.add_parser("serve"); h.add_argument("round"); h.add_argument("--port", type=int, default=8765); h.add_argument("--host", default="127.0.0.1"); h.set_defaults(f=cmd_human)
    h = hs.add_parser("import"); h.add_argument("round"); h.add_argument("votes"); h.set_defaults(f=cmd_human)
    h = hs.add_parser("report"); h.add_argument("round"); h.add_argument("--key", help="private key.json of the round"); h.add_argument("--record", action="store_true"); h.set_defaults(f=cmd_human)
    h = hs.add_parser("build-text", help="blind A/B of storylines or deck outlines (text, before rendering)")
    h.add_argument("-o", "--out", required=True); h.add_argument("--pair", action="append", required=True, help="NAME=DIR_A,DIR_B (each with <case>.md)")
    h.add_argument("--context", help="dir with <case>.md: business question + short source summary"); h.add_argument("--kind", choices=["storyline", "outline"], default="storyline")
    h.add_argument("--repeats", type=int, default=2); h.add_argument("--purpose"); h.set_defaults(f=cmd_human)
    h = hs.add_parser("package"); h.add_argument("round"); h.add_argument("-o", "--out", required=True); h.set_defaults(f=cmd_human)
    h = hs.add_parser("close"); h.add_argument("round"); h.add_argument("--key"); h.set_defaults(f=cmd_human)
    h = hs.add_parser("status"); h.add_argument("round"); h.add_argument("--record", action="store_true"); h.set_defaults(f=cmd_human)
    h = hs.add_parser("mark-used"); h.add_argument("round"); h.add_argument("--change", required=True, help="what the votes were used to change"); h.set_defaults(f=cmd_human)
    s = sub.add_parser("holdout"); s.add_argument("which", choices=["private", "external", "intake", "external-seal", "external-run", "corporate-run"]); s.add_argument("--root"); s.add_argument("-o", "--out"); s.add_argument("--record", action="store_true"); s.set_defaults(f=cmd_holdout)
    s = sub.add_parser("deck", help="existing deck: ingest structure and conventions, plan the update against new facts (v1.8)")
    ds = s.add_subparsers(dest="deck_cmd", required=True)
    d = ds.add_parser("ingest"); d.add_argument("pptx"); d.add_argument("-o", "--out", required=True); d.set_defaults(f=cmd_deck)
    d = ds.add_parser("stale", help="old deck numbers vs the new fact model: current / outdated / untraced"); d.add_argument("work"); d.set_defaults(f=cmd_deck)
    d = ds.add_parser("edits", help="(v1.9) proposed in-place edits from the update plan, to approve"); d.add_argument("work")
    d.add_argument("--derive", action="store_true", help="(v2.0) keep the reviewed edits.json and recompute totals, ratios and repeated figures from the approved ones")
    d.set_defaults(f=cmd_deck)
    d = ds.add_parser("patch", help="(v1.9) apply approved edits to the ORIGINAL pptx; everything else stays as it was")
    d.add_argument("pptx"); d.add_argument("edits"); d.add_argument("-o", "--out", required=True)
    d.add_argument("--mark", action="store_true", help="highlight changed text for review")
    d.add_argument("--accept-proposed", action="store_true", help="apply every proposed edit, approved or not (dry run)"); d.set_defaults(f=cmd_deck)
    s = sub.add_parser("update", help="(v2.0) update an existing deck: OLD.pptx SOURCES -o WORK, then --apply WORK -o NEW.pptx")
    s.add_argument("pptx", nargs="?"); s.add_argument("sources", nargs="?"); s.add_argument("-o", "--out")
    s.add_argument("--apply", metavar="WORK", help="apply the approved edits of WORK to the original deck")
    s.add_argument("--mark", action="store_true", help="highlight changed text for review")
    s.add_argument("--accept-derived", action="store_true", help="also apply totals / ratios recomputed from approved values")
    s.add_argument("--sheet", metavar="WORK", help="(v2.1) write WORK/review.xlsx: approve, correct, dismiss in a spreadsheet")
    s.add_argument("--read-sheet", metavar="WORK", help="(v2.1) read WORK/review.xlsx back into the edits and conflicts")
    s.add_argument("--rebuild", metavar="WORK", help="build the slides whose message no longer holds in the old deck's style (replace_slide edits)")
    s.add_argument("--slides", help="with --rebuild: which old slides (e.g. 4,6)")
    s.add_argument("--no-render", action="store_true", help="with --rebuild: skip rendering")
    s.set_defaults(f=cmd_update)
    s = sub.add_parser("reason", help="source-to-deck reasoning artifacts (v1.7)"); rs = s.add_subparsers(dest="reason_cmd", required=True)
    r = rs.add_parser("facts"); r.add_argument("sources"); r.add_argument("-o", "--out", required=True); r.set_defaults(f=cmd_reason)
    r = rs.add_parser("conflicts", help="(v2.1) ranked conflicts between sources; --dismiss / --resolve one")
    r.add_argument("work"); r.add_argument("--dismiss", metavar="ID"); r.add_argument("--resolve", metavar="ID")
    r.add_argument("--why", help="the reason (required to dismiss)"); r.add_argument("--use", metavar="FACT", help="with --resolve: the fact whose value is used")
    r.set_defaults(f=cmd_reason)
    r = rs.add_parser("check"); r.add_argument("work"); r.add_argument("--sources"); r.add_argument("--enrich", action="store_true",
                                                                                                    help="fill deck.json evidence from the cited facts first"); r.set_defaults(f=cmd_reason)
    r = rs.add_parser("ghost"); r.add_argument("work"); r.set_defaults(f=cmd_reason)
    r = rs.add_parser("analyze", help="run an analysis script over row-level sources; its CSV outputs become citable facts")
    r.add_argument("work"); r.add_argument("script"); r.add_argument("--sources", required=True); r.set_defaults(f=cmd_reason)
    r = rs.add_parser("trace"); r.add_argument("work"); r.add_argument("text"); r.add_argument("--json", action="store_true"); r.set_defaults(f=cmd_reason)
    r = rs.add_parser("eval-text", help="any system's storyline.md vs the case reference (conclusions, traps, untraced numbers)")
    r.add_argument("case"); r.add_argument("storyline"); r.set_defaults(f=cmd_reason)
    r = rs.add_parser("eval"); r.add_argument("case"); r.add_argument("work"); r.add_argument("--model", default="unrecorded"); r.add_argument("--skill", default="unrecorded")
    r.add_argument("--author", default="unrecorded", help="who produced the work (agent/session/person)"); r.set_defaults(f=cmd_reason)
    s = sub.add_parser("quality"); s.add_argument("report", help="eval_report.json"); s.add_argument("--check", action="store_true", help="exit 1 on an enforced gate failure")
    s.set_defaults(f=cmd_quality)
    s = sub.add_parser("results"); s.add_argument("what", choices=["readme", "verify"]); s.add_argument("--check", action="store_true"); s.set_defaults(f=cmd_results)
    s = sub.add_parser("brand"); bs = s.add_subparsers(dest="brand_cmd", required=True)
    s2 = bs.add_parser("ingest"); s2.add_argument("template"); s2.add_argument("-o", "--out", required=True); s2.add_argument("--name"); s2.add_argument("--base-theme", default="meridian"); s2.set_defaults(f=cmd_brand)
    s2 = bs.add_parser("fidelity", help="corporate usage per slide (native / adaptive / engine fallback) and brand-fidelity metrics of a run folder")
    s2.add_argument("run"); s2.add_argument("--brand", help="brand directory, if it moved since the run"); s2.set_defaults(f=cmd_brand_fidelity)
    s = sub.add_parser("measure"); s.add_argument("run_dir"); s.add_argument("--name", default="deck"); s.set_defaults(f=cmd_measure)
    s = sub.add_parser("catalog"); s.add_argument("-o", "--out"); s.set_defaults(f=cmd_catalog)
    s = sub.add_parser("themes"); s.set_defaults(f=cmd_themes)
    a = ap.parse_args(argv)
    return a.f(a) or 0


if __name__ == "__main__":
    sys.exit(main())
